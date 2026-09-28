"""Scopus author profiles: the office's workbook, linked to the people in it.

The workbook is one sheet per faculty member (plus "General", which in the
file the college has is itself a profile): "Scopus ID" in A1/B1, a Metric |
Value table on the left and a second one far to the right, a Year |
Publications table under the right one, and the author's document list. Every
workbook here is built in the test in that shape; the real one names real
people and stays out of the repository.
"""
from __future__ import annotations

import io
import tempfile
from datetime import datetime
from io import StringIO
from pathlib import Path

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, TestCase

from core.hod import MONEY_KEYS
from core.models import AuditLog, FacultyMaster, Role, User

DOC_HEADER = [
    "Year", "Title", "Journal", "Document Type", "Citations", "DOI", "EID",
    "Publisher", "ISSN", "Volume", "Issue", "Pages", "Open Access",
    "Total Authors", "Author Position", "All Authors", "Affiliations", "Scopus Link",
]


def _row(left=(), right=()):
    """A row with the left table in A.. and the right one in V/W."""
    return list(left) + [None] * (21 - len(left)) + list(right)


def fill_profile(ws, *, scopus_id, pubs=26.0, cites=166.0, h=8.0, blank_second_row=False,
                 by_year=((2013.0, 1.0), (2015.0, 1.0), (2025.0, 3.0)), docs=None):
    ws.append(["Scopus ID", scopus_id])
    if blank_second_row:
        ws.append([])
    ws.append(_row(["Metric", "Value"], ["Metric", "Value"]))
    ws.append(_row(["Author Name", "N/A"], ["Total Publications", pubs]))
    ws.append(_row(["Scopus Author ID", scopus_id], ["Total Citations", cites]))
    ws.append(_row(["Affiliation", "N/A"], ["Average Citations per Paper", cites / pubs]))
    ws.append(["Total Publications", pubs])
    ws.append(_row(["Total Citations", cites], ["Year", "Publications"]))
    pairs = list(by_year)
    ws.append(_row(["H-Index", h], pairs.pop(0) if pairs else ()))
    docs = docs if docs is not None else [
        [2025.0, "Optical glass for LEDs", "Optics and Laser Technology", "Article", 9.0,
         "10.1016/j.optlastec.2024.112111", "2-s2.0-85209087350", None, 303992.0, 182.0,
         None, None, "No", 0.0, None, None, "Saveetha Engineering College", "Scopus"],
        [2022.0, "Fuzzy UPQC", "Lecture Notes in Electrical Engineering", "Conference Paper",
         5.0, "10.1007/978-981-16-9239-0_1", "2-s2.0-85131925904", None, 18761100.0, 852.0,
         None, datetime(2026, 1, 13), "No", 0.0, None, None,
         "Saveetha Engineering College", "Scopus"],
    ]
    ws.append(_row(DOC_HEADER, pairs.pop(0) if pairs else ()))
    for doc in docs:
        ws.append(_row(doc, pairs.pop(0) if pairs else ()))
    for pair in pairs:
        ws.append(_row((), pair))


def profiles_workbook(sheets) -> bytes:
    """`sheets` is [(sheet name, fill_profile kwargs)]."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, kw in sheets:
        fill_profile(wb.create_sheet(name), **kw)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


JOYAL = "57983494200"
GENERAL = "57527550200"

TWO_SHEETS = [
    ("General", {"scopus_id": 57527550200.0, "pubs": 49.0, "cites": 170.0, "h": 6.0,
                 "blank_second_row": True}),
    ("Mr. S. Joyal Isac", {"scopus_id": 57983494200.0}),
]


def _person(email, name, *, role=Role.FACULTY, **extra):
    return User.objects.create_user(email=email, password="pass", name=name, role=role, **extra)


class ScopusIdTests(TestCase):
    def test_every_way_the_id_arrives_comes_out_as_digits(self):
        from core.services.scopus_profiles import normalize_scopus_id as n

        self.assertEqual(n(57527550200.0), GENERAL)
        self.assertEqual(n("57527550200.0"), GENERAL)
        self.assertEqual(n(57527550200), GENERAL)
        self.assertEqual(n(" 57527550200 "), GENERAL)
        self.assertEqual(
            n("https://www.scopus.com/authid/detail.uri?authorId=57527550200"), GENERAL
        )
        for nothing in (None, "", "N/A", "not an id"):
            self.assertIsNone(n(nothing), nothing)


class ReadProfilesTests(TestCase):
    def _read(self, sheets=TWO_SHEETS):
        from core.services.scopus_profiles import read_profiles

        return read_profiles(io.BytesIO(profiles_workbook(sheets)))

    def test_one_profile_per_sheet_including_general(self):
        profiles = self._read()
        self.assertEqual(
            [(p.sheet, p.scopus_id) for p in profiles],
            [("General", GENERAL), ("Mr. S. Joyal Isac", JOYAL)],
        )

    def test_the_metrics_are_read_from_both_tables(self):
        joyal = self._read()[1]
        self.assertEqual(
            (joyal.total_publications, joyal.total_citations, joyal.h_index), (26, 166, 8)
        )
        self.assertIsNone(joyal.author_name, "N/A is not a name")
        self.assertIsNone(joyal.affiliation)
        self.assertAlmostEqual(joyal.metrics["Average Citations per Paper"], 166 / 26)
        self.assertEqual(
            joyal.metrics["Publications by year"], {"2013": 1, "2015": 1, "2025": 3}
        )

    def test_the_document_list_is_kept_with_its_cells_repaired(self):
        docs = self._read()[1].documents
        self.assertEqual(len(docs), 2)
        self.assertEqual(docs[0]["Title"], "Optical glass for LEDs")
        self.assertEqual(docs[0]["Year"], 2025)
        self.assertEqual(docs[0]["Citations"], 9)
        # Excel turned 0030-3992 into the number 303992; the check digit
        # confirms which zeros it lost.
        self.assertEqual(docs[0]["ISSN"], "0030-3992")
        # A page range Excel read as a date is kept as text, not crashed on.
        self.assertIsInstance(docs[1]["Pages"], str)

    def test_a_workbook_with_no_profile_in_it_is_refused(self):
        from core.services.scopus_profiles import ProfileWorkbookError, read_profiles

        wb = openpyxl.Workbook()
        wb.active.append(["Nothing", "here"])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(ProfileWorkbookError):
            read_profiles(io.BytesIO(buf.getvalue()))

    def test_something_that_is_not_a_workbook_is_refused(self):
        from core.services.scopus_profiles import ProfileWorkbookError, read_profiles

        with self.assertRaises(ProfileWorkbookError):
            read_profiles(io.BytesIO(b"a,b\n1,2\n"))


class ImportProfilesTests(TestCase):
    def _import(self, sheets=TWO_SHEETS):
        from core.services.scopus_profiles import import_profiles, read_profiles

        return import_profiles(read_profiles(io.BytesIO(profiles_workbook(sheets))))

    def test_a_profile_is_linked_to_the_account_carrying_its_id_however_it_was_stored(self):
        from core.models import ScopusProfile

        joyal = _person("joyal@test.edu", "Joyal Isac S", scopus_author_id="57983494200.0")
        result = self._import()
        self.assertEqual((result["created"], result["updated"]), (2, 0))
        self.assertEqual(ScopusProfile.objects.get(scopus_id=JOYAL).user_id, joyal.id)
        self.assertEqual([u["scopus_id"] for u in result["unmatched"]], [GENERAL])

    def test_the_link_can_come_from_the_profile_url_alone(self):
        from core.models import ScopusProfile

        joyal = _person(
            "joyal@test.edu", "Joyal Isac S",
            scopus_author_url=f"https://www.scopus.com/authid/detail.uri?authorId={JOYAL}",
        )
        self._import()
        self.assertEqual(ScopusProfile.objects.get(scopus_id=JOYAL).user_id, joyal.id)

    def test_the_faculty_master_s_scopus_id_links_through_the_staff_id(self):
        from core.models import ScopusProfile

        general = _person("gen@test.edu", "Someone General", staff_id="TSEE001")
        FacultyMaster.objects.create(name="Someone General", staff_id="TSEE001",
                                     scopus_author_id="57527550200.0")
        self._import()
        self.assertEqual(ScopusProfile.objects.get(scopus_id=GENERAL).user_id, general.id)

    def test_two_accounts_on_one_id_link_neither_and_say_so(self):
        from core.models import ScopusProfile

        _person("a@test.edu", "A", scopus_author_id=JOYAL)
        _person("b@test.edu", "B", scopus_author_id=JOYAL)
        result = self._import()
        self.assertIsNone(ScopusProfile.objects.get(scopus_id=JOYAL).user_id)
        self.assertEqual([a["scopus_id"] for a in result["ambiguous"]], [JOYAL])

    def test_re_importing_updates_the_same_rows(self):
        from core.models import ScopusProfile

        self._import()
        first = ScopusProfile.objects.get(scopus_id=JOYAL).imported_at
        again = self._import([("Mr. S. Joyal Isac", {"scopus_id": 57983494200.0, "cites": 200.0})])
        self.assertEqual((again["created"], again["updated"]), (0, 1))
        self.assertEqual(ScopusProfile.objects.count(), 2)
        row = ScopusProfile.objects.get(scopus_id=JOYAL)
        self.assertEqual(row.total_citations, 200)
        self.assertGreaterEqual(row.imported_at, first)
        self.assertEqual(row.source_sheet, "Mr. S. Joyal Isac")
        self.assertEqual(len(row.documents), 2)

    def test_the_command_prints_the_counts_and_the_unmatched_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "profiles.xlsx"
            path.write_bytes(profiles_workbook(TWO_SHEETS))
            out = StringIO()
            call_command("import_scopus_profiles", str(path), stdout=out)
        text = out.getvalue()
        self.assertIn("created: 2", text)
        self.assertIn("unmatched: 2", text)
        self.assertIn(GENERAL, text)
        self.assertIn(JOYAL, text)


class OnePersonOneProfileTests(TestCase):
    """A person is linked through their own Scopus id, and only through the
    faculty master's when they carry none -- never through both."""

    def _import(self, sheets=TWO_SHEETS):
        from core.services.scopus_profiles import import_profiles, read_profiles

        return import_profiles(read_profiles(io.BytesIO(profiles_workbook(sheets))))

    def test_a_stale_master_id_does_not_give_an_account_a_second_profile(self):
        from core.models import ScopusProfile
        from core.services.scopus_profiles import department_totals

        person = _person("p@test.edu", "Joyal Isac S", department="EEE",
                         scopus_author_id=JOYAL, staff_id="TSEE001")
        FacultyMaster.objects.create(name="Joyal Isac S", staff_id="TSEE001",
                                     scopus_author_id=GENERAL)  # stale
        self._import()
        self.assertEqual(ScopusProfile.objects.get(scopus_id=JOYAL).user_id, person.id)
        self.assertIsNone(ScopusProfile.objects.get(scopus_id=GENERAL).user_id)
        rows = department_totals()
        self.assertEqual([(r["department"], r["people_with_profile"]) for r in rows], [("EEE", 1)])

    def test_a_corrected_id_is_what_shows_not_the_link_made_at_import(self):
        person = _person("p@test.edu", "Somebody", scopus_author_id=JOYAL)
        self._import()
        person.scopus_author_id = "99999999999"  # corrected; no sheet for it yet
        person.save()
        client = Client()
        client.force_login(person)
        self.assertIsNone(client.get("/api/me/scopus").json()["profile"])

    def test_the_master_is_matched_on_staff_id_the_same_way_everywhere(self):
        from core.services.scopus_profiles import ids_for

        person = _person("p@test.edu", "Spaced Id", staff_id="TSEE 009")
        FacultyMaster.objects.create(name="Spaced Id", staff_id="tsee009",
                                     scopus_author_id="57527550200.0")
        self.assertEqual(ids_for(person), {GENERAL})

    def test_a_deactivated_account_is_not_counted_in_its_department(self):
        from core.services.scopus_profiles import department_totals

        _person("gone@test.edu", "Left Last Year", department="EEE",
                scopus_author_id=GENERAL, active=False)
        _person("here@test.edu", "Still Here", department="EEE", scopus_author_id=JOYAL)
        self._import()
        (row,) = department_totals()
        self.assertEqual((row["people_with_profile"], row["citations"]), (1, 166))


class ProfileUploadTests(TestCase):
    def setUp(self):
        self.client = Client()

    def _upload(self, content, name="profiles.xlsx"):
        return self.client.post(
            "/api/admin/scopus-profiles/import",
            data={"file": SimpleUploadedFile(name, content)},
        )

    def test_the_office_uploads_and_nobody_else_may(self):
        content = profiles_workbook(TWO_SHEETS)
        for role, allowed in (
            (Role.SUPER_ADMIN, True), (Role.RESEARCH_CELL, True),
            (Role.RESEARCH_COORDINATOR, True), (Role.FACULTY, False),
            (Role.HOD, False), (Role.PRINCIPAL, False), (Role.FINANCE, False),
        ):
            self.client.force_login(_person(f"{role.lower()}@test.edu", role.title(), role=role))
            r = self._upload(content)
            self.assertEqual(r.status_code == 200, allowed, f"{role}: {r.status_code}")
        self.assertTrue(AuditLog.objects.filter(action="SCOPUS_PROFILES_IMPORT").exists())

    def test_a_file_that_is_not_a_profile_workbook_is_a_400(self):
        self.client.force_login(_person("cell@test.edu", "Cell", role=Role.RESEARCH_CELL))
        self.assertEqual(self._upload(b"a,b\n1,2\n", "probe.csv").status_code, 400)


class VerificationReportTests(TestCase):
    """The office's list of things to fix, three kinds of them."""

    def setUp(self):
        from core.services.scopus_profiles import import_profiles, read_profiles

        self.office = _person("cell@test.edu", "Cell", role=Role.RESEARCH_CELL)
        # On the sheet named for him, but his account carries another id.
        self.joyal = _person("joyal@test.edu", "Joyal Isac S", department="EEE",
                             scopus_author_id="11111111111")
        # No Scopus id anywhere on the account; the faculty master has one.
        self.bare = _person("bare@test.edu", "No Id Yet", department="EEE", staff_id="TSEE009")
        FacultyMaster.objects.create(name="No Id Yet", staff_id="TSEE009",
                                     scopus_author_id="22222222222")
        # Has one, and nothing is wrong with it.
        _person("fine@test.edu", "Fine Person", department="EEE", scopus_author_id="33333")
        import_profiles(read_profiles(io.BytesIO(profiles_workbook(TWO_SHEETS))))
        self.client = Client()

    def _report(self, user=None):
        self.client.force_login(user or self.office)
        return self.client.get("/api/admin/scopus-profiles/verification")

    def test_profiles_whose_id_matches_no_account(self):
        body = self._report().json()
        self.assertEqual(
            sorted(p["scopus_id"] for p in body["profiles_without_account"]), [GENERAL, JOYAL]
        )

    def test_faculty_accounts_with_no_scopus_id_with_the_master_s_hint(self):
        rows = {r["name"]: r for r in self._report().json()["faculty_without_scopus"]}
        self.assertIn("No Id Yet", rows)
        self.assertEqual(rows["No Id Yet"]["faculty_master_scopus_id"], "22222222222")
        self.assertNotIn("Fine Person", rows)
        self.assertNotIn("Cell", rows, "office accounts are not faculty")

    def test_an_account_whose_id_differs_from_the_sheet_named_for_them(self):
        rows = self._report().json()["name_mismatches"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["user_id"], self.joyal.id)
        self.assertEqual(rows[0]["stored_scopus_id"], "11111111111")
        self.assertEqual(rows[0]["sheet_scopus_id"], JOYAL)
        self.assertEqual(rows[0]["sheet"], "Mr. S. Joyal Isac")

    def test_only_the_office_reads_it(self):
        self.assertEqual(self._report(self.joyal).status_code, 403)


class WhereProfilesAreShownTests(TestCase):
    def setUp(self):
        from core.services.scopus_profiles import import_profiles, read_profiles

        self.joyal = _person("joyal@test.edu", "Joyal Isac S", department="EEE",
                             scopus_author_id=JOYAL)
        self.general = _person("gen@test.edu", "General Person", department="EEE",
                               scopus_author_id=GENERAL)
        self.other_dept = _person("mech@test.edu", "Mech Person", department="MECH")
        import_profiles(read_profiles(io.BytesIO(profiles_workbook(TWO_SHEETS))))
        self.client = Client()

    def test_a_person_sees_their_own_profile_with_the_link_to_scopus(self):
        self.client.force_login(self.joyal)
        body = self.client.get("/api/me/scopus").json()
        profile = body["profile"]
        self.assertEqual(profile["scopus_id"], JOYAL)
        self.assertEqual(
            profile["url"], f"https://www.scopus.com/authid/detail.uri?authorId={JOYAL}"
        )
        self.assertEqual((profile["publications"], profile["citations"], profile["h_index"]),
                         (26, 166, 8))
        self.assertTrue(profile["imported_at"])

    def test_somebody_with_no_profile_is_told_so_rather_than_shown_zeros(self):
        self.client.force_login(self.other_dept)
        self.assertIsNone(self.client.get("/api/me/scopus").json()["profile"])

    def test_the_office_s_person_page_carries_it(self):
        self.client.force_login(_person("cell@test.edu", "Cell", role=Role.RESEARCH_CELL))
        report = self.client.get(f"/api/faculty/{self.joyal.id}/report").json()
        self.assertEqual(report["scopus_profile"]["citations"], 166)

    def test_the_head_sees_citations_and_scopus_publications_for_the_department(self):
        hod = _person("hod@test.edu", "Head EEE", role=Role.HOD, department="EEE")
        self.client.force_login(hod)
        body = self.client.get("/api/hod/overview").json()
        self.assertEqual(body["scopus"]["publications"], 26 + 49)
        self.assertEqual(body["scopus"]["citations"], 166 + 170)
        self.assertEqual(body["scopus"]["people_with_profile"], 2)
        people = {p["name"]: p for p in body["people"]}
        self.assertEqual(people["Joyal Isac S"]["scopus_citations"], 166)
        self.assertEqual(people["Joyal Isac S"]["scopus_h_index"], 8)

        def keys(value):
            if isinstance(value, dict):
                return set(value) | {k for v in value.values() for k in keys(v)}
            if isinstance(value, list):
                return {k for v in value for k in keys(v)}
            return set()

        self.assertFalse(keys(body) & MONEY_KEYS, "a head is shown no money")

    def test_the_reports_break_citations_down_by_department(self):
        self.client.force_login(_person("super@test.edu", "Super", role=Role.SUPER_ADMIN))
        rows = {r["department"]: r for r in self.client.get("/api/reports").json()["scopus_by_department"]}
        self.assertEqual(rows["EEE"]["citations"], 336)
        self.assertEqual(rows["EEE"]["publications"], 75)
        self.assertEqual(rows["EEE"]["people_with_profile"], 2)
        self.assertNotIn("MECH", rows)
