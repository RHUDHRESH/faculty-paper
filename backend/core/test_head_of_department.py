"""A head of department is a faculty member who also heads the department.

The college decided this on 2026-09-23. Until then a head was a money-blind
viewer who filed nothing; now they file, edit, withdraw and track their own
papers exactly as a faculty member does, and keep being money-blind about
everybody else.

So the rule these tests pin is no longer "no figure, anywhere" but "no figure
that is not yours":

- a head sees the amount on their own claims, and the claimant's journey;
- a head never sees the amount on a colleague's claim, by any route that
  returns claims -- the list, the detail, search, the reports, the department
  screens and the department's export.

The colleague's figure is a distinctive number, and the assertions look for
that number in the raw response. A test that only checks one key name passes
for a leak that arrives under a different one.
"""
from __future__ import annotations

import csv
import io
import json

from django.core.cache import cache
from django.test import Client, TestCase
from django.utils import timezone

from core import visibility
from core.hod import MONEY_KEYS
from core.models import Claim, ClaimStatus, FormulaConfig, Role, User
from core.services import rbac
from core.services.remuneration import DEFAULT_AUTHOR_POINTS

#: What the head was paid for their own paper, and what a colleague was paid
#: for theirs. Chosen to appear nowhere else in any response.
OWN_FIGURE = 42137.0
COLLEAGUE_FIGURE = 73519.75


def _person(email, name, role, **extra):
    return User.objects.create_user(email=email, password=None, name=name, role=role, **extra)


class HeadBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        FormulaConfig.objects.create(
            author_point_json=json.dumps(DEFAULT_AUTHOR_POINTS), active=True
        )
        cls.head = _person(
            "hd-head@test.edu", "Hema Head", Role.HOD,
            department="CSE", staff_id="STF-HD1",
        )
        cls.colleague = _person(
            "hd-colleague@test.edu", "Kiran Colleague", Role.FACULTY,
            department="CSE", staff_id="STF-HD2",
        )
        cls.office = _person("hd-office@test.edu", "Office Admin", Role.SUPER_ADMIN)

        cls.own_paid = Claim.objects.create(
            owner=cls.head, paper_title="The Head's Own Paid Paper",
            journal_title="Journal of Heads", status=ClaimStatus.PAID,
            remuneration=OWN_FIGURE, ticket_number="FP-2025-000901",
            publication_year=2025, quartile="Q1", submitted_at=timezone.now(),
            paid_at=timezone.now(),
            # Somebody else's payment, quoted inside the head's own claim by
            # the duplicate check. The claim is theirs; that figure is not.
            duplicate_matches_json=json.dumps([
                {"source": "claim", "id": "x", "title": "Similar", "amount": COLLEAGUE_FIGURE,
                 "reference": "FP-2024-000001", "who": "Kiran Colleague", "when": "2024-05"},
            ]),
            verification_snapshot_json=json.dumps({
                "issues": [],
                "paid": {"warning": True, "matches": [
                    {"source": "prior", "id": "y", "title": "Similar", "amount": COLLEAGUE_FIGURE,
                     "reference": "ERP-1", "who": "Kiran Colleague", "when": "2024-05"},
                ]},
            }),
        )
        cls.own_moving = Claim.objects.create(
            owner=cls.head, paper_title="The Head's Paper Under Review",
            journal_title="Journal of Heads", status=ClaimStatus.SUBMITTED,
            remuneration=11111.0, ticket_number="FP-2026-000902",
            publication_year=2026, submitted_at=timezone.now(),
        )
        cls.colleague_paid = Claim.objects.create(
            owner=cls.colleague, paper_title="Colleague Paper On Lattice Struts",
            journal_title="Journal of Lattices", status=ClaimStatus.PAID,
            remuneration=COLLEAGUE_FIGURE, ticket_number="FP-2025-000903",
            publication_year=2025, quartile="Q1", submitted_at=timezone.now(),
            paid_at=timezone.now(), author_position=1,
        )

    def setUp(self):
        cache.clear()
        self.client = Client()

    def as_head(self):
        self.client.force_login(self.head)
        return self.client

    def assertNoColleagueFigure(self, raw: str, where: str):
        for spelling in ("73519.75", "73,519.75"):
            self.assertNotIn(spelling, raw, f"a colleague's payment reached a head via {where}")


# --------------------------------------------------------------------------- #
# 1. A head is a claimant                                                     #
# --------------------------------------------------------------------------- #


class HeadFilesTheirOwnPapersTests(HeadBase):
    def test_rbac_lets_a_head_issue_claims_like_faculty(self):
        self.assertTrue(rbac.can_issue_claims(Role.HOD))
        self.assertTrue(rbac.can_faculty_portal(Role.HOD))
        self.assertIn(Role.HOD, rbac.CLAIMANT_ROLES)
        self.assertIn(Role.FACULTY, rbac.CLAIMANT_ROLES)

    def test_a_head_files_a_draft_and_it_is_theirs(self):
        r = self.as_head().post(
            "/api/claims",
            data=json.dumps({
                "paper_title": "A Paper the Head Wrote",
                "journal_title": "Journal of Heads",
                "self_reported_quartile": "Q1",
                "self_reported_snip": 1.0,
                "total_authors": 1,
                "author_position": 1,
            }),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        claim = Claim.objects.get(pk=body["id"])
        self.assertEqual(claim.owner_id, self.head.id)
        self.assertEqual(claim.status, ClaimStatus.DRAFT)
        # Their own claim, so the figure the draft was priced at comes back.
        self.assertIn("remuneration", body)
        self.assertEqual(body["faculty_stage"], "Draft")

    def test_a_head_edits_their_own_draft(self):
        draft = Claim.objects.create(
            owner=self.head, paper_title="Half Typed", status=ClaimStatus.DRAFT
        )
        r = self.as_head().patch(
            f"/api/claims/{draft.id}",
            data=json.dumps({"paper_title": "Fully Typed"}),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        draft.refresh_from_db()
        self.assertEqual(draft.paper_title, "Fully Typed")

    def test_a_head_withdraws_their_own_filed_paper(self):
        r = self.as_head().post(
            f"/api/claims/{self.own_moving.id}/withdraw",
            data=json.dumps({}), content_type="application/json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.own_moving.refresh_from_db()
        self.assertEqual(self.own_moving.status, ClaimStatus.DRAFT)

    def test_a_head_can_price_their_own_paper_while_filing(self):
        """The wizard's estimate is the head's own prospective paper."""
        r = self.as_head().post(
            "/api/calculate",
            data=json.dumps({"snip": 1.0, "quartile": "Q1", "total_authors": 1,
                             "author_position": 1, "indexing_level": "Scopus",
                             "publication_type": "Journal"}),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertIn("remuneration", r.json())

    def test_the_office_may_file_on_a_head_s_behalf(self):
        self.client.force_login(self.office)
        r = self.client.post(
            "/api/claims",
            data=json.dumps({"owner_id": self.head.id, "paper_title": "Filed For The Head"}),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(Claim.objects.get(pk=r.json()["id"]).owner_id, self.head.id)

    def test_the_office_s_owner_list_offers_a_head(self):
        """`/admin/faculty-options` is where the office finds whom a claim can
        belong to; a head is there with their account, like any faculty."""
        from core.models import FacultyMaster

        FacultyMaster.objects.create(
            staff_id="STF-HD1", name="Hema Head", department="CSE", email=self.head.email
        )
        self.client.force_login(self.office)
        rows = self.client.get("/api/admin/faculty-options?q=Hema").json()
        self.assertEqual([(r["owner_id"], r["has_user_account"]) for r in rows],
                         [(self.head.id, True)])
        rows = self.client.get("/api/admin/faculty-options?q=hd-head").json()
        self.assertIn(self.head.id, [r["owner_id"] for r in rows])

    def test_a_claim_may_be_reassigned_to_a_head(self):
        self.client.force_login(self.office)
        r = self.client.post(
            f"/api/admin/claims/{self.colleague_paid.id}/reassign",
            data=json.dumps({"owner_email": self.head.email,
                             "reason": "Filed against the wrong colleague by the import"}),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 200, r.content)


# --------------------------------------------------------------------------- #
# 2. Their own money, and nobody else's                                       #
# --------------------------------------------------------------------------- #


class HeadSeesTheirOwnAmountTests(HeadBase):
    def test_own_claims_list_carries_their_amounts_and_journey(self):
        r = self.as_head().get("/api/claims?limit=200")
        self.assertEqual(r.status_code, 200, r.content)
        rows = {c["id"]: c for c in r.json()["results"]}
        self.assertEqual(set(rows), {self.own_paid.id, self.own_moving.id})
        self.assertEqual(rows[self.own_paid.id]["remuneration"], OWN_FIGURE)
        self.assertEqual(rows[self.own_paid.id]["faculty_stage"], "Paid")
        self.assertEqual(rows[self.own_moving.id]["faculty_stage"], "Under review")
        self.assertIsNotNone(rows[self.own_moving.id]["days_waiting"])

    def test_own_claim_detail_carries_the_amount(self):
        r = self.as_head().get(f"/api/claims/{self.own_paid.id}")
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        self.assertEqual(body["remuneration"], OWN_FIGURE)
        self.assertEqual(body["faculty_stage"], "Paid")

    def test_other_people_s_payments_quoted_inside_their_own_claim_are_stripped(self):
        """The duplicate check quotes what a colleague was paid for a similar
        paper. The claim is the head's; that figure is not."""
        r = self.as_head().get(f"/api/claims/{self.own_paid.id}")
        raw = r.content.decode()
        self.assertIn(str(int(OWN_FIGURE)), raw)
        self.assertNoColleagueFigure(raw, "their own claim's payment-history match")
        body = r.json()
        matches = json.loads(body["duplicate_matches_json"])
        self.assertEqual(matches[0]["who"], "Kiran Colleague", "only the figure goes")
        self.assertNotIn("amount", matches[0])
        snapshot = json.loads(body["verification_snapshot_json"])
        self.assertNotIn("amount", snapshot["paid"]["matches"][0])


class HeadSeesNoColleagueAmountTests(HeadBase):
    """Every endpoint that returns claims, asked as the head."""

    def test_list(self):
        r = self.as_head().get("/api/claims?limit=200")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn(self.colleague_paid.id, r.content.decode())
        self.assertNoColleagueFigure(r.content.decode(), "/claims")

    def test_list_searched_by_the_colleague_s_title(self):
        r = self.as_head().get("/api/claims?q=Lattice")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["total"], 0)

    def test_detail(self):
        r = self.as_head().get(f"/api/claims/{self.colleague_paid.id}")
        self.assertIn(r.status_code, (403, 404), r.content)
        self.assertNoColleagueFigure(r.content.decode(), "/claims/{id}")

    def test_search(self):
        r = self.as_head().get("/api/search?q=Lattice%20Struts&kinds=people")
        self.assertEqual(r.status_code, 200, r.content)
        tickets = r.json()["tickets"]
        self.assertEqual([t["id"] for t in tickets], [self.colleague_paid.id],
                         "the head can still find the paper, which is department business")
        self.assertNotIn("amount", tickets[0])
        self.assertNoColleagueFigure(r.content.decode(), "/search")

    def test_reports_refuse_a_head(self):
        for path in (
            "/api/reports",
            "/api/reports/search",
            "/api/reports/export?fmt=csv",
            "/api/reports/search/export",
            f"/api/faculty/{self.colleague.id}/report",
            f"/api/faculty/{self.colleague.id}/report/export?fmt=csv",
            "/api/dashboard",
            "/api/lookup/ticket?q=FP-2025-000903",
        ):
            with self.subTest(path=path):
                r = self.as_head().get(path)
                self.assertEqual(r.status_code, 403, f"{path}: {r.content[:200]}")
                self.assertNoColleagueFigure(r.content.decode(errors="ignore"), path)

    def test_department_screens(self):
        for path in (
            "/api/hod/overview",
            "/api/hod/publications?limit=200",
            f"/api/hod/people/{self.colleague.id}",
            "/api/hod/opportunities",
            "/api/hod/standing",
            "/api/hod/targets",
            "/api/claims/counts",
        ):
            with self.subTest(path=path):
                r = self.as_head().get(path)
                self.assertEqual(r.status_code, 200, f"{path}: {r.content[:200]}")
                raw = r.content.decode()
                self.assertNoColleagueFigure(raw, path)
                # The department's own screens are money-free for everybody's
                # papers, the head's included.
                for key in MONEY_KEYS:
                    self.assertNotIn(f'"{key}"', raw, f"{path} carries {key}")

    def test_department_publications_still_list_the_colleague_s_paper(self):
        r = self.as_head().get("/api/hod/publications?limit=200")
        ids = [row["id"] for row in r.json()["results"]]
        self.assertIn(self.colleague_paid.id, ids)
        row = next(x for x in r.json()["results"] if x["id"] == self.colleague_paid.id)
        self.assertEqual(row["progress"], "Completed", "PAID is translated, not passed through")

    def test_department_export_csv(self):
        r = self.as_head().get("/api/hod/export?fmt=csv")
        self.assertEqual(r.status_code, 200)
        raw = r.content.decode()
        self.assertIn("Colleague Paper On Lattice Struts", raw)
        self.assertNoColleagueFigure(raw, "/hod/export csv")
        header = next(csv.reader(io.StringIO(raw)))
        self.assertFalse(
            [h for h in header if "amount" in h.lower() or "remuneration" in h.lower()],
            header,
        )

    def test_department_export_xlsx(self):
        from openpyxl import load_workbook

        r = self.as_head().get("/api/hod/export?fmt=xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        cells = [
            str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
            for v in row if v is not None
        ]
        self.assertIn("Colleague Paper On Lattice Struts", cells)
        for v in cells:
            self.assertNotIn("73519.75", v, "a colleague's payment reached the department workbook")


class TheHeadCountsAsAMemberOfTheirDepartmentTests(HeadBase):
    """Their papers are the department's output, so they are one of its people.

    The department screens listed `role=FACULTY` only while counting the
    head's own papers in the totals -- "2 of 1 have published" is what that
    reads as once a head files."""

    def test_overview_lists_the_head_among_the_department_s_people(self):
        body = self.as_head().get("/api/hod/overview").json()
        people = {p["id"]: p for p in body["people"]}
        self.assertIn(self.head.id, people)
        self.assertEqual(people[self.head.id]["publications"], 2)
        totals = body["totals"]
        self.assertLessEqual(totals["faculty_who_published"], totals["faculty_in_department"])
        self.assertEqual(totals["faculty_in_department"], 2)

    def test_standing_counts_the_head_as_faculty(self):
        body = self.as_head().get("/api/hod/standing").json()
        self.assertEqual(body["mine"]["faculty"], 2)

    def test_opportunities_consider_the_head_too(self):
        Claim.objects.filter(owner=self.head).update(author_position=2)
        body = self.as_head().get("/api/hod/opportunities").json()
        never_led = next(g for g in body["groups"] if g["key"] == "never_led")
        self.assertIn(self.head.id, [p["id"] for p in never_led["people"]],
                      "the head has filed, and led none of it")


class TheRuleLivesInTheRendererTests(HeadBase):
    """`visibility.for_viewer` is the one place a head's view of money is
    decided. Asked directly, with a payload no endpoint builds today, so the
    rule is proven for the endpoint somebody writes next year."""

    def _claim_dict(self, claim):
        from core.api.deps import claim_to_dict

        return claim_to_dict(claim)

    def test_own_keeps_its_figures_and_a_colleague_s_loses_every_one(self):
        payload = {
            "results": [self._claim_dict(self.own_paid), self._claim_dict(self.colleague_paid)],
            "total": 2,
        }
        shaped = visibility.for_viewer(self.head, payload)
        own, theirs = shaped["results"]
        self.assertEqual(own["remuneration"], OWN_FIGURE)
        self.assertEqual(own["faculty_stage"], "Paid", "their own claim reads as a claimant's")
        for key in MONEY_KEYS:
            self.assertNotIn(key, theirs, f"a colleague's claim kept {key}")
        self.assertNoColleagueFigure(json.dumps(shaped), "the renderer")

    def test_a_row_naming_somebody_else_as_owner_is_stripped_whatever_its_shape(self):
        row = {"owner_id": self.colleague.id, "title": "x", "amount": COLLEAGUE_FIGURE,
               "nested": {"total_amount": COLLEAGUE_FIGURE}}
        self.assertEqual(
            visibility.for_viewer(self.head, [row]),
            [{"owner_id": self.colleague.id, "title": "x", "nested": {}}],
        )

    def test_other_roles_are_unaffected(self):
        payload = [self._claim_dict(self.colleague_paid)]
        principal = _person("hd-prin@test.edu", "P", Role.PRINCIPAL)
        self.assertEqual(
            visibility.for_viewer(principal, payload)[0]["remuneration"], COLLEAGUE_FIGURE
        )
