"""Record accuracy: Scopus ids from the ERP workbook, own-name variants,
roster name formats, placeholder venues, the Scopus AU-ID sync."""
from __future__ import annotations

import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings

from core.models import Authorship, Publication, Role, User
from core.services import coauthors as graph
from core.services import publications as P
from core.services.author_names import name_score
from core.services.erp_scopus import link_scopus_ids, read_erp_scopus
from core.services.normalize import clean_venue
from core.services.research_picture import my_research
from core.services.scopus_sync import sync_scopus_authors


def workbook() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faculty_Data"
    ws.append(["S.No", "Department", "Bio-ID", "Staff-ID", "Scopus ID", "Name of the Staff", "Email ID", "Scopus Profile"])
    ws.append([1, "EEE", 101.0, "TSEE023", 57983494200.0, "Mr. S. Joyal Isac", "joyal@x.edu", "View Profile"])
    ws.append([2, "S&H", 102.0, "TSSH208", 59009649300.0, "Dr. R. Subhashini", "subha@x.edu", "View Profile"])
    ws.append([3, "ECE", 103.0, "TSEC001", 11111111111.0, "Dr. Conflict", "c@x.edu", "View Profile"])
    ws.append([4, "ECE", 104.0, "TSEC999", 22222222222.0, "Nobody Here", "n@x.edu", "View Profile"])
    ws.append([5, "ECE", 105.0, "", 33333333333.0, "Bio Only", "", "View Profile"])
    ws.append([6, "ECE", 106.0, "TSEC002", None, "From Papers", "fp@x.edu", "View Profile"])
    p = wb.create_sheet("Processed")
    p.append(["S.No", "Faculty ID", "Faculty Name", "Biometric ID", "Scopus ID"])
    p.append([1, "TSEC002", "From Papers", 106, 44444444444])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def mk(name, **kw):
    return User.objects.create_user(email=f"{name.split()[-1].lower()}{User.objects.count()}@x.edu",
                                    password="p", name=name, role=kw.pop("role", Role.FACULTY), **kw)


class ErpScopusLinkTests(TestCase):
    def setUp(self):
        self.joyal = mk("Mr. S. Joyal Isac", staff_id="TSEE023")
        self.subha = mk("Dr. R. Subhashini", staff_id="tssh208", scopus_author_id="59009649300")
        self.conf = mk("Dr. Conflict", staff_id="TSEC001", scopus_author_id="99999999999")
        self.bio = mk("Bio Only", biometric_id="105")
        self.fp = mk("From Papers", staff_id="TSEC002")

    def test_links_by_staff_then_bio_reports_conflicts_and_is_idempotent(self):
        entries = read_erp_scopus(workbook())
        r = link_scopus_ids(entries)
        self.assertEqual((r["set"], r["same"], len(r["conflicts"]), len(r["unmatched"])), (3, 1, 1, 1))
        self.joyal.refresh_from_db()
        self.assertEqual(self.joyal.scopus_author_id, "57983494200")
        self.assertIn("authorId=57983494200", self.joyal.scopus_author_url)
        self.bio.refresh_from_db()
        self.assertEqual(self.bio.scopus_author_id, "33333333333")
        self.fp.refresh_from_db()
        self.assertEqual(self.fp.scopus_author_id, "44444444444")
        self.conf.refresh_from_db()
        self.assertEqual(self.conf.scopus_author_id, "99999999999")  # never overwritten
        again = link_scopus_ids(read_erp_scopus(workbook()))
        self.assertEqual((again["set"], again["same"]), (0, 4))

    def test_office_imports_action(self):
        admin = mk("Office", role=Role.SUPER_ADMIN)
        c = Client()
        c.force_login(admin)
        f = SimpleUploadedFile("erp.xlsx", workbook().read())
        r = c.post("/api/admin/scopus-ids/link", {"file": f})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["set"], 3)
        c.force_login(self.joyal)
        f = SimpleUploadedFile("erp.xlsx", workbook().read())
        self.assertEqual(c.post("/api/admin/scopus-ids/link", {"file": f}).status_code, 403)


def pub(title, **kw):
    return Publication.objects.create(title=title, normalized_title=title.lower(), year=2024, **kw)


class AuthorAccuracyTests(TestCase):
    def setUp(self):
        self.subha = mk("Dr. R. Subhashini", department="S&H-MATHS")
        self.kamala = mk("Dr. Kamaladevi K", department="ECE")

    def test_run_together_roster_name_matches_split_paper_name(self):
        self.assertGreaterEqual(name_score("K. Kamala Devi", "Dr. Kamaladevi K"), 0.85)
        self.assertEqual(name_score("Kamala Devi R", "Dr. Kamaladevi K"), 0.0)
        self.assertGreaterEqual(name_score("Dr. R. Subhashini", "R. Subashini"), 0.85)  # one letter apart
        self.assertEqual(name_score("Dr. R. Subhashini", "K. Subashini"), 0.0)
        self.assertEqual(name_score("Kumaran R", "Kumar R"), 0.0)
        p = pub("A long enough title about kamala devi things")
        Authorship.objects.create(publication=p, position=1, display_name="K. Kamala Devi",
                                  author_key="n:k", is_college=True)
        P.match_authors()
        self.assertEqual(Authorship.objects.get(publication=p).user, self.kamala)

    def test_own_name_variant_is_not_a_coauthor(self):
        p = pub("A paper on graphs with a sufficiently long title")
        # The record put her on (no position); OpenAlex spelt her name its own way, not college.
        Authorship.objects.create(publication=p, position=None, display_name=self.subha.name, author_key="n:s",
                                  user=self.subha, match_method="record", match_confidence=0.95, is_college=True)
        Authorship.objects.create(publication=p, position=1, display_name="R. Subhashini", author_key="A9",
                                  openalex_author_id="A9", institution_name="Saveetha Engineering College")
        Authorship.objects.create(publication=p, position=2, display_name="K. Kamala Devi", author_key="n:kd",
                                  is_college=True)
        out = P.match_authors()
        self.assertEqual(out["own_variants_absorbed"], 1)
        rows = list(p.authorships.order_by("position"))
        self.assertEqual([(r.position, r.user_id) for r in rows], [(1, self.subha.id), (2, self.kamala.id)])
        co = graph.coauthors(self.subha)
        names = [i["name"] for i in co["inside"] + co["outside"]]
        self.assertEqual(names, [self.kamala.name])

    def test_coauthor_list_drops_self_even_when_rows_disagree(self):
        p = pub("Another sufficiently long title for the test")
        Authorship.objects.create(publication=p, position=1, display_name="Dr. R. Subhashini", author_key="A1",
                                  user=self.subha, match_method="manual", match_confidence=1, match_locked=True)
        Authorship.objects.create(publication=p, position=3, display_name="Subhashini R.", author_key="n:x",
                                  match_locked=True)
        Authorship.objects.create(publication=p, position=2, display_name="Kamaladevi K.", author_key="n:y",
                                  match_locked=True)
        Authorship.objects.create(publication=p, position=4, display_name="Dr. K. Kamaladevi", author_key="n:z",
                                  user=self.kamala, match_locked=True, match_confidence=1, match_method="manual")
        co = graph.coauthors(self.subha)
        self.assertEqual([(i["name"], i["user_id"]) for i in co["inside"]], [(self.kamala.name, self.kamala.id)])
        self.assertEqual(co["outside"], [])


class VenueTests(TestCase):
    def test_placeholders_are_unknown(self):
        for v in ("-", "N/A", "NA", "", None, " n.a. ", "—"):
            self.assertEqual(clean_venue(v), "", v)
        self.assertEqual(clean_venue(" IEEE  Access "), "IEEE Access")

    def test_where_you_publish_skips_dash(self):
        u = mk("Asha Rao")
        for i, venue in enumerate(["-", "N/A", "IEEE Access"]):
            p = pub(f"A paper number {i} with a long enough title", venue=venue)
            Authorship.objects.create(publication=p, position=1, display_name=u.name, author_key=f"u{i}",
                                      user=u, is_college=True, match_method="manual", match_confidence=1)
        body = my_research(u)
        names = [v["name"] for v in body["venues"]]
        self.assertEqual(names, ["IEEE Access"])


class ScopusIdMatchTests(TestCase):
    def test_openalex_author_with_members_scopus_id_is_them(self):
        u = mk("Mr. S. Joyal Isac", scopus_author_id="57983494200")
        p = pub("Some paper where OpenAlex has a different spelling")
        Authorship.objects.create(publication=p, position=1, display_name="Joyal Isac Selvaraj", author_key="A5",
                                  openalex_author_id="A5", is_college=False)
        calls = []

        def fetch(path, params):
            calls.append(params)
            return {"results": [{"id": "https://openalex.org/A5",
                                 "ids": {"scopus": "http://www.scopus.com/inward/authorDetails.url?authorID=57983494200&partnerID=MN8TOARS"}}]}

        # Only rows that are college or matched are asked about: make it college.
        Authorship.objects.update(is_college=True)
        got = P.link_by_scopus_ids(fetch=fetch)
        self.assertEqual(got["rows_linked"], 1)
        P.match_authors()  # survives a re-match
        self.assertEqual(Authorship.objects.get().user, u)


def entry(n, *, doi=None, title=None):
    return {"dc:title": title or f"Scopus paper number {n} about inverters", "prism:doi": doi,
            "eid": f"2-s2.0-{n}", "prism:publicationName": "IEEE Access", "prism:issn": "21693536",
            "prism:coverDate": "2024-05-01", "subtypeDescription": "Article", "citedby-count": str(n)}


@override_settings(SCOPUS_API_KEY="k")
class ScopusSyncTests(TestCase):
    def test_pages_upserts_links_and_resumes(self):
        u = mk("Mr. S. Joyal Isac", scopus_author_id="57983494200")
        other = mk("Dr. R. Subhashini", scopus_author_id="59009649300")
        existing = pub("An existing paper already on the record", doi="10.1/x")
        pages = {0: [entry(i) for i in range(1, 26)] + [], 25: [entry(26, doi="10.1/x")]}

        def search(query, start):
            if "59009649300" in query:
                from core.services.scopus import ScopusError

                raise ScopusError("quota", code="rate_limit")
            return {"search-results": {"opensearch:totalResults": "26", "entry": pages[start]}}

        out = sync_scopus_authors(users=[u, other], search=search)
        self.assertEqual(out["people"], 1)
        self.assertTrue(out["stopped"].startswith("rate_limit"))
        self.assertEqual(Publication.objects.filter(authorships__user=u).distinct().count(), 26)
        existing.refresh_from_db()
        self.assertTrue(existing.scopus_indexed)
        self.assertEqual(existing.scopus_citations, 26)
        u.refresh_from_db()
        self.assertIsNotNone(u.scopus_synced_at)
        before = Publication.objects.count()
        sync_scopus_authors(users=[u], search=search)
        self.assertEqual(Publication.objects.count(), before)
        self.assertEqual(Authorship.objects.filter(user=u).count(), 26)
        P.match_authors()
        self.assertEqual(Authorship.objects.filter(user=u).count(), 26)
