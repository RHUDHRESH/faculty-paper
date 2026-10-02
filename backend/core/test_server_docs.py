"""Server-generated documents: real rupee sign, Indian grouping, typed cells, readable CSVs.

Each test reads the file the way its reader would (text out of the PDF, cells
out of the workbook, rows out of the CSV) and asserts on what they would see.
"""
import csv
import io
from datetime import date

import openpyxl
from django.test import TestCase
from django.utils import timezone
from pypdf import PdfReader

from core.models import AuditLog, Claim, ClaimStatus, PaidLedger, Role, User
from core.services import exporters, pdf_fonts

LONG_TITLE = ("Advanced machine learning assisted prediction of multi-fiber hybrid natural "
              "fiber/epoxy composites based on jute, coconut coir, banana, and pineapple leaf fibers")


def pdf_text(body: bytes) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(body)).pages)


def csv_rows(body: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


class GroupingTests(TestCase):
    def test_indian_grouping(self):
        self.assertEqual(pdf_fonts.group_in(1586), "1,586")
        self.assertEqual(pdf_fonts.group_in(12865956), "1,28,65,956")
        self.assertEqual(pdf_fonts.group_in(999), "999")
        self.assertEqual(pdf_fonts.group_in(None), "–")


class DocsTests(TestCase):
    def setUp(self):
        self.fin = User.objects.create_user(email="sd-fin@t.edu", password="p", name="Fin Officer",
                                            role=Role.FINANCE)
        self.pr = User.objects.create_user(email="sd-pr@t.edu", password="p", name="Principal",
                                           role=Role.PRINCIPAL)
        self.admin = User.objects.create_user(email="sd-sa@t.edu", password="p", name="Admin",
                                              role=Role.SUPER_ADMIN)
        self.fac = User.objects.create_user(email="sd-fac@t.edu", password="p", name="Priya Raman",
                                            role=Role.FACULTY, department="S&H-PHY", staff_id="S1")
        aug = date(2026, 8, 1)
        now = timezone.now()
        self.claim = Claim.objects.create(owner=self.fac, status=ClaimStatus.PAID, remuneration=1234567,
                                          payout_month=aug, ticket_number="FP-2026-000001",
                                          paper_title=LONG_TITLE, publication_year=2026,
                                          director_approved_at=now, paid_at=now)
        PaidLedger.objects.create(claim=self.claim, payout_month=aug, amount=1234567, department="S&H-PHY",
                                  faculty_name="Priya Raman", staff_id="S1", paper_title=LONG_TITLE)

    # -- payout statement ------------------------------------------------

    def test_statement_pdf_prints_rupee_sign_and_whole_titles(self):
        self.client.force_login(self.fin)
        r = self.client.get("/api/payouts/statement.pdf", {"month": "2026-08"})
        self.assertEqual(r.status_code, 200)
        # A font that has the rupee glyph is embedded: the statement is set in
        # Inter (which carries U+20B9); older documents set the sign in DejaVu.
        self.assertTrue(b"DejaVuSans" in r.content or b"Inter" in r.content)
        text = pdf_text(r.content)
        self.assertIn("₹12,34,567", text)
        self.assertNotIn("Rs.", text)
        # The title is wrapped, never cut at a character limit.
        self.assertIn("pineapple leaf", text.replace("\n", " "))
        self.assertIn("S&H-PHY", text)

    def test_bank_csv_total_matches_statement(self):
        self.client.force_login(self.fin)
        st = self.client.get("/api/payouts/statement", {"month": "2026-08"}).json()
        rows = csv_rows(self.client.get("/api/payouts/statement.csv", {"month": "2026-08"}).content)
        self.assertEqual(float(rows[-1][7]), st["total"])

    # -- principal brief -------------------------------------------------

    def test_brief_pdf_and_workbook(self):
        self.client.force_login(self.pr)
        brief = self.client.get("/api/reports/brief", {"year": 2026}).json()
        r = self.client.get("/api/reports/brief/export", {"year": 2026, "fmt": "pdf"})
        self.assertEqual(r["Content-Disposition"], 'attachment; filename="research-brief-2026.pdf"')
        text = pdf_text(r.content)
        self.assertIn("₹" + pdf_fonts.group_in(brief["totals"]["paid"]), text.replace("\n", " "))
        self.assertNotIn("Rs ", text)
        wb = openpyxl.load_workbook(io.BytesIO(
            self.client.get("/api/reports/brief/export", {"year": 2026, "fmt": "xlsx"}).content))
        ws = wb["Summary"]
        self.assertEqual(ws.freeze_panes, "A5")
        self.assertEqual(ws["A9"].value, "Incentives paid, FY (₹)")
        self.assertEqual(ws["B9"].value, brief["totals"]["paid"])
        self.assertIn("##\\,##\\,##0", ws["B9"].number_format)  # lakh grouping
        self.assertIn("(₹)", wb["Departments"]["I4"].value)

    # -- CSVs ------------------------------------------------------------

    def test_my_payment_statement_csv(self):
        self.client.force_login(self.fac)
        r = self.client.get("/api/me/payments/statement", {"format": "csv"})
        self.assertTrue(r.content.startswith("\ufeff".encode("utf-8")))
        rows = csv_rows(r.content)
        body = [x for x in rows if x and x[0] == "Aug 2026"]
        self.assertEqual(body[0][-1], "1234567.00")
        self.assertEqual(rows[-1], ["Total", "", "", "", "", "1234567.00"])

    def test_leaderboard_csv_reads_like_a_table(self):
        self.client.force_login(self.fac)
        r = self.client.get("/api/leaderboard", {"fmt": "csv"})
        head = csv_rows(r.content)[0]
        self.assertEqual(head[0], "Rank")
        self.assertNotIn("id", [h.lower() for h in head])
        self.assertNotIn("me", [h.lower() for h in head])

    def test_audit_csv_local_time(self):
        AuditLog.objects.create(actor=self.admin, action="X", entity="User", entity_id="1")
        self.client.force_login(self.admin)
        r = self.client.get("/api/admin/audit.csv")
        self.assertTrue(r.content.startswith("\ufeff".encode("utf-8")))
        rows = csv_rows(r.content)
        self.assertEqual(rows[0][0], "When (IST)")
        self.assertRegex(rows[1][0], r"^\d{4}-\d\d-\d\d \d\d:\d\d:\d\d$")

    def test_ledger_csv_amounts_two_decimals(self):
        self.client.force_login(self.fin)
        rows = csv_rows(self.client.get("/api/admin/ledger/export").content)
        self.assertEqual(rows[1][rows[0].index("amount")], "1234567.00")


class ExporterFormatTests(TestCase):
    def test_years_ratios_and_money(self):
        pack = {"S&H": {"columns": ["Year of publication", "Papers per head", "Amount paid"],
                        "rows": [[2025, 1.88, 1234567.5]]}}
        ws = openpyxl.load_workbook(io.BytesIO(exporters.render(pack, "xlsx", title="t")))["S&H"]
        row = [c for c in ws[2]]
        self.assertEqual(row[0].number_format, "0")  # 2025, not 2,025
        self.assertEqual(row[1].number_format, "#,##0.00")  # 1.88, not 2
        self.assertIn('"₹"##\\,##\\,##0.00', row[2].number_format)
        text = pdf_text(exporters.render(pack, "pdf", title="Report & pack"))
        self.assertIn("₹12,34,567.50", text)
        self.assertIn("1.88", text)
        self.assertIn("2025", text)
