"""The head's brief and report: on track, who needs a push, who writes with whom.

Counted from the college's publication record, the same one the Principal's
pages use, so a head and the Principal see the same number for one department.
"""
from __future__ import annotations

import io
import json
from datetime import date

from django.core.cache import cache
from django.test import Client, TestCase
from django.utils import timezone

from core.api.hod_brief import _elapsed, _months_left, _verdict
from core.models import (
    Authorship, Claim, ClaimStatus, DepartmentAssignment, DepartmentTarget, FormulaConfig, Publication, Role, User,
)
from core.services import principal_brief
from core.services.remuneration import DEFAULT_AUTHOR_POINTS

import itertools

_N = itertools.count(1)
FIGURE = 61234.5  # a colleague's amount; must appear in no response or file


def _u(email, name, role=Role.FACULTY, dept="CSE"):
    return User.objects.create_user(email=email, password=None, name=name, role=role, department=dept)


def _paper(owner, year, quartile=None, position=1, area="ML", **extra):
    return Claim.objects.create(
        owner=owner, status=extra.pop("status", ClaimStatus.PAID), paper_title=f"P {owner.name} {year} {next(_N)}",
        journal_title="J", publication_year=year, publication_date=f"{year}-01-05",
        quartile=quartile, author_position=position, total_authors=3, subject_category=area,
        remuneration=FIGURE, issn="1234-5678", doi=f"10.5555/{owner.id}{year}{quartile}{position}/{next(_N)}",
        submitted_at=timezone.now(), **extra,
    )


def _record(user, year, title, quartile="", topics=None, doi="10.1/x", issn="1111-2222", position=1, on=None):
    pub = Publication.objects.create(
        title=title, normalized_title=title.lower(), year=year, date=on or date(year, 2, 1), venue="Journal of Tests", type="article",
        quartile=quartile, doi=doi, issn=issn, topics_json=json.dumps(topics or []),
    )
    Authorship.objects.create(publication=pub, position=position, display_name=user.name, user=user, is_college=True)
    return pub


class BriefTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        FormulaConfig.objects.create(author_point_json=json.dumps(DEFAULT_AUTHOR_POINTS), active=True)
        cls.year = timezone.localdate().year
        cls.head = _u("h@t.edu", "Head", Role.HOD)
        cls.star = _u("s@t.edu", "Star")
        cls.quiet = _u("q@t.edu", "Quiet")
        cls.other = _u("o@t.edu", "Elsewhere", dept="ECE")
        for q in ("Q1", "Q1", "Q1"):
            _paper(cls.star, cls.year, q)
        _paper(cls.quiet, cls.year - 1, "Q3")
        _paper(cls.quiet, cls.year, "Q2", status=ClaimStatus.REJECTED, rejected_outright=True)
        _paper(cls.other, cls.year, "Q1")
        DepartmentTarget.objects.create(
            department="CSE", year=cls.year, metric=DepartmentTarget.Metric.PUBLICATIONS, target=100,
        )

    def setUp(self):
        cache.clear()
        self.c = Client()
        self.c.force_login(self.head)

    def brief(self, query=""):
        r = self.c.get(f"/api/hod/brief{query}")
        self.assertEqual(r.status_code, 200, r.content)
        return r.json()

    def test_a_claim_the_chain_refused_is_not_a_paper(self):
        b = self.brief()
        self.assertEqual(b["totals"]["publications"], 3)
        self.assertNotIn("rejected_outright", b["totals"])
        t = self.c.get("/api/hod/targets").json()["department_targets"][0]
        self.assertEqual(t["done"], 3)

    def test_papers_come_from_the_record_and_agree_with_the_principal(self):
        """The bug this replaces: a head saw 19 papers where the record held 373."""
        _record(self.star, self.year, "Record paper one", "Q1")
        _record(self.quiet, self.year, "Record paper two", "Q2")
        _record(self.other, self.year, "Not ours", "Q1")
        cache.clear()
        b = self.brief()
        self.assertEqual(b["totals"]["publications"], 5)  # 3 claims + 2 record papers, none from ECE
        pb = principal_brief.brief(self.year)
        row = next(d for d in pb["departments"] if d["department"] == "CSE")
        self.assertEqual(row["papers"], b["totals"]["publications"])
        self.assertEqual(b["by_year"][-1]["publications"], 5)

    def test_pace_months_left_and_verdict(self):
        b = self.brief()
        pubs = b["targets"][0]
        self.assertAlmostEqual(pubs["expected_by_now"], round(100 * b["elapsed"], 1))
        self.assertEqual(pubs["to_go"], 97)
        self.assertEqual(b["months_left"], 12 - timezone.localdate().month)
        if b["months_left"]:
            self.assertAlmostEqual(pubs["per_month_needed"], round(97 / b["months_left"], 1))
        self.assertEqual(_verdict(50, 60, 100), "behind")
        self.assertEqual(_verdict(55, 60, 100), "close")
        self.assertEqual(_verdict(61, 60, 100), "on_track")
        self.assertEqual(_verdict(100, 60, 100), "met")
        self.assertEqual(_elapsed(2020, date(2026, 6, 1)), 1.0)
        self.assertEqual(_elapsed(2030, date(2026, 6, 1)), 0.0)
        self.assertEqual(_months_left(2026, date(2026, 9, 30)), 3)
        self.assertEqual(_months_left(2026, date(2026, 12, 31)), 0)
        self.assertEqual(_months_left(2025, date(2026, 9, 30)), 0)

    def test_same_date_last_year_uses_papers_that_carry_a_date(self):
        today = timezone.localdate()
        _record(self.star, self.year - 1, "Early last year", on=date(self.year - 1, 1, 2))
        _record(self.star, self.year - 1, "Late last year", on=date(self.year - 1, 12, 30))
        _record(self.star, self.year - 1, "No date", on=None)
        Publication.objects.filter(title="No date").update(date=None)
        cache.clear()
        t = self.brief()["totals"]
        self.assertEqual(t["last_year_full"], 4)  # the claim from last year and three record papers
        if (today.month, today.day) != (12, 30):
            self.assertLess(t["last_year_to_date"], t["last_year_full"])

    def test_push_lists_reason_and_next_step(self):
        b = self.brief()
        names = [r["person"]["name"] for r in b["push"]]
        self.assertIn("Quiet", names)
        self.assertNotIn("Head", names)
        self.assertNotIn("Elsewhere", [p["name"] for p in b["people"]])
        quiet = next(r for r in b["push"] if r["person"]["name"] == "Quiet")
        self.assertEqual(quiet["kind"], "slipped")
        self.assertIn(f"No paper in {self.year}; 1 in {self.year - 1}", quiet["reasons"][0])
        self.assertTrue(quiet["next_step"])
        self.assertIn("reminder from your head of department", quiet["draft"])
        self.assertNotIn(str(int(FIGURE)), quiet["draft"])

    def test_a_person_with_no_paper_and_no_scopus_id_is_a_data_gap_first(self):
        _u("n@t.edu", "Newcomer")
        b = self.brief()
        row = next(r for r in b["push"] if r["person"]["name"] == "Newcomer")
        self.assertEqual(row["kind"], "never")
        self.assertIn("no Scopus ID on file", row["reasons"][0])
        self.assertIn("Scopus ID", row["next_step"])

    def test_pairs_need_a_shared_area(self):
        b = self.brief()
        # Claims carry no topics, so nothing to pair on until the record has them.
        self.assertEqual(b["pairs"], [])
        for i in range(2):
            _record(self.star, self.year - i, f"Q1 in ML {i}", "Q1", topics=["Machine learning"], doi=f"10.1/s{i}")
        _record(self.quiet, self.year - 1, "Older in ML", "", topics=["Machine learning"], doi="10.1/q")
        cache.clear()
        b = self.brief()
        self.assertEqual(len(b["pairs"]), 1)
        self.assertEqual(b["pairs"][0]["mentee"]["name"], "Quiet")
        self.assertEqual(b["pairs"][0]["mentor"]["name"], "Star")
        self.assertEqual(b["pairs"][0]["area"], "Machine learning")

    def test_existing_pairing_is_not_suggested_again(self):
        for i in range(2):
            _record(self.star, self.year - i, f"Q1 in ML {i}", "Q1", topics=["Machine learning"], doi=f"10.1/s{i}")
        _record(self.quiet, self.year - 1, "Older in ML", "", topics=["Machine learning"], doi="10.1/q")
        DepartmentAssignment.objects.create(
            department="CSE", kind=DepartmentAssignment.Kind.PAIRING, title="x",
            assignee=self.quiet, partner=self.star, created_by=self.head,
        )
        cache.clear()
        self.assertEqual(self.brief()["pairs"], [])

    def test_the_department_sits_at_a_place_without_naming_others(self):
        b = self.brief()
        raw = json.dumps(b)
        self.assertIn("college", b)
        self.assertNotIn("Elsewhere", raw)
        self.assertNotIn("ECE", raw)

    def test_no_money_anywhere(self):
        self.assertNotIn(str(FIGURE), self.c.get("/api/hod/brief").content.decode())
        from openpyxl import load_workbook
        r = self.c.get("/api/hod/report?fmt=xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames[:3], ["Summary", "Per teacher", "NAAC 3.3 papers"])
        cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row]
        self.assertFalse(any(str(int(FIGURE)) in v for v in cells))
        naac = wb["NAAC 3.3 papers"]
        self.assertIn("Link to article (DOI)", [c.value for c in naac[4]])

    def test_workbook_lists_record_papers_with_their_links(self):
        _record(self.star, self.year, "Record with a DOI", "Q1", doi="10.9/abc")
        _record(self.star, self.year, "Record without a DOI", "", doi="")
        cache.clear()
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.c.get("/api/hod/report?fmt=xlsx").content))
        rows = [[c.value for c in row] for row in wb["NAAC 3.3 papers"].iter_rows(min_row=5)]
        titles = [r[0] for r in rows]
        self.assertIn("Record with a DOI", titles)
        row = rows[titles.index("Record with a DOI")]
        self.assertIn("https://doi.org/10.9/abc", row)
        missing = [[c.value for c in row] for row in wb["Missing data"].iter_rows(min_row=4)]
        self.assertTrue(any(r[1] == "Record without a DOI" and "DOI" in r[4] for r in missing))

    def test_pdf_is_a4(self):
        r = self.c.get("/api/hod/report?fmt=pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        from pypdf import PdfReader
        page = PdfReader(io.BytesIO(r.content)).pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), 595.27, places=0)
        self.assertNotIn(str(int(FIGURE)), page.extract_text())

    def test_pdf_reads_like_a_note_to_the_principal(self):
        from pypdf import PdfReader
        pages = PdfReader(io.BytesIO(self.c.get("/api/hod/report?fmt=pdf").content)).pages
        first = pages[0].extract_text()
        for part in ("Saveetha Engineering College", "Research note for the Principal", "Summary",
                     "Pace", "Who needs a push", "Next step", "Records to fix", f"Page 1 of {len(pages)}"):
            self.assertIn(part, first)
        text = " ".join(p.extract_text() for p in pages)
        self.assertIn("Per teacher", text)
        self.assertIn(f"Page {len(pages)} of {len(pages)}", text)

    def test_faculty_cannot(self):
        c = Client()
        c.force_login(self.star)
        self.assertEqual(c.get("/api/hod/brief").status_code, 403)
        self.assertEqual(c.get("/api/hod/report").status_code, 403)
        self.assertEqual(c.get("/api/hod/department/papers").status_code, 403)
        self.assertEqual(c.get("/api/hod/department/records").status_code, 403)


class DepartmentPapersTests(TestCase):
    """The list behind every figure, and the export of it."""

    @classmethod
    def setUpTestData(cls):
        cls.year = timezone.localdate().year
        cls.head = _u("h2@t.edu", "Head", Role.HOD)
        cls.a = _u("a@t.edu", "Asha")
        cls.b = _u("b@t.edu", "Bala")
        cls.other = _u("o2@t.edu", "Elsewhere", dept="ECE")
        _record(cls.a, cls.year, "Alpha paper", "Q1", doi="10.1/a")
        _record(cls.a, cls.year, "Beta paper", "Q3", doi="")
        _record(cls.b, cls.year, "Gamma paper", "", issn="")
        _record(cls.other, cls.year, "Delta elsewhere", "Q1")

    def setUp(self):
        cache.clear()
        self.c = Client()
        self.c.force_login(self.head)

    def test_lists_only_this_department_and_counts_match_the_brief(self):
        r = self.c.get(f"/api/hod/department/papers?year={self.year}").json()
        titles = [p["title"] for p in r["results"]]
        self.assertEqual(sorted(titles), ["Alpha paper", "Beta paper", "Gamma paper"])
        self.assertEqual(r["total"], self.c.get("/api/hod/brief").json()["totals"]["publications"])
        self.assertNotIn("claim_id", json.dumps(r))

    def test_filters(self):
        get = lambda q: self.c.get(f"/api/hod/department/papers?{q}").json()  # noqa: E731
        self.assertEqual(get("quartile=Q1")["total"], 1)
        self.assertEqual(get("quartile=none")["total"], 1)
        self.assertEqual(get("quartile=low")["total"], 1)
        self.assertEqual(get("quartile=top")["total"], 1)
        self.assertEqual(get("missing=doi")["total"], 1)
        self.assertEqual(get("missing=issn")["total"], 1)
        self.assertEqual(get("q=gamma")["total"], 1)
        self.assertEqual(get(f"person={self.a.id}")["total"], 2)
        self.assertEqual(get("q=Asha")["total"], 2)

    def test_another_departments_person_is_refused(self):
        r = self.c.get(f"/api/hod/department/papers?person={self.other.id}")
        self.assertEqual(r.status_code, 403)

    def test_export_carries_the_filter_and_no_money(self):
        from openpyxl import load_workbook
        r = self.c.get(f"/api/hod/department/papers/export?year={self.year}&quartile=Q1")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        cells = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None]
        self.assertTrue(any("Department of CSE" in v and "Q1" in v for v in cells))
        self.assertIn("Alpha paper", cells)
        self.assertNotIn("Beta paper", cells)

    def test_records_to_fix_names_papers_and_people(self):
        r = self.c.get(f"/api/hod/department/records?year={self.year}").json()
        by = {p["title"]: p["missing"] for p in r["papers"]}
        self.assertEqual(by["Beta paper"], ["DOI"])
        self.assertEqual(by["Gamma paper"], ["ISSN"])
        self.assertNotIn("Alpha paper", by)
        self.assertIn("Asha", [p["name"] for p in r["no_scopus_id"]])

    def test_a_record_paper_opens_for_the_head_and_nothing_else_does(self):
        pub = Publication.objects.get(title="Alpha paper")
        r = self.c.get(f"/api/hod/papers/{pub.id}")
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        self.assertEqual(body["paper_title"], "Alpha paper")
        self.assertEqual(body["source"], "record")
        self.assertIsNone(body["progress"])
        elsewhere = Publication.objects.get(title="Delta elsewhere")
        self.assertEqual(self.c.get(f"/api/hod/papers/{elsewhere.id}").status_code, 404)


class HeadSeesNoPayment(TestCase):
    """A colleague's record, opened by their head: a paper paid through the ledger
    reads "Completed", never "Paid", and no figure comes with it."""

    def test_a_ledger_paid_paper_reads_completed(self):
        from core.models import PaidLedger

        head = _u("h3@t.edu", "Head", Role.HOD)
        mate = _u("m3@t.edu", "Colleague")
        pub = _record(mate, timezone.localdate().year - 1, "Paid through the ledger", "Q1", doi="10.7/paid")
        row = PaidLedger.objects.create(
            payout_month=date(timezone.localdate().year - 1, 6, 1), faculty_name=mate.name, staff_id="S-1",
            paper_title="Paid through the ledger", amount=FIGURE, raw_json=json.dumps({"DOI": "10.7/paid"}),
        )
        row.publications.add(pub)
        c = Client()
        c.force_login(head)
        r = c.get(f"/api/directory/faculty/{mate.id}")
        self.assertEqual(r.status_code, 200, r.content)
        raw = r.content.decode()
        body = r.json()
        stages = [p["claim"]["stage"] for p in body["papers"] if p["claim"]]
        self.assertEqual(stages, ["Completed"])
        self.assertNotIn('"Paid"', raw)
        self.assertNotIn(str(int(FIGURE)), raw)
