"""The year on one page: what the Principal takes to the governing council.

The questions, in the order she asks them (docs/jtbd/principal.md):

1. Did the college publish more or less than last year, and per teacher?
2. Which departments carry it, and which lag, measured per teacher so a big
   department does not win by size alone?
3. What did the scheme cost, against the budget, and per paper?
4. Is the quality moving (share of papers in Q1/Q2 journals)?
5. Where does NAAC metric 3.3.1 (papers per teacher over five years) stand?

Sources are the two truths `college_totals` already defines: papers from the
publication record (plus recognised claim papers it does not hold), money
from the ledger. Papers are counted by calendar publication year, money by
India's April-March financial year, because that is how each is reported:
NAAC/NIRF count publications by year, the budget is set per financial year.

Teachers are today's active faculty and heads on the roll, per department.
No history of headcount is stored, so earlier years are divided by today's
roll; the brief says so rather than presenting it as exact.
"""
from __future__ import annotations

from core.services.cell_safe import safe_append
import io
from collections import defaultdict
from datetime import date
from typing import Any, Optional
from urllib.parse import quote

from django.db.models import Q

from core.models import Authorship, Budget, Role, User
from core.services import college_totals, pdf_fonts
from core.services.college_totals import NO_DEPARTMENT

TEACHER_ROLES = (Role.FACULTY, Role.HOD)
TOP_QUARTILES = {"Q1", "Q2"}


def fy_label(year: int) -> str:
    """The financial year that starts in April of `year`: 2025 -> "2025-26"."""
    return f"{year}-{str(year + 1)[-2:]}"


def _fy_start(d: date) -> int:
    return d.year if d.month >= 4 else d.year - 1


def _ratio(a: float, b: float, places: int = 2) -> Optional[float]:
    return round(a / b, places) if b else None


def _change(now: float, before: float) -> Optional[float]:
    """Percentage change, or None when there is no base to compare against."""
    return round((now - before) * 100 / before, 1) if before else None


def naac_331_band(per_teacher: Optional[float]) -> int:
    """NAAC 3.3.1's scale for papers per teacher over five years."""
    if not per_teacher:
        return 0
    if per_teacher >= 10:
        return 4
    if per_teacher >= 5:
        return 3
    if per_teacher >= 3:
        return 2
    return 1


def teachers_by_department() -> dict[str, int]:
    canon = college_totals.canonical_departments()
    out: dict[str, int] = defaultdict(int)
    for d in User.objects.filter(active=True, role__in=TEACHER_ROLES).values_list("department", flat=True):
        out[canon(d)] += 1
    return dict(out)




def cached_brief(year: Optional[int] = None) -> dict[str, Any]:
    """The brief for every reader, worked out once per change to the data."""
    from core.services.aggregate_cache import shared as cached

    return cached("reports-brief", {"year": year}, lambda: brief(year))


def brief(year: Optional[int] = None) -> dict[str, Any]:
    today = date.today()
    # The last complete year by default: a year still running compared with a
    # whole one reads as a fall that has not happened.
    year = year or today.year - 1
    years = list(range(year - 4, year + 1))
    fy = fy_label(year)

    teachers = teachers_by_department()
    college_teachers = sum(v for k, v in teachers.items() if k != NO_DEPARTMENT) or sum(teachers.values())

    record = college_totals.papers()
    by_year: dict[int, int] = defaultdict(int)
    by_dept_year: dict[str, dict[int, int]] = defaultdict(lambda: defaultdict(int))
    for p in record:
        if not p["year"]:
            continue
        by_year[p["year"]] += 1
        for d in p["departments"]:
            by_dept_year[d][p["year"]] += 1

    pays = college_totals.payments()
    paid_fy: dict[int, float] = defaultdict(float)
    paid_dept_fy: dict[str, dict[int, float]] = defaultdict(lambda: defaultdict(float))
    for p in pays:
        if not p["month"]:
            continue
        s = _fy_start(p["month"])
        paid_fy[s] += p["amount"]
        paid_dept_fy[p["department"]][s] += p["amount"]

    budgets: dict[tuple[str, Optional[str]], float] = {
        (b.financial_year, (b.department or "").strip() or None): b.amount
        for b in Budget.objects.all()
    }
    # From the same list of papers the counts come from, so "58% in Q1 or Q2"
    # and the list it opens (?quartile=top) can never disagree.
    quart: dict[int, list[int]] = defaultdict(lambda: [0, 0])
    for p in record:
        if p["year"] and p.get("quartile"):
            quart[p["year"]][0] += 1
            quart[p["year"]][1] += p["quartile"] in TOP_QUARTILES

    trend = []
    for y in years:
        qn, qt = quart.get(y, (0, 0))
        trend.append({
            "year": y,
            "papers": by_year.get(y, 0),
            "per_teacher": _ratio(by_year.get(y, 0), college_teachers),
            "top_quartile_share": _ratio(qt * 100, qn, 0),
            "quartile_known": qn,
            "financial_year": fy_label(y),
            "paid": round(paid_fy.get(y, 0), 2),
            "budget": budgets.get((fy_label(y), None)),
        })

    now, prev = trend[-1], trend[-2]
    five_year = sum(t["papers"] for t in trend)

    dept_names = {d for d in teachers if d != NO_DEPARTMENT}
    dept_names |= {d for d in by_dept_year if d != NO_DEPARTMENT}
    departments = []
    for d in dept_names:
        n_t = teachers.get(d, 0)
        ys = by_dept_year.get(d, {})
        papers_now, papers_prev = ys.get(year, 0), ys.get(year - 1, 0)
        paid = round(paid_dept_fy.get(d, {}).get(year, 0), 2)
        five = sum(ys.get(y, 0) for y in years)
        departments.append({
            "department": d,
            "teachers": n_t,
            "papers": papers_now,
            "papers_prev": papers_prev,
            "change": _change(papers_now, papers_prev),
            "per_teacher": _ratio(papers_now, n_t),
            "five_year": five,
            "five_year_per_teacher": _ratio(five, n_t),
            "paid": paid,
            "budget": budgets.get((fy, d)),
            "cost_per_paper": _ratio(paid, papers_now, 0),
        })
    departments.sort(key=lambda r: (-(r["per_teacher"] or 0), -r["papers"], r["department"]))

    unassigned = by_dept_year.get(NO_DEPARTMENT, {}).get(year, 0)
    per5 = _ratio(five_year, college_teachers)
    totals = {
        "papers": now["papers"],
        "papers_prev": prev["papers"],
        "change": _change(now["papers"], prev["papers"]),
        "teachers": college_teachers,
        "per_teacher": now["per_teacher"],
        "per_teacher_prev": prev["per_teacher"],
        "top_quartile_share": now["top_quartile_share"],
        "top_quartile_share_prev": prev["top_quartile_share"],
        "quartile_known": now["quartile_known"],
        "paid": now["paid"],
        "paid_prev": prev["paid"],
        "budget": now["budget"],
        "budget_used": _ratio(now["paid"] * 100, now["budget"] or 0, 0),
        "cost_per_paper": _ratio(now["paid"], now["papers"], 0),
    }
    brief = {
        "year": year,
        "partial": year >= today.year,
        "financial_year": fy,
        "years_available": sorted({y for y in by_year if y} | {today.year}, reverse=True),
        "totals": totals,
        "trend": trend,
        "departments": departments,
        "naac_331": {
            "from": years[0], "to": years[-1], "papers": five_year,
            "teachers": college_teachers, "per_teacher": per5, "band": naac_331_band(per5),
        },
        "unassigned_papers": unassigned,
        "notes": [
            "Papers: the publication record plus recognised claim papers it does not hold, "
            "counted once per paper by calendar publication year. A paper shared by two "
            "departments counts in each department, once for the college.",
            f"Money: every payment on the ledger, by financial year (April to March); "
            f"{year} is read with FY {fy}.",
            "Teachers: today's active faculty and heads on the roll. Earlier years are "
            "divided by today's roll; no headcount history is stored.",
            "NAAC 3.3.1 counts only UGC-CARE journals. This figure counts every paper on "
            "record, so it is an upper bound until the UGC-CARE list is checked.",
            "Q1/Q2 share: of the papers whose journal quartile is recorded.",
        ],
    }
    brief["headline"] = headline(brief)
    brief["push"] = needs_a_push(brief)
    brief["rising"] = rising(brief)
    brief["pack"] = pack_checks(brief)
    return brief


#: A department this small has too few teachers for a rate to mean anything.
MIN_TEACHERS = 3


def needs_a_push(b: dict[str, Any], limit: int = 6) -> list[dict[str, Any]]:
    """Departments to call the head about, worst first, each with the reason.

    Three reasons, all read off the department table so the page, the PDF and
    the Excel agree: no papers at all; fewer papers than last year by a fifth
    or more (from a base of five or more); a rate under half the college's.
    """
    college = b["totals"]["per_teacher"] or 0
    out = []
    for d in b["departments"]:
        if d["teachers"] < MIN_TEACHERS:
            continue
        reasons = []
        if d["papers"] == 0:
            reasons.append("No papers")
        elif d["papers_prev"] >= 5 and d["change"] is not None and d["change"] <= -20:
            reasons.append(f"Down {abs(d['change']):g}% ({d['papers_prev']} to {d['papers']})")
        if d["papers"] > 0 and college and (d["per_teacher"] or 0) < college / 2:
            reasons.append(f"{d['per_teacher']:g} per teacher, under half the college's {college:g}")
        elif d["papers"] == 0 and college:
            reasons.append(f"The college averages {college:g} per teacher")
        if reasons:
            out.append({**d, "reasons": reasons})
    out.sort(key=lambda d: (d["per_teacher"] or 0, -d["teachers"]))
    return out[:limit]


def rising(b: dict[str, Any], limit: int = 3) -> list[dict[str, Any]]:
    """Departments to praise: the biggest gain in papers on last year, from a base of five or more."""
    rows = [d for d in b["departments"]
            if d["teachers"] >= MIN_TEACHERS and d["papers_prev"] >= 5 and (d["change"] or 0) > 0]
    rows.sort(key=lambda d: -(d["change"] or 0))
    return rows[:limit]


def pack_checks(b: dict[str, Any]) -> list[dict[str, Any]]:
    """What to settle before this goes to the council: each check says what is
    wrong and where to fix it, or that it is fine."""
    from core.models import JournalStanding

    t, y = b["totals"], b["year"]
    checks = []

    def add(key, label, ok, detail, to):
        checks.append({"key": key, "label": label, "ok": bool(ok), "detail": detail, "to": to})

    add("partial", f"{y} is a full year", not b["partial"],
        "The year is not over, so every figure is to date and none is compared with last year."
        if b["partial"] else f"{y} is complete.", None)
    add("budget", f"A budget is set for FY {b['financial_year']}", t["budget"] is not None,
        f"Budget {_inr(t['budget'])}; {t['budget_used']:g}% used." if t["budget"] is not None
        else "No budget is set, so the council cannot see spend against it.", "/budget")
    add("department", "Every paper has a department", b["unassigned_papers"] == 0,
        "All papers are in a department." if b["unassigned_papers"] == 0
        else f"{pdf_fonts.group_in(b['unassigned_papers'])} papers of {y} have no department, so no row above holds them.",
        f"/reports/papers?year={y}&department={quote(NO_DEPARTMENT)}")
    known = t.get("quartile_known") or 0
    add("quartile", "Every paper has a journal quartile", known >= t["papers"] * 0.9 if t["papers"] else True,
        f"A quartile is recorded for {pdf_fonts.group_in(known)} of {pdf_fonts.group_in(t['papers'])} papers, "
        f"so the Q1/Q2 share ({t['top_quartile_share'] if t['top_quartile_share'] is not None else 'not known'}%) "
        "is of those only.", f"/reports/papers?year={y}&quartile=none")
    ugc = JournalStanding.objects.filter(source=JournalStanding.Source.UGC_CARE).exists()
    add("ugc", "The UGC-CARE list is loaded", ugc,
        "Loaded." if ugc else "Not loaded, so NAAC 3.3.1 below is an upper bound: it counts every paper, NAAC counts UGC-CARE only.",
        "/accreditation")
    return checks


def _inr(n: Optional[float], sign: str = "₹") -> str:
    """Indian grouping: 3250610 -> "₹32,50,610"."""
    if n is None:
        return "not set"
    s = str(int(round(n)))
    neg = s.startswith("-")
    s = s.lstrip("-")
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts + [tail])
    return ("-" if neg else "") + sign + s


def headline(b: dict[str, Any]) -> str:
    t, y = b["totals"], b["year"]
    parts = [f"In {y} the college published {pdf_fonts.group_in(t['papers'])} papers"]
    if t["change"] is not None and not b.get("partial"):
        word = "up" if t["change"] > 0 else "down" if t["change"] < 0 else "level"
        parts.append(
            f", {word} {abs(t['change']):g}% on {y - 1}" if word != "level" else f", the same as {y - 1}"
        )
    if t["per_teacher"] is not None:
        parts.append(f": {t['per_teacher']:g} per teacher across {pdf_fonts.group_in(t['teachers'])} teachers")
    s = "".join(parts) + "."
    if b.get("partial"):
        s = f"{y} is not over, so these are figures to date. " + s.replace(
            f"In {y} the college published", f"So far in {y} the college has published")
    if t["budget"]:
        s += f" It paid {_inr(t['paid'])} in incentives in FY {b['financial_year']}, {t['budget_used']:g}% of the {_inr(t['budget'])} budget."
    else:
        s += f" It paid {_inr(t['paid'])} in incentives in FY {b['financial_year']}; no budget is set for that year."
    ranked = [d for d in b["departments"] if d["teachers"] >= MIN_TEACHERS and d["per_teacher"] is not None]
    if len(ranked) >= 2:
        top, low = ranked[0], ranked[-1]
        s += (f" {top['department']} leads with {top['per_teacher']:g} papers per teacher;"
              f" {low['department']} is lowest at {low['per_teacher']:g}.")
    return s


# ------------------------------------------------------------ downloads ----

def xlsx(b: dict[str, Any], college: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F2430")

    def sheet(ws, title: str, columns: list[str], rows: list[list[Any]], money_cols=(), widths=None):
        ws.title = title
        safe_append(ws, [college])
        ws["A1"].font = Font(bold=True, size=13)
        safe_append(ws, [f"Research publications and incentive spend, {b['year']} (FY {b['financial_year']})"])
        safe_append(ws, [])
        safe_append(ws, columns)
        for c in ws[4]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            safe_append(ws, r)
        for col in money_cols:
            for row in ws.iter_rows(min_row=5, min_col=col, max_col=col):
                for c in row:
                    c.number_format = pdf_fonts.INR_XLSX
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = "A5"

    t = b["totals"]
    sheet(wb.active, "Summary", ["Measure", str(b["year"]), str(b["year"] - 1)], [
        ["Papers published", t["papers"], t["papers_prev"]],
        ["Teachers on roll (today)", t["teachers"], t["teachers"]],
        ["Papers per teacher", t["per_teacher"], t["per_teacher_prev"]],
        ["Share of papers in Q1/Q2 journals (%)", t["top_quartile_share"], t["top_quartile_share_prev"]],
        ["Incentives paid, FY (₹)", t["paid"], t["paid_prev"]],
        ["Budget, FY (₹)", t["budget"], b["trend"][-2]["budget"]],
        ["Budget used (%)", t["budget_used"], None],
        ["Cost per paper (₹)", t["cost_per_paper"], None],
        [],
        ["Headline", b["headline"]],
    ], widths=[40, 18, 18])
    ws = wb.active
    for r in (9, 10, 12):  # paid, budget, cost per paper
        for col in "BC":
            ws[f"{col}{r}"].number_format = pdf_fonts.INR_XLSX
    ws.merge_cells("B14:C14")
    ws["B14"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[14].height = 110
    sheet(wb.create_sheet(), "Five years", [
        "Year", "Papers", "Papers per teacher", "Q1/Q2 share (%)", "Financial year",
        "Incentives paid (₹)", "Budget (₹)",
    ], [[r["year"], r["papers"], r["per_teacher"], r["top_quartile_share"], r["financial_year"],
         r["paid"], r["budget"]] for r in b["trend"]], money_cols=(6, 7), widths=[8, 10, 12, 12, 12, 16, 16])
    pushed = {p["department"]: p["reasons"] for p in b.get("push", [])}
    sheet(wb.create_sheet(), "Departments", [
        "Department", "Teachers", f"Papers {b['year']}", f"Papers {b['year'] - 1}", "Change (%)",
        "Papers per teacher", "Papers, five years", "Five-year papers per teacher",
        f"Paid FY {b['financial_year']} (₹)", "Budget (₹)", "Cost per paper (₹)", "Needs a push",
    ], [[d["department"], d["teachers"], d["papers"], d["papers_prev"],
         d["change"] if d["change"] is not None else "None",
         d["per_teacher"] if d["teachers"] else "No teachers",
         d["five_year"], d["five_year_per_teacher"] if d["teachers"] else "None", d["paid"],
         d["budget"] if d["budget"] is not None else "Not set",
         d["cost_per_paper"] if d["cost_per_paper"] else "None",
         "; ".join(pushed.get(d["department"], []))]
        for d in b["departments"]], money_cols=(9, 10, 11), widths=[16, 9, 10, 10, 10, 11, 11, 12, 16, 14, 14, 60])
    n = b["naac_331"]
    sheet(wb.create_sheet(), "NAAC 3.3.1", [
        "Year", "Number of research papers", "Full-time teachers",
    ], [[r["year"], r["papers"], b["totals"]["teachers"]] for r in b["trend"]] + [
        [], ["Five-year total", n["papers"], n["teachers"]],
        ["Papers per teacher", n["per_teacher"]], ["NAAC band (0-4)", n["band"]],
    ], widths=[20, 24, 18])
    ws = wb.create_sheet("Notes")
    for line in b["notes"]:
        safe_append(ws, [line])
    ws.column_dimensions["A"].width = 120
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pdf(b: dict[str, Any], college: str) -> bytes:
    """Two A4 pages for the governing council (see core/services/principal_pdf.py)."""
    from core.services import principal_pdf

    return principal_pdf.build(b, college, _inr, _change)


