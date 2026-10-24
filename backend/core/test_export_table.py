"""POST /api/export/table: the table on screen, out as Excel, PDF or CSV.

Each test opens the file the way its reader would (cells out of the workbook,
text out of the PDF, rows out of the CSV) and asserts on what they would see.
"""
import csv
import io
import json
from datetime import datetime
from unittest.mock import patch

import openpyxl
from django.test import TestCase
from django.utils import timezone
from pypdf import PdfReader

from core.models import AuditLog, Role, User
from core.services import exporters

ENDPOINT = "/api/export/table"


def payload(**overrides):
    body = {
        "title": "Faculty output",
        "subtitle": "Saveetha Engineering College · 2026",
        "format": "xlsx",
        "filename": None,
        "columns": [
            {"key": "name", "label": "Name", "kind": "text"},
            {"key": "papers", "label": "Papers", "kind": "number"},
            {"key": "amount", "label": "Amount paid", "kind": "money"},
            {"key": "when", "label": "Paid on", "kind": "date"},
        ],
        "rows": [
            ["Priya Raman", 3, 1234567.5, "2026-03-14"],
            ["Arun K", 1, -9000, "not a date"],
        ],
    }
    body.update(overrides)
    return body


def csv_rows(body: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


class ExportTableTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email="et-fac@t.edu", password="p", name="Priya Raman", role=Role.FACULTY,
            department="CSE", staff_id="ET1",
        )
        self.client.force_login(self.user)

    def post(self, body):
        return self.client.post(ENDPOINT, data=json.dumps(body), content_type="application/json")

    # ------------------------------------------------------------- xlsx ----

    def test_xlsx_has_title_sheet_bold_header_frozen_pane_and_filter(self):
        r = self.post(payload())
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], exporters.CONTENT_TYPES["xlsx"])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Faculty output"])
        ws = wb["Faculty output"]
        self.assertEqual([c.value for c in ws[1]], ["Name", "Papers", "Amount paid", "Paid on"])
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A2"].value, "Priya Raman")
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:D3")

    def test_xlsx_numbers_are_numbers_and_money_keeps_its_grouping(self):
        ws = openpyxl.load_workbook(io.BytesIO(self.post(payload()).content)).active
        self.assertEqual(ws["B2"].value, 3)
        self.assertEqual(ws["B2"].data_type, "n")
        self.assertEqual(ws["C2"].value, 1234567.5)
        self.assertEqual(ws["C2"].data_type, "n")
        self.assertEqual(ws["C2"].number_format, exporters.INR_MONEY_XLSX)
        self.assertEqual(ws["C3"].value, -9000)

    def test_xlsx_iso_dates_become_dates_and_other_text_stays_text(self):
        ws = openpyxl.load_workbook(io.BytesIO(self.post(payload()).content)).active
        self.assertEqual(ws["D2"].value, datetime(2026, 3, 14))
        self.assertEqual(ws["D2"].number_format, "dd-mmm-yyyy")
        self.assertEqual(ws["D3"].value, "not a date")
        self.assertEqual(ws["D3"].data_type, "s")

    def test_xlsx_formula_lookalikes_come_out_as_text(self):
        body = payload(rows=[["=1+1", 1, 1, None], ["@SUM(A1)", 2, 2, None]])
        ws = openpyxl.load_workbook(io.BytesIO(self.post(body).content)).active
        for ref in ("A2", "A3"):
            self.assertEqual(ws[ref].data_type, "s", ref)
        self.assertEqual(ws["A2"].value, "'=1+1")
        self.assertEqual(ws["A3"].value, "'@SUM(A1)")

    def test_sheet_name_is_the_title_made_safe_for_excel(self):
        title = "Ledger: Aug/2026 [draft] a title that runs well past thirty-one characters"
        r = self.post(payload(title=title))
        name = openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames[0]
        self.assertLessEqual(len(name), 31)
        self.assertTrue(name.startswith("Ledger"))
        self.assertFalse(set(name) & set("[]:*?/\\"))

    # -------------------------------------------------------------- csv ----

    def test_csv_starts_with_bom_and_reads_back_with_accents_and_rupees(self):
        body = payload(format="csv", rows=[["Ṕriya ₹", 2, 1, None]])
        r = self.post(body)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], exporters.CONTENT_TYPES["csv"])
        self.assertTrue(r.content.startswith(b"\xef\xbb\xbf"))
        rows = csv_rows(r.content)
        self.assertEqual(rows[0], ["Name", "Papers", "Amount paid", "Paid on"])
        self.assertEqual(rows[1], ["Ṕriya ₹", "2", "1", ""])

    def test_csv_is_formula_safe_and_keeps_negative_numbers_plain(self):
        body = payload(format="csv", rows=[
            ["=1+1", 3, 1234567.5, "2026-03-14"],
            ["@SUM(A1)", 1, -9000, "not a date"],
        ])
        rows = csv_rows(self.post(body).content)
        self.assertEqual(rows[1], ["'=1+1", "3", "1234567.5", "2026-03-14"])
        self.assertEqual(rows[2], ["'@SUM(A1)", "1", "-9000", "not a date"])

    def test_csv_writes_booleans_as_yes_and_no(self):
        body = payload(format="csv", columns=[{"key": "ok", "label": "Approved", "kind": None}],
                       rows=[[True], [False]])
        rows = csv_rows(self.post(body).content)
        self.assertEqual(rows[1:], [["Yes"], ["No"]])

    def test_cell_text_is_cut_to_2000_characters(self):
        rows = csv_rows(self.post(payload(format="csv", rows=[["a" * 2500, 1, 1, None]])).content)
        self.assertEqual(len(rows[1][0]), 2000)

    def test_a_non_scalar_cell_becomes_its_string(self):
        rows = csv_rows(self.post(payload(format="csv", rows=[[{"k": 1}, 1, 1, None]])).content)
        self.assertEqual(rows[1][0], "{'k': 1}")

    # -------------------------------------------------------------- pdf ----

    def test_pdf_is_landscape_with_title_generated_line_repeated_header_and_page_numbers(self):
        rows = [[f"Person {i:03d}", i, 1000.0 + i, "2026-01-01"] for i in range(120)]
        r = self.post(payload(format="pdf", rows=rows))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], exporters.CONTENT_TYPES["pdf"])
        self.assertTrue(r.content.startswith(b"%PDF"))
        pages = PdfReader(io.BytesIO(r.content)).pages
        self.assertGreater(len(pages), 2)
        self.assertGreater(float(pages[0].mediabox.width), float(pages[0].mediabox.height))
        first = pages[0].extract_text()
        for part in ("Faculty output", "Saveetha Engineering College", "generated",
                     "by Priya Raman", f"Page 1 of {len(pages)}"):
            self.assertIn(part, first)
        self.assertIn("Amount paid", pages[1].extract_text())
        self.assertIn(f"Page {len(pages)} of {len(pages)}", pages[-1].extract_text())

    def test_pdf_shows_money_in_indian_grouping_as_on_screen(self):
        r = self.post(payload(format="pdf"))
        text = PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
        self.assertIn("₹12,34,567.50", text)
        self.assertIn("-₹9,000", text)

    def test_pdf_splits_a_wide_table_into_column_groups_and_keeps_every_column(self):
        columns = [{"key": f"f{i}", "label": f"Field{i:02d}", "kind": None} for i in range(1, 61)]
        rows = [[f"value-{r}-{c}" for c in range(60)] for r in range(3)]
        r = self.post(payload(format="pdf", columns=columns, rows=rows))
        self.assertEqual(r.status_code, 200)
        pages = PdfReader(io.BytesIO(r.content)).pages
        text = " ".join(p.extract_text() for p in pages)
        for column in columns:
            self.assertIn(column["label"], text)
        self.assertIn("value-2-59", text)

    def test_zero_rows_still_gives_a_file_in_every_format(self):
        for fmt in ("xlsx", "pdf", "csv"):
            r = self.post(payload(format=fmt, rows=[]))
            self.assertEqual(r.status_code, 200, fmt)
            self.assertTrue(r.content, fmt)

    # ------------------------------------------------------------ file name --

    def test_download_name_is_the_title_slug_and_date_when_no_filename_is_given(self):
        r = self.post(payload())
        expected = f'attachment; filename="faculty-output-{timezone.localdate():%Y%m%d}.xlsx"'
        self.assertEqual(r["Content-Disposition"], expected)

    def test_a_given_filename_is_made_safe_and_keeps_the_right_extension(self):
        r = self.post(payload(format="csv", filename="../../Faculty report 2026.csv"))
        expected = f'attachment; filename="faculty-report-2026-{timezone.localdate():%Y%m%d}.csv"'
        self.assertEqual(r["Content-Disposition"], expected)

    # ------------------------------------------------------------ audit --

    def test_the_audit_row_records_title_format_and_count_but_no_cell_data(self):
        marker = "SECRET-CELL-MARKER-7731"
        body = payload(rows=[[marker, 1, 2, None], ["Arun K", 1, 2, None]])
        self.assertEqual(self.post(body).status_code, 200)
        rows = AuditLog.objects.filter(action="DATA_EXPORT", entity="Table")
        self.assertEqual(rows.count(), 1)
        entry = rows.get()
        self.assertEqual(entry.actor, self.user)
        self.assertEqual(json.loads(entry.detail_json), {"title": "Faculty output", "format": "xlsx", "rows": 2})
        self.assertNotIn(marker, json.dumps(entry.__dict__, default=str))

    # ------------------------------------------------------------ refusals --

    def test_unauthenticated_request_gets_401(self):
        self.client.logout()
        self.assertEqual(self.post(payload()).status_code, 401)

    def test_more_than_20000_rows_is_refused_with_a_clear_message(self):
        r = self.post(payload(format="csv", rows=[["x", 1, 1.0, None]] * 20001))
        self.assertEqual(r.status_code, 400)
        self.assertIn("20,000", r.json()["detail"])

    def test_exactly_20000_rows_is_accepted(self):
        r = self.post(payload(format="csv", rows=[["x", 1, 1.0, None]] * 20000))
        self.assertEqual(r.status_code, 200)

    def test_a_full_20000_row_table_of_ordinary_text_is_accepted(self):
        """Django refuses request bodies over 2.5 MB before any code runs. A full
        table of ordinary text is bigger than that, and must still get through."""
        text_row = ["Department of Computer Science and Engineering", "Priya Raman", 3,
                    1234567.5, "2026-03-14", "Journal of Something Quite Useful"]
        columns = [{"key": f"c{i}", "label": f"Column {i}", "kind": None} for i in range(6)]
        r = self.post(payload(format="csv", columns=columns, rows=[text_row] * 20000))
        self.assertEqual(r.status_code, 200)

    def test_more_than_60_columns_is_refused(self):
        columns = [{"key": f"c{i}", "label": f"C{i}", "kind": None} for i in range(61)]
        r = self.post(payload(columns=columns, rows=[[1] * 61]))
        self.assertEqual(r.status_code, 400)
        self.assertIn("60", r.json()["detail"])

    def test_unknown_format_is_refused(self):
        r = self.post(payload(format="docx"))
        self.assertEqual(r.status_code, 400)
        self.assertIn("xlsx", r.json()["detail"])

    def test_title_must_be_between_1_and_120_characters(self):
        self.assertEqual(self.post(payload(title="")).status_code, 400)
        self.assertEqual(self.post(payload(title="t" * 121)).status_code, 400)

    def test_a_row_that_does_not_match_the_columns_is_refused(self):
        r = self.post(payload(rows=[["Priya Raman", 3]]))
        self.assertEqual(r.status_code, 400)
        self.assertIn("row 1", r.json()["detail"])

    def test_an_unknown_column_kind_is_refused(self):
        columns = [{"key": "a", "label": "A", "kind": "percent"}]
        self.assertEqual(self.post(payload(columns=columns, rows=[[1]])).status_code, 400)

    def test_a_body_that_is_not_json_is_refused_with_400(self):
        r = self.client.post(ENDPOINT, data="not json at all", content_type="application/json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Not a table export", r.json()["detail"])

    def test_a_missing_field_is_refused_and_named(self):
        body = payload()
        del body["format"]
        r = self.post(body)
        self.assertEqual(r.status_code, 400)
        self.assertIn("format", r.json()["detail"])

    def test_a_body_over_the_size_cap_is_refused(self):
        with patch("core.api.export_table.MAX_BODY_BYTES", 1000):
            r = self.post(payload(format="csv", rows=[["x" * 100, 1, 1.0, None]] * 20))
        self.assertEqual(r.status_code, 400)
        self.assertIn("larger than", r.json()["detail"])
