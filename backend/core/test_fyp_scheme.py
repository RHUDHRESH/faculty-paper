"""The final-year project scheme: the roster of teams, and the claim filed on one.

The college decided (2026-09-23) that a student-project conference paper is
not paid under the faculty publication remuneration scheme. It is a scheme of
its own: a fixed amount per team per conference paper, claimed once, by the
team's mentor, from a roster the office imports out of the department's
workbook.

Every workbook here is built in the test with openpyxl, in the shape of the
real one -- a blank first row, the header on row 2, register numbers as the
whole numbers Excel stores them as. The real file carries students' names and
is never copied into the repository.
"""
from __future__ import annotations

import io
import json
from io import StringIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, TestCase

from core.models import AuditLog, Role, Team, TeamMember, User
from core.tests import CONFIRMED

HEADER = [
    "Department", "Team ID", "Name", "Faculty ID",
    "Reg No - 1", "Name - 1", "Reg No - 2", "Name - 2",
    "Reg No - 3", "Name - 3", "Reg No - 4", "Name - 4",
    "Project Title",
]


def roster_row(code, *, faculty_id="TSCH001", mentor="Dr. Kumar A",
               department="Chemical", students=(), title="A project"):
    """One line of the roster, students padded out to the four slots."""
    slots: list = []
    for reg, name in list(students)[:4]:
        slots += [reg, name]
    slots += [None] * (8 - len(slots))
    return [department, code, mentor, faculty_id, *slots, title]


def roster_workbook(rows, *, sheet="25-26") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append([None] * len(HEADER))  # the real file's first row is blank
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _person(email, name, *, role=Role.FACULTY, **extra):
    return User.objects.create_user(
        email=email, password="pass", name=name, role=role, **extra
    )


class RosterImportTests(TestCase):
    """`import_fyp_teams`: the department's workbook, loaded as teams."""

    def setUp(self):
        self.kumar = _person(
            "kumar@test.edu", "Kumar A", department="Chemical", staff_id="TSCH001"
        )

    def _import(self, rows, **kw):
        from core.services.fyp_roster import import_roster, read_roster

        return import_roster(read_roster(io.BytesIO(roster_workbook(rows, **kw))))

    def test_a_team_arrives_with_its_students_and_its_mentor_linked_by_staff_id(self):
        result = self._import([
            roster_row(
                "PR26CH0001",
                students=[(212222210012, "Mohamed Ahamed Meeran"), (212222210013, "Basith Ali A")],
                title="Phytochemical analysis of Canna indica",
            )
        ])

        self.assertEqual(result["created"], 1)
        team = Team.objects.get(code="PR26CH0001")
        self.assertEqual(team.mentor_id, self.kumar.id)
        self.assertEqual(team.mentor_staff_id, "TSCH001")
        self.assertEqual(team.mentor_name, "Dr. Kumar A")
        self.assertEqual(team.department, "Chemical")
        self.assertEqual(team.title, "Phytochemical analysis of Canna indica")
        self.assertIsNotNone(team.imported_at)
        self.assertEqual(
            sorted(team.members.values_list("register_number", "name")),
            [("212222210012", "Mohamed Ahamed Meeran"), ("212222210013", "Basith Ali A")],
        )

    def test_register_numbers_are_digits_whether_excel_gave_an_int_or_a_float(self):
        """Excel stores a register number as a number. As a float it would
        read back "212222210012.0", which is nobody's register number."""
        self._import([
            roster_row("PR26CH0002", students=[(212222210006, "Hemamalini S"),
                                               (212222210009.0, "Kowsika G")])
        ])
        self.assertEqual(
            sorted(TeamMember.objects.values_list("register_number", flat=True)),
            ["212222210006", "212222210009"],
        )

    def test_the_academic_year_is_read_off_the_sheet_name(self):
        self._import([roster_row("PR26CH0003", students=[(1, "A")])], sheet="25-26")
        self.assertEqual(Team.objects.get(code="PR26CH0003").academic_year, "2025-26")

    def test_every_year_sheet_is_read_with_its_own_academic_year(self):
        """One sheet per academic year. Reading only the first would load
        last year's teams and silently leave this year's mentors unable to
        claim."""
        from core.services.fyp_roster import import_roster, read_roster

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet, code in (("24-25", "PR25CH0001"), ("25-26", "PR26CH0001")):
            ws = wb.create_sheet(sheet)
            ws.append([None] * len(HEADER))
            ws.append(HEADER)
            ws.append(roster_row(code, students=[(1, "A")]))
        buf = io.BytesIO()
        wb.save(buf)

        result = import_roster(read_roster(io.BytesIO(buf.getvalue())))
        self.assertEqual(result["created"], 2)
        self.assertEqual(
            dict(Team.objects.values_list("code", "academic_year")),
            {"PR25CH0001": "2024-25", "PR26CH0001": "2025-26"},
        )

    def test_an_unmatched_mentor_keeps_the_raw_id_and_name_and_is_reported(self):
        result = self._import([
            roster_row("PR26ME0029", faculty_id="TSME134", mentor="DR SELVAM",
                       department="MECHANICAL", students=[(212222086001, "SHREERAM RG")]),
        ])
        team = Team.objects.get(code="PR26ME0029")
        self.assertIsNone(team.mentor_id)
        self.assertEqual(team.mentor_staff_id, "TSME134")
        self.assertEqual(team.mentor_name, "DR SELVAM")
        self.assertEqual(
            result["mentors_unmatched"],
            [{"code": "PR26ME0029", "faculty_id": "TSME134", "mentor_name": "DR SELVAM",
              "department": "MECHANICAL"}],
        )

    def test_a_mentor_the_office_linked_by_hand_survives_a_re_import(self):
        """The roster's Faculty ID matched nobody, so the office linked the
        mentor by hand. Loading the same roster again must not undo that."""
        row = roster_row("PR26ME0029", faculty_id="TSME134", mentor="DR SELVAM",
                         students=[(1, "A")])
        self._import([row])
        selvam = _person("selvam@test.edu", "Selvam", staff_id="TSME-134")
        Team.objects.filter(code="PR26ME0029").update(mentor=selvam)

        again = self._import([row])
        self.assertEqual(Team.objects.get(code="PR26ME0029").mentor_id, selvam.id)
        self.assertEqual(again["mentors_unmatched"], [])

    def test_a_changed_faculty_id_is_not_kept_on_the_old_hand_made_link(self):
        self._import([roster_row("PR26ME0029", faculty_id="TSME134", students=[(1, "A")])])
        selvam = _person("selvam@test.edu", "Selvam", staff_id="TSME-134")
        Team.objects.filter(code="PR26ME0029").update(mentor=selvam)

        self._import([roster_row("PR26ME0029", faculty_id="TSME999", students=[(1, "A")])])
        self.assertIsNone(Team.objects.get(code="PR26ME0029").mentor_id)

    def test_the_faculty_id_is_matched_however_it_was_typed(self):
        self._import([roster_row("PR26CH0004", faculty_id=" tsch001 ", students=[(1, "A")])])
        self.assertEqual(Team.objects.get(code="PR26CH0004").mentor_id, self.kumar.id)

    def test_a_second_import_of_the_same_file_changes_nothing(self):
        rows = [roster_row("PR26CH0001", students=[(1, "A"), (2, "B")])]
        first = self._import(rows)
        again = self._import(rows)
        self.assertEqual((first["created"], first["updated"]), (1, 0))
        self.assertEqual((again["created"], again["updated"], again["unchanged"]), (0, 0, 1))
        self.assertEqual(Team.objects.count(), 1)
        self.assertEqual(TeamMember.objects.count(), 2)

    def test_a_re_import_updates_the_team_by_its_id(self):
        self._import([roster_row("PR26CH0001", students=[(1, "A"), (2, "B")], title="Old")])
        result = self._import(
            [roster_row("pr26ch0001", students=[(1, "A"), (3, "C")], title="New")]
        )
        self.assertEqual((result["created"], result["updated"]), (0, 1))
        team = Team.objects.get()
        self.assertEqual(team.title, "New")
        self.assertEqual(
            sorted(team.members.values_list("name", flat=True)), ["A", "C"]
        )

    def test_a_row_with_no_team_id_is_skipped_and_said_so(self):
        result = self._import([
            roster_row("", students=[(1, "Orphan")]),
            roster_row("PR26CH0005", students=[(2, "Kept")]),
        ])
        self.assertEqual(result["created"], 1)
        self.assertEqual(len(result["skipped"]), 1)
        self.assertIn("Team ID", result["skipped"][0])

    def test_a_workbook_without_the_roster_header_is_refused(self):
        from core.services.fyp_roster import RosterError, read_roster

        wb = openpyxl.Workbook()
        wb.active.append(["Something", "Else"])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(RosterError):
            read_roster(io.BytesIO(buf.getvalue()))

    def test_something_that_is_not_a_workbook_is_refused(self):
        from core.services.fyp_roster import RosterError, read_roster

        with self.assertRaises(RosterError):
            read_roster(io.BytesIO(b"Department,Team ID\nx,y\n"))


class RosterCommandTests(TestCase):
    def test_the_command_imports_and_prints_the_counts_and_the_unmatched(self):
        _person("kumar@test.edu", "Kumar A", staff_id="TSCH001")
        import tempfile
        from pathlib import Path

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "roster.xlsx"
            path.write_bytes(roster_workbook([
                roster_row("PR26CH0001", students=[(1, "A")]),
                roster_row("PR26ME0029", faculty_id="TSME134", mentor="DR SELVAM",
                           students=[(2, "B")]),
            ]))
            out = StringIO()
            call_command("import_fyp_teams", str(path), stdout=out)

        text = out.getvalue()
        self.assertIn("created: 2", text)
        self.assertIn("updated: 0", text)
        self.assertIn("mentors unmatched: 1", text)
        self.assertIn("PR26ME0029", text)
        self.assertIn("TSME134", text)
        self.assertEqual(Team.objects.count(), 2)


class RosterUploadTests(TestCase):
    """The office's upload: same importer, over HTTP, behind the office's door."""

    def setUp(self):
        self.office = _person("cell@test.edu", "Research Office", role=Role.RESEARCH_CELL)
        self.faculty = _person("fac@test.edu", "A Faculty", staff_id="TSCH001")
        self.client = Client()

    def _upload(self, content: bytes, name="roster.xlsx"):
        return self.client.post(
            "/api/admin/fyp-teams/import",
            data={"file": SimpleUploadedFile(name, content)},
        )

    def test_the_office_uploads_the_roster_and_gets_the_counts_back(self):
        self.client.force_login(self.office)
        r = self._upload(roster_workbook([
            roster_row("PR26CH0001", students=[(1, "A")]),
            roster_row("PR26ME0029", faculty_id="TSME134", mentor="DR SELVAM",
                       students=[(2, "B")]),
        ]))
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        self.assertEqual((body["created"], body["updated"]), (2, 0))
        self.assertEqual(
            [m["code"] for m in body["mentors_unmatched"]], ["PR26ME0029"]
        )
        self.assertTrue(AuditLog.objects.filter(action="FYP_TEAMS_IMPORT").exists())

    def test_every_office_role_may_upload_and_nobody_else(self):
        content = roster_workbook([roster_row("PR26CH0001", students=[(1, "A")])])
        for role, allowed in (
            (Role.SUPER_ADMIN, True),
            (Role.RESEARCH_CELL, True),
            (Role.RESEARCH_COORDINATOR, True),
            (Role.FACULTY, False),
            (Role.HOD, False),
            (Role.PRINCIPAL, False),
            (Role.FINANCE, False),
        ):
            user = _person(f"{role.lower()}-up@test.edu", role.title(), role=role)
            self.client.force_login(user)
            r = self._upload(content)
            self.assertEqual(r.status_code == 200, allowed, f"{role}: {r.status_code}")

    def test_a_file_that_is_not_the_roster_is_a_400_not_a_500(self):
        self.client.force_login(self.office)
        r = self._upload(b"a,b\n1,2\n", name="probe.csv")
        self.assertEqual(r.status_code, 400, r.content)

    def test_the_office_can_read_what_is_loaded_and_who_is_unmatched(self):
        self.client.force_login(self.office)
        self._upload(roster_workbook([
            roster_row("PR26CH0001", students=[(1, "A")]),
            roster_row("PR26ME0029", faculty_id="TSME134", mentor="DR SELVAM",
                       students=[(2, "B")]),
        ]))
        r = self.client.get("/api/admin/fyp-teams")
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        self.assertEqual(body["teams"], 2)
        self.assertEqual(body["academic_years"], ["2025-26"])
        self.assertEqual([m["code"] for m in body["mentors_unmatched"]], ["PR26ME0029"])

        self.client.force_login(self.faculty)
        self.assertEqual(self.client.get("/api/admin/fyp-teams").status_code, 403)


class TeamWritesAreTheOfficesTests(TestCase):
    """A team decides who may claim ₹15,000, so it is not self-service.

    `POST /teams` let anybody signed in create a team naming themselves as
    its mentor, or re-point an existing team's mentor at themselves -- which
    under the new scheme is a way to mint a claimable team. The roster is the
    office's; so is correcting it.
    """

    def setUp(self):
        self.faculty = _person("self@test.edu", "Self Service")
        self.office = _person("office@test.edu", "Office", role=Role.RESEARCH_CELL)
        self.client = Client()

    def _post(self, **payload):
        body = {"code": "SELF-1", "members": [{"name": "A"}], **payload}
        return self.client.post(
            "/api/teams", data=json.dumps(body), content_type="application/json"
        )

    def test_faculty_cannot_create_a_team(self):
        self.client.force_login(self.faculty)
        r = self._post(mentor_id=self.faculty.id)
        self.assertEqual(r.status_code, 403, r.content)
        self.assertFalse(Team.objects.exists())

    def test_faculty_cannot_take_over_an_imported_team(self):
        team = Team.objects.create(code="PR26CH0001", mentor=self.office)
        self.client.force_login(self.faculty)
        r = self._post(code="PR26CH0001", mentor_id=self.faculty.id)
        self.assertEqual(r.status_code, 403, r.content)
        team.refresh_from_db()
        self.assertEqual(team.mentor_id, self.office.id)

    def test_the_office_can_still_add_a_team_by_hand(self):
        self.client.force_login(self.office)
        r = self._post(mentor_id=self.faculty.id)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(Team.objects.get().mentor_id, self.faculty.id)


# ---------------------------------------------------------------------------
# The claim: filed by the mentor, once per team, on a conference paper, for a
# fixed amount outside the faculty publication formula.
# ---------------------------------------------------------------------------


def _paper_attachment():
    return {
        "kind": "PUBLISHED_PAPER",
        "url": "/media/claims/f0f0f0f0f0f0f0f0f0f0f0f0f0f0f0f0.pdf",
        "filename": "paper.pdf",
        "size_bytes": 20,
    }


class StudentProjectClaimRuleTests(TestCase):
    def setUp(self):
        from core.models import FormulaConfig
        from core.services.remuneration import DEFAULT_AUTHOR_POINTS

        self.cfg = FormulaConfig.objects.create(
            author_point_json=json.dumps(DEFAULT_AUTHOR_POINTS), active=True
        )
        self.mentor = _person(
            "mentor@test.edu", "Kumar A", department="Chemical", staff_id="TSCH001",
            biometric_id="B1", designation="Professor",
        )
        self.other = _person(
            "other@test.edu", "Other Faculty", department="Chemical", staff_id="TSCH099",
            biometric_id="B2", designation="Professor",
        )
        self.team = Team.objects.create(
            code="PR26CH0001", title="Canna indica", department="Chemical",
            academic_year="2025-26", mentor=self.mentor, mentor_name="Dr. Kumar A",
            mentor_staff_id="TSCH001",
        )
        TeamMember.objects.create(team=self.team, name="Meeran", register_number="212222210012")
        TeamMember.objects.create(team=self.team, name="Basith", register_number="212222210013")
        self.second = Team.objects.create(code="PR26CH0002", mentor=self.mentor)
        self.theirs = Team.objects.create(code="PR26CH0003", mentor=self.other)
        self.client = Client()

        from core.tests import _verify_hit, patch_api

        self.verify = patch_api("verify_publication", return_value=_verify_hit())
        self.verify.start()
        self.addCleanup(self.verify.stop)

    def payload(self, **over):
        body = {
            "paper_title": "Phytochemical analysis of Canna indica",
            "journal_title": "Proceedings of a Conference",
            "issn": "1234-5678",
            "publication_date": "2026-03-01",
            "publication_type": "Conference Proceeding",
            "indexing_level": "Scopus",
            "yukthi_id": "YK-1",
            "scopus_author_url": "https://www.scopus.com/authid/detail.uri?authorId=1",
            "affiliation_ok": True,
            "claim_reason": "STUDENT_PROJECT",
            "team_code": "PR26CH0001",
            "total_authors": 5,
            "author_position": 3,
            "attachments": [_paper_attachment()],
            "submit": True,
            "confirmations": CONFIRMED,
        }
        body.update(over)
        return body

    def file(self, user=None, **over):
        self.client.force_login(user or self.mentor)
        return self.client.post(
            "/api/claims", data=json.dumps(self.payload(**over)),
            content_type="application/json",
        )

    # -- the amount ------------------------------------------------------

    def test_the_mentor_is_paid_the_fixed_amount_whatever_their_author_position(self):
        from core.models import Claim, ClaimStatus

        r = self.file()
        self.assertEqual(r.status_code, 200, r.content)
        claim = Claim.objects.get(pk=r.json()["id"])
        self.assertEqual(claim.status, ClaimStatus.SUBMITTED)
        self.assertEqual(claim.remuneration, 15000)
        self.assertEqual(claim.base_amount, 15000)
        self.assertEqual(claim.remuneration_category, "FYP")
        self.assertIsNone(claim.author_point, "the scheme has no author-position split")
        self.assertEqual(claim.qf_amount, 0)
        self.assertIn("final-year project", (claim.remuneration_note or "").lower())

    def test_the_research_cell_can_clear_one_when_scopus_cannot_be_asked(self):
        """The amount is fixed per team; an unreachable (or unconfigured)
        Scopus must not leave the claim unclearable."""
        from unittest.mock import patch

        from core.models import Claim, ClaimStatus

        r = self.file()
        self.assertEqual(r.status_code, 200, r.content)
        claim = Claim.objects.get(pk=r.json()["id"])
        Claim.objects.filter(pk=claim.pk).update(
            snip_source=None, quartile_source=None, manual_verified_at=None
        )
        cell = _person("cell@test.edu", "Cell", role=Role.RESEARCH_CELL)
        self.client.force_login(cell)
        with patch("core.api.journals.verify_publication", return_value={"ok": False}):
            r = self.client.post(f"/api/claims/{claim.id}/recalculate", content_type="application/json")
            self.assertEqual(r.status_code, 200, r.content)
            self.assertEqual(r.json()["remuneration"], 15000)
            r = self.client.post(
                f"/api/claims/{claim.id}/clear",
                data=json.dumps({"expected_amount": 15000}),
                content_type="application/json",
            )
        self.assertEqual(r.status_code, 200, r.content)
        claim.refresh_from_db()
        self.assertEqual(claim.status, ClaimStatus.CLEARED)

    def test_the_amount_is_the_policy_setting_not_a_constant(self):
        from core.models import Claim

        self.cfg.student_project_amount = 20000
        self.cfg.save()
        r = self.file()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(Claim.objects.get(pk=r.json()["id"]).remuneration, 20000)

    def test_no_sec_affiliated_references_are_needed_for_this_scheme(self):
        """The two-reference minimum is the faculty scheme's eligibility rule.
        The payload above attaches the paper and nothing else."""
        r = self.file()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["remuneration"], 15000)

    def test_a_research_faculty_member_is_not_zeroed_by_the_threshold_and_spends_none_of_it(self):
        from datetime import date
        from core.models import Claim, ResearchThreshold

        self.mentor.faculty_type = "RESEARCH"
        self.mentor.save()
        ResearchThreshold.objects.create(user=self.mentor, amount=10_000_000, effective_from=date(2020, 6, 1))
        r = self.file()
        self.assertEqual(r.status_code, 200, r.content)
        claim = Claim.objects.get(pk=r.json()["id"])
        self.assertEqual(claim.remuneration, 15000)
        self.assertFalse(claim.quota_applied)
        self.assertEqual(claim.research_absorbed, 0)

    # -- who may file ----------------------------------------------------

    def test_somebody_who_is_not_the_mentor_is_refused_and_told_whose_team_it_is(self):
        from core.models import Claim

        r = self.file(self.other)
        self.assertEqual(r.status_code, 403, r.content)
        self.assertIn("mentor", r.json()["detail"])
        self.assertIn("Kumar A", r.json()["detail"])
        self.assertFalse(Claim.objects.exists())

    def test_the_office_filing_for_somebody_is_held_to_the_same_rule(self):
        office = _person("cell@test.edu", "Cell", role=Role.RESEARCH_CELL)
        r = self.file(office, owner_id=self.other.id)
        self.assertEqual(r.status_code, 403, r.content)

    def test_a_faculty_member_who_mentors_no_team_is_told_why(self):
        loner = _person("loner@test.edu", "Loner", staff_id="TSX1", biometric_id="B3",
                        designation="Professor")
        r = self.file(loner, team_code="")
        self.assertEqual(r.status_code, 403, r.content)
        self.assertIn("not the mentor of any", r.json()["detail"])

    # -- once per team -----------------------------------------------------

    def test_a_second_claim_on_a_claimed_team_is_refused_naming_the_ticket(self):
        first = self.file()
        self.assertEqual(first.status_code, 200, first.content)
        ticket = first.json()["ticket_number"]
        self.assertTrue(ticket)

        again = self.file(paper_title="Another paper from the same project")
        self.assertEqual(again.status_code, 409, again.content)
        self.assertIn(ticket, again.json()["detail"])

    def test_a_claimed_team_cannot_even_be_put_on_a_draft(self):
        self.file()
        draft = self.file(paper_title="A draft", submit=False)
        self.assertEqual(draft.status_code, 409, draft.content)

    def test_another_of_the_mentor_s_teams_is_still_free(self):
        self.file()
        r = self.file(paper_title="The other project's paper", team_code="PR26CH0002")
        self.assertEqual(r.status_code, 200, r.content)

    def test_withdrawing_the_claim_frees_the_team(self):
        first = self.file().json()
        w = self.client.post(
            f"/api/claims/{first['id']}/withdraw", data=json.dumps({}),
            content_type="application/json",
        )
        self.assertEqual(w.status_code, 200, w.content)
        r = self.file(paper_title="Filed again instead")
        self.assertEqual(r.status_code, 200, r.content)

    def test_a_claim_rejected_outright_frees_the_team(self):
        from core.models import Claim, ClaimStatus

        first = self.file().json()
        Claim.objects.filter(pk=first["id"]).update(
            status=ClaimStatus.REJECTED, rejected_outright=True
        )
        r = self.file(paper_title="A different paper")
        self.assertEqual(r.status_code, 200, r.content)

    def test_a_claim_sent_back_to_fix_still_holds_the_team(self):
        from core.models import Claim, ClaimStatus

        first = self.file().json()
        Claim.objects.filter(pk=first["id"]).update(status=ClaimStatus.REJECTED)
        r = self.file(paper_title="A different paper")
        self.assertEqual(r.status_code, 409, r.content)

    def test_the_withdrawn_claim_cannot_come_back_once_another_took_the_team(self):
        first = self.file().json()
        self.client.post(
            f"/api/claims/{first['id']}/withdraw", data=json.dumps({}),
            content_type="application/json",
        )
        self.assertEqual(self.file(paper_title="Instead").status_code, 200)
        refile = self.client.patch(
            f"/api/claims/{first['id']}", data=json.dumps({"submit": True, "confirmations": CONFIRMED}),
            content_type="application/json",
        )
        self.assertEqual(refile.status_code, 409, refile.content)

    def test_a_status_override_cannot_give_a_team_a_second_filed_claim(self):
        """The rescue for stranded statuses must not be a way round the rule:
        overriding a withdrawn claim back to SUBMITTED while another claim
        holds its team would have both paid."""
        first = self.file().json()
        self.client.post(
            f"/api/claims/{first['id']}/withdraw", data=json.dumps({}),
            content_type="application/json",
        )
        second = self.file(paper_title="Filed instead").json()

        self.client.force_login(_person("rescue@test.edu", "Rescue", role=Role.SUPER_ADMIN))
        r = self.client.post(
            f"/api/admin/claims/{first['id']}/override-status",
            data=json.dumps({"to_status": "SUBMITTED", "note": "Put it back where it was"}),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 409, r.content)
        self.assertIn(second["ticket_number"], r.json()["detail"])

    # -- conference papers only ---------------------------------------------

    def test_a_journal_article_is_refused_because_the_scheme_is_for_conference_papers(self):
        r = self.file(publication_type="Journal")
        self.assertEqual(r.status_code, 400, r.content)
        self.assertIn("conference", r.json()["detail"].lower())

    def test_the_refusal_comes_as_soon_as_the_type_is_known(self):
        r = self.file(publication_type="Book Series", submit=False)
        self.assertEqual(r.status_code, 400, r.content)

    def test_a_draft_may_leave_the_type_for_later_but_cannot_be_filed_without_it(self):
        draft = self.file(publication_type="", submit=False)
        self.assertEqual(draft.status_code, 200, draft.content)
        filed = self.client.patch(
            f"/api/claims/{draft.json()['id']}", data=json.dumps({"submit": True, "confirmations": CONFIRMED}),
            content_type="application/json",
        )
        self.assertEqual(filed.status_code, 400, filed.content)
        self.assertIn("conference", filed.json()["detail"].lower())

    # -- the mentor's picker ---------------------------------------------------

    def test_the_mentor_sees_their_own_teams_with_students_and_what_holds_them(self):
        self.client.force_login(self.mentor)
        mine = self.client.get("/api/teams", {"mine": "true"}).json()["results"]
        self.assertEqual([t["code"] for t in mine], ["PR26CH0001", "PR26CH0002"])
        self.assertEqual(sorted(m["name"] for m in mine[0]["members"]), ["Basith", "Meeran"])
        self.assertIsNone(mine[0]["claimed_by"])

        ticket = self.file().json()["ticket_number"]
        mine = self.client.get("/api/teams", {"mine": "true"}).json()["results"]
        self.assertEqual(mine[0]["claimed_by"]["ticket_number"], ticket)

    def test_the_office_filing_for_a_mentor_is_shown_that_mentor_s_teams(self):
        office = _person("proxy@test.edu", "Proxy", role=Role.RESEARCH_CELL)
        self.client.force_login(office)
        teams = self.client.get(
            "/api/teams", {"mine": "true", "owner_id": self.mentor.id}
        ).json()["results"]
        self.assertEqual([t["code"] for t in teams], ["PR26CH0001", "PR26CH0002"])

    def test_a_claimant_cannot_list_somebody_else_s_teams_that_way(self):
        self.client.force_login(self.other)
        teams = self.client.get(
            "/api/teams", {"mine": "true", "owner_id": self.mentor.id}
        ).json()["results"]
        self.assertEqual([t["code"] for t in teams], ["PR26CH0003"])

    def test_mine_means_mine_even_for_the_office(self):
        office = _person("cell2@test.edu", "Cell", role=Role.RESEARCH_CELL)
        self.client.force_login(office)
        self.assertEqual(self.client.get("/api/teams", {"mine": "true"}).json()["results"], [])

    # -- the estimate, the rules, the policy -----------------------------------

    def test_the_estimate_is_the_fixed_amount_for_a_conference_paper_and_nothing_otherwise(self):
        self.client.force_login(self.mentor)
        body = {"claim_reason": "STUDENT_PROJECT", "publication_type": "Conference Proceeding",
                "total_authors": 5, "author_position": 3}
        r = self.client.post("/api/calculate", data=json.dumps(body),
                             content_type="application/json").json()
        self.assertEqual((r["remuneration"], r["category"]), (15000, "FYP"))

        body["publication_type"] = "Journal"
        r = self.client.post("/api/calculate", data=json.dumps(body),
                             content_type="application/json").json()
        self.assertEqual(r["remuneration"], 0)
        self.assertIn("conference", r["note"].lower())

    def test_the_filing_rules_carry_the_amount_and_the_conference_rule(self):
        self.client.force_login(self.mentor)
        rules = self.client.get("/api/meta/filing-rules").json()
        self.assertEqual(rules["student_project_amount"], 15000)
        self.assertIn("conference", rules["why"]["student_project"].lower())

    def test_the_policy_screen_edits_the_amount(self):
        from core.services.remuneration import DEFAULT_AUTHOR_POINTS

        admin = _person("super@test.edu", "Super", role=Role.SUPER_ADMIN)
        self.client.force_login(admin)
        policy = {
            "snip_multiplier": 55000, "qf_q1": 50000, "qf_q2": 30000, "qf_q3": 15000,
            "qf_q4": 7000, "author_point_json": json.dumps(DEFAULT_AUTHOR_POINTS),
            "student_project_amount": 18000,
        }
        r = self.client.put("/api/admin/formula", data=json.dumps(policy),
                            content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(self.client.get("/api/admin/formula").json()["student_project_amount"], 18000)

        policy["student_project_amount"] = -1
        r = self.client.put("/api/admin/formula", data=json.dumps(policy),
                            content_type="application/json")
        self.assertEqual(r.status_code, 400, r.content)

    # -- reports and the ledger keep the scheme apart ---------------------------

    def test_reports_and_the_ledger_say_which_scheme_paid(self):
        from datetime import date

        from core.models import Claim, ClaimStatus, PaidLedger

        claim = Claim.objects.get(pk=self.file().json()["id"])
        Claim.objects.filter(pk=claim.pk).update(
            status=ClaimStatus.PAID, payout_month=date(2026, 9, 1)
        )
        PaidLedger.objects.create(claim=claim, payout_month=date(2026, 9, 1), amount=15000,
                                  faculty_name="Kumar A")
        PaidLedger.objects.create(payout_month=date(2026, 9, 1), amount=4000,
                                  faculty_name="Somebody Else")

        admin = _person("super2@test.edu", "Super", role=Role.SUPER_ADMIN)
        self.client.force_login(admin)
        categories = {r["key"]: r for r in self.client.get("/api/reports").json()["by_category"]}
        self.assertIn("FYP", categories)
        self.assertIn("final-year project", categories["FYP"]["label"].lower())

        ledger = self.client.get("/api/admin/ledger").json()["results"]
        self.assertEqual(
            sorted((r["faculty_name"], r["scheme"]) for r in ledger),
            [("Kumar A", "FYP"), ("Somebody Else", "FACULTY")],
        )
        only = self.client.get("/api/admin/ledger", {"scheme": "FYP"}).json()
        self.assertEqual([r["faculty_name"] for r in only["results"]], ["Kumar A"])
        self.assertEqual(only["total_amount"], 15000)
        export = self.client.get("/api/admin/ledger/export", {"scheme": "FYP"})
        self.assertEqual(export.status_code, 200)
        text = export.content.decode()
        self.assertIn("scheme", text.splitlines()[0])
        self.assertIn("FYP", text)
