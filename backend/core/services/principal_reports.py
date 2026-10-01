"""The lists behind the Principal's figures.

Every number on the year brief opens the papers or the people it counts. The
counts come from `college_totals.papers()` (the one definition of "a college
paper"), so a figure and the list it opens can never disagree.

- `papers_list`: the papers behind any figure, filtered by year, department
  and quartile, with the college authors of each.
- `department_detail`: one department in a year: how it compares with the
  college, its five years, its people and where they publish.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from typing import Any, Optional

from django.db.models import Q

from core.models import Authorship, Publication, Role, User
from core.services import college_totals, principal_brief
from core.services.college_totals import NO_DEPARTMENT

QUARTILE_ORDER = {"Q1": 0, "Q2": 1, "Q3": 2, "Q4": 3, "": 4}


def _match_quartile(p: dict[str, Any], quartile: Optional[str]) -> bool:
    if not quartile:
        return True
    q = quartile.strip().upper()
    if q == "NONE":
        return not p.get("quartile")
    if q == "TOP":
        return p.get("quartile") in ("Q1", "Q2")
    return p.get("quartile") == q


def papers_list(
    year: Optional[int] = None,
    department: Optional[str] = None,
    quartile: Optional[str] = None,
    q: Optional[str] = None,
    limit: int = 25,
    offset: int = 0,
) -> dict[str, Any]:
    rows = [p for p in college_totals.papers(year, department) if _match_quartile(p, quartile)]

    term = (q or "").strip()
    if term:
        low = term.casefold()
        record_ids = {p["id"] for p in rows if p["source"] == "record"}
        hit_ids = set(
            Publication.objects.filter(id__in=record_ids)
            .filter(Q(title__icontains=term) | Q(venue__icontains=term) | Q(doi__icontains=term)
                    | Q(authorships__user__name__icontains=term))
            .values_list("id", flat=True)
        )
        rows = [
            p for p in rows
            if (p["source"] == "record" and p["id"] in hit_ids)
            or (p["source"] == "claims" and any(
                low in (p.get(k) or "").casefold() for k in ("title", "journal", "author")))
        ]

    rows.sort(key=lambda p: (-(p["year"] or 0), QUARTILE_ORDER.get(p.get("quartile") or "", 4), p["id"]))
    total = len(rows)
    by_quartile = Counter(p.get("quartile") or "none" for p in rows)
    page = rows[offset: offset + limit]

    record_ids = [p["id"] for p in page if p["source"] == "record"]
    pubs = {p.id: p for p in Publication.objects.filter(id__in=record_ids)}
    authors: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for a in (
        Authorship.objects.filter(publication_id__in=record_ids, user__isnull=False)
        .order_by("publication_id", "position")
        .values("publication_id", "user_id", "user__name", "user__department")
    ):
        if all(x["user_id"] != a["user_id"] for x in authors[a["publication_id"]]):
            authors[a["publication_id"]].append(
                {"user_id": a["user_id"], "name": a["user__name"], "department": a["user__department"] or ""})
    # Claim links for the record papers, so a row can open its claim.
    claim_of: dict[str, str] = {}
    for pid, cid in Publication.claims.through.objects.filter(publication_id__in=record_ids).values_list(
        "publication_id", "claim_id"
    ):
        claim_of.setdefault(pid, cid)

    from core.models import Claim

    wanted = set(claim_of.values()) | {p.get("claim_id") for p in page if p.get("claim_id")}
    numbers = dict(Claim.objects.filter(id__in=wanted).values_list("id", "ticket_number"))
    out = []
    for p in page:
        depts = sorted(d for d in p["departments"] if d != NO_DEPARTMENT)
        if p["source"] == "record":
            pub = pubs.get(p["id"])
            out.append({
                "id": p["id"], "source": "record",
                "title": (pub.title if pub else "") or "",
                "journal": (pub.venue if pub else "") or "",
                "year": p["year"], "quartile": p.get("quartile") or None,
                "doi": (pub.doi if pub else None) or None,
                "departments": depts,
                "authors": authors.get(p["id"], []),
                "claim_id": claim_of.get(p["id"]), "claim_no": numbers.get(claim_of.get(p["id"])),
            })
        else:
            out.append({
                "id": p["id"], "source": "claims",
                "title": p.get("title") or "", "journal": p.get("journal") or "",
                "year": p["year"], "quartile": p.get("quartile") or None,
                "doi": p.get("doi"),
                "departments": depts,
                "authors": ([{"user_id": p["author_id"], "name": p["author"], "department": depts[0] if depts else ""}]
                            if p.get("author_id") and p.get("author") else
                            ([{"user_id": None, "name": p["author"], "department": ""}] if p.get("author") else [])),
                "claim_id": p.get("claim_id"), "claim_no": numbers.get(p.get("claim_id")),
            })
    return {
        "total": total, "limit": limit, "offset": offset, "results": out,
        "by_quartile": dict(by_quartile),
        "year": year, "department": department, "quartile": quartile,
    }


def accreditation_summary(year: Optional[int] = None) -> dict[str, Any]:
    """The five years an assessor asks about, in the terms NAAC and NIRF use.

    NAAC 3.3.1 wants papers per teacher over five years (UGC-CARE journals
    only, which this system cannot yet tell). NIRF reads Scopus and Web of
    Science only, takes marks off for retracted papers, and rewards the
    top-quartile share. So each year shows the papers, the rate per teacher,
    how many the record lists in Scopus, how many are in Q1 or Q2 journals, and
    how many carry a retraction signal in the title. The Scopus, quartile and
    retraction columns count the publication record (not claim-only papers),
    and the table says so.
    """
    from core.models import JournalStanding
    from core.services import retraction

    b = principal_brief.cached_brief(year)
    y = b["year"]
    years = [t["year"] for t in b["trend"]]
    rows = (
        Publication.objects.filter(year__in=years)
        .filter(Q(authorships__is_college=True) | Q(authorships__user__isnull=False))
        .distinct()
        .values_list("id", "year", "scopus_indexed", "quartile", "title")
    )
    scopus: Counter = Counter()
    top: Counter = Counter()
    record: Counter = Counter()
    retracted: Counter = Counter()
    flagged_titles: list[dict[str, Any]] = []
    for pid, py, sc, quart, title in rows:
        record[py] += 1
        scopus[py] += bool(sc)
        top[py] += (quart or "").strip().upper() in ("Q1", "Q2")
        hit = retraction.looks_retracted(title)
        if hit:
            retracted[py] += 1
            flagged_titles.append({"id": pid, "year": py, "title": title, "phrase": hit})
    teachers = b["totals"]["teachers"]
    table = []
    for t in b["trend"]:
        py = t["year"]
        table.append({
            "year": py, "papers": t["papers"], "per_teacher": t["per_teacher"],
            "record_papers": record[py], "scopus": scopus[py], "top_quartile": top[py],
            "retraction_signals": retracted[py],
        })
    five = b["naac_331"]
    return {
        "year": y, "years_available": b["years_available"], "teachers": teachers,
        "naac_331": five, "table": table,
        "scopus_share": round(100 * sum(scopus.values()) / max(1, sum(record.values()))),
        "retraction_signals": sum(retracted.values()),
        "retraction_examples": flagged_titles[:5],
        "ugc_list_loaded": JournalStanding.objects.filter(source=JournalStanding.Source.UGC_CARE).exists(),
    }


def papers_xlsx(college: str, scope: str, rows: list[dict[str, Any]]) -> bytes:
    """The list on screen as a sheet with the college's name on top and the
    filter it was cut with, so a page torn out of the pack says what it is."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from core.services import pdf_fonts
    from core.services.cell_safe import safe_append

    wb = Workbook()
    ws = wb.active
    ws.title = "Papers"
    safe_append(ws, [college])
    ws["A1"].font = Font(bold=True, size=13)
    safe_append(ws, [f"Papers: {scope}"])
    safe_append(ws, [f"{pdf_fonts.group_in(len(rows))} papers"])
    safe_append(ws, [])
    head = ["No.", "Title", "Authors", "Department", "Journal", "Year", "Quartile", "DOI", "Claim no."]
    safe_append(ws, head)
    for c in ws[5]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2430")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows, start=1):
        safe_append(ws, [
            i, r["title"], "; ".join(a["name"] for a in r["authors"] if a.get("name")),
            "; ".join(r["departments"]) or "Not recorded", r["journal"] or "Not recorded",
            r["year"] or "Not recorded", r["quartile"] or "Not recorded",
            r["doi"] or "", r.get("claim_no") or "",
        ])
    for col, w in zip("ABCDEFGHI", (6, 70, 34, 22, 40, 8, 10, 30, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def department_detail(name: str, year: Optional[int] = None) -> Optional[dict[str, Any]]:
    b = principal_brief.cached_brief(year)
    want = name.strip().casefold()
    row = next((d for d in b["departments"] if d["department"].casefold() == want), None)
    if row is None:
        return None
    dept = row["department"]
    y = b["year"]
    years = [t["year"] for t in b["trend"]]

    # The people: today's teachers of the department, and their papers by year.
    canon = college_totals.canonical_departments()
    people = [
        p for p in User.objects.filter(active=True, role__in=principal_brief.TEACHER_ROLES)
        .values("id", "name", "designation", "role", "department")
        if canon(p["department"]) == dept
    ]
    ids = [p["id"] for p in people]
    seen: set[tuple[str, str]] = set()
    per_person: dict[str, dict[int, int]] = defaultdict(lambda: defaultdict(int))
    for uid, pid, py in Authorship.objects.filter(user_id__in=ids).values_list(
        "user_id", "publication_id", "publication__year"
    ):
        if (uid, pid) in seen or not py:
            continue
        seen.add((uid, pid))
        per_person[uid][py] += 1
    roll = []
    for p in people:
        ys = per_person.get(p["id"], {})
        roll.append({
            "user_id": p["id"], "name": p["name"], "designation": p["designation"] or "",
            "head": p["role"] == Role.HOD,
            "papers": ys.get(y, 0), "papers_prev": ys.get(y - 1, 0),
            "five_year": sum(ys.get(k, 0) for k in years),
        })
    roll.sort(key=lambda r: (-r["papers"], -r["five_year"], r["name"]))

    # Where the department's papers of the year appear, and their quartiles.
    papers = college_totals.papers(y, dept)
    record_ids = [p["id"] for p in papers if p["source"] == "record"]
    venues = Counter(
        v.strip() for v in Publication.objects.filter(id__in=record_ids).values_list("venue", flat=True) if v and v.strip()
    )
    quartiles = Counter(p.get("quartile") or "none" for p in papers)

    # The department's own five years, by year.
    by_year: dict[int, int] = defaultdict(int)
    for p in college_totals.papers(None, dept):
        if p["year"]:
            by_year[p["year"]] += 1
    teachers = row["teachers"]
    trend = [{"year": k, "papers": by_year.get(k, 0),
              "per_teacher": round(by_year.get(k, 0) / teachers, 2) if teachers else None} for k in years]

    ranked = [d for d in b["departments"] if d["teachers"] and d["per_teacher"] is not None]
    rank = next((i + 1 for i, d in enumerate(ranked) if d["department"] == dept), None)
    return {
        "department": dept, "year": y, "financial_year": b["financial_year"], "partial": b["partial"],
        "row": row, "rank": rank, "of": len(ranked),
        "college": {"per_teacher": b["totals"]["per_teacher"], "papers": b["totals"]["papers"],
                    "top_quartile_share": b["totals"]["top_quartile_share"]},
        "trend": trend,
        "people": roll,
        "silent": sum(1 for r in roll if r["five_year"] == 0),
        "journals": [{"journal": j, "papers": n} for j, n in venues.most_common(8)],
        "quartiles": dict(quartiles),
        "years_available": b["years_available"],
    }
