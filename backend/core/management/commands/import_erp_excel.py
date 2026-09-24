"""Import master tables and historical claims from Publication_Processing_ERP Excel."""
from __future__ import annotations

import json
import secrets
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from core.models import (
    Claim,
    ClaimStatus,
    FacultyMaster,
    PaidLedger,
    PriorImport,
    PriorPayment,
    Role,
    ScimagoJournal,
    SnipSource,
    User,
)
from core.management.commands.rebuild_from_erp import ledger_amount
from core.services.erp_import import find_existing_claim, map_excel_status, stable_ticket
from core.services.normalize import normalize_doi, normalize_title
from core.services.scimago import parse_categories_field


def _snip_upsert(*, print_issn, year, defaults):
    """update_or_create for a table with no unique constraint on the key."""
    row = SnipSource.objects.filter(print_issn=print_issn, year=year).first()
    if row is None:
        return SnipSource.objects.create(print_issn=print_issn, year=year, **defaults)
    for k, v in defaults.items():
        setattr(row, k, v)
    row.save()
    return row


def _month(val):
    """First of the month for a payout column; None when unparseable."""
    if isinstance(val, datetime):
        return val.date().replace(day=1)
    if not val:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(val).strip()[:10], fmt).date().replace(day=1)
        except ValueError:
            continue
    return None


def _cell(row, *keys):
    for k in keys:
        if k in row and row[k] is not None and str(row[k]).strip() != "":
            return row[k]
    return None


def _f(val):
    if val is None or val == "":
        return None
    try:
        return float(str(val).replace(",", "").strip())
    except ValueError:
        return None


def _i(val):
    f = _f(val)
    return int(f) if f is not None else None


def _s(val, n=None, *, drop_na=True):
    """Trimmed string, with placeholder values normalised to None.

    Pass drop_na=False where the sheet asks the user to type "NA" on purpose —
    the annexure reference column instructs exactly that, so stripping it turned
    a correctly-filled cell into a missing one.
    """
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    if drop_na and s.lower() in ("nan", "none", "na", "nil", "n/a"):
        return None
    return s[:n] if n else s


def _id(val, n=None):
    """An identifier from a spreadsheet cell.

    Excel hands back whole numbers as floats, so a Scopus author ID arrived as
    "57306678000.0" — which is not an author ID, and every linkage check
    against it failed. Identifiers are digits, never decimals.
    """
    s = _s(val, None)
    if s is None:
        return None
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s[:n] if n else s


def _truthy(val) -> bool:
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    s = str(val).strip().lower()
    return s in ("1", "true", "yes", "y", "duplicate", "dup")


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(it))]
    for row in it:
        if not any(c is not None and str(c).strip() for c in row):
            continue
        yield {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}


def _pub_year(val) -> int | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.year
    s = str(val).strip()
    m = __import__("re").search(r"(20\d{2}|19\d{2})", s)
    return int(m.group(1)) if m else None


class Command(BaseCommand):
    help = (
        "Import Faculty_Data, Master_List_Accounts, SJR_Data, SNIP_2025, "
        "and historical Claims from Processed / Raw_Data / Accounts"
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx", type=str, help="Path to Publication_Processing_ERP *.xlsx")
        parser.add_argument("--year", type=int, default=2025, help="Scimago/SNIP dataset year")
        parser.add_argument("--skip-sjr", action="store_true")
        parser.add_argument("--skip-snip", action="store_true")
        parser.add_argument("--skip-faculty", action="store_true")
        parser.add_argument("--skip-accounts", action="store_true")
        parser.add_argument("--skip-claims", action="store_true", help="Skip Processed/Raw_Data/Accounts claims")
        parser.add_argument(
            "--claims-only",
            action="store_true",
            help="Only import Processed/Raw_Data/Accounts claims (skip masters)",
        )
        parser.add_argument("--limit", type=int, default=0, help="Limit rows per sheet (0=all)")

    def handle(self, *args, **options):
        path = Path(options["xlsx"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise CommandError("openpyxl required") from e

        self.stdout.write(f"Loading {path} …")
        wb = load_workbook(str(path), read_only=True, data_only=True)
        year = options["year"]
        limit = options["limit"] or 0
        actor = User.objects.filter(role="SUPER_ADMIN").first()
        if actor is None:
            # Every batch records who loaded it (PriorImport.imported_by is not
            # nullable). On a fresh install the faculty accounts are created by
            # this very import, so there is nobody to fall back to -- it used to
            # crash half-way with an IntegrityError instead of saying so.
            raise CommandError(
                "No super admin account exists yet. Create the first administrator "
                "(open /setup, or run manage.py createsuperuser) and run the import again."
            )
        claims_only = options["claims_only"]

        if not claims_only:
            if not options["skip_faculty"] and "Faculty_Data" in wb.sheetnames:
                n = self._import_faculty(wb["Faculty_Data"], limit)
                self.stdout.write(self.style.SUCCESS(f"FacultyMaster: {n}"))

            if not options["skip_accounts"] and "Master_List_Accounts" in wb.sheetnames:
                n = self._import_accounts(wb["Master_List_Accounts"], actor, limit)
                self.stdout.write(self.style.SUCCESS(f"PriorPayment/PaidLedger: {n}"))

            if not options["skip_sjr"] and "SJR_Data" in wb.sheetnames:
                n = self._import_sjr(wb["SJR_Data"], year, limit)
                self.stdout.write(self.style.SUCCESS(f"ScimagoJournal: {n}"))

            if not options["skip_snip"] and "SNIP_2025" in wb.sheetnames:
                n = self._import_snip(wb["SNIP_2025"], year, limit)
                self.stdout.write(self.style.SUCCESS(f"SnipSource: {n}"))

        if not options["skip_claims"]:
            if "Processed" in wb.sheetnames:
                n = self._import_processed(wb["Processed"], limit)
                self.stdout.write(self.style.SUCCESS(f"Claims from Processed: {n}"))
            if "Raw_Data" in wb.sheetnames:
                n = self._import_raw_data(wb["Raw_Data"], limit)
                self.stdout.write(self.style.SUCCESS(f"Claims from Raw_Data (new/filled): {n}"))
            if "Accounts" in wb.sheetnames:
                n = self._import_accounts_claims(wb["Accounts"], actor, limit)
                self.stdout.write(self.style.SUCCESS(f"Accounts -> PAID claims/ledger: {n}"))

        wb.close()
        self.stdout.write(self.style.SUCCESS("Import complete."))

    # ── masters ──────────────────────────────────────────────────────────

    def _import_faculty(self, ws, limit: int) -> int:
        n = 0
        for row in _rows(ws):
            staff_id = _s(_cell(row, "Staff-ID", "Staff ID", "staff_id"), 64)
            name = _s(_cell(row, "Name of the Staff", "Name", "Faculty Name"), 255)
            if not staff_id or not name:
                continue
            FacultyMaster.objects.update_or_create(
                staff_id=staff_id,
                defaults={
                    "department": _s(_cell(row, "Department"), 128),
                    "biometric_id": _id(_cell(row, "Bio-ID", "Biometric ID"), 64),
                    "scopus_author_id": _id(_cell(row, "Scopus ID", "Scopus Author ID"), 64),
                    "name": name,
                    "designation": _s(_cell(row, "Designation"), 128),
                    "email": _s(_cell(row, "Email ID", "Email"), 254),
                    "phone": _s(_cell(row, "Phone Number", "Phone"), 64),
                    "raw_json": json.dumps(
                        {k: str(v)[:200] if v is not None else None for k, v in row.items()}
                    )[:50000],
                },
            )
            n += 1
            if limit and n >= limit:
                break
        return n

    def _import_accounts(self, ws, actor, limit: int) -> int:
        batch = PriorImport.objects.create(
            filename="Master_List_Accounts",
            row_count=0,
            mapping_json="{}",
            imported_by=actor or User.objects.first(),
        )
        n = 0
        for row in _rows(ws):
            title = _s(_cell(row, "Scopus Article Title", "Article Title", "paper_title"))
            if not title:
                continue
            doi = _s(_cell(row, "DOI"), 255)
            # Not the column headed "Amount": in the sheet's newest block that
            # holds the author count, and the payout is the ERP's own working
            # in col27. One rule for both importers (rebuild_from_erp).
            amount = ledger_amount(row)
            if amount is None:
                amount = _f(_cell(row, "amount"))
            payout = _month(_cell(row, "Month"))
            norm_title = normalize_title(title)
            norm_doi = normalize_doi(doi) if doi else None

            # These rows drive the already-paid duplicate warning on every future
            # claim, so importing the workbook twice used to make legitimate new
            # submissions look like duplicates.
            if PriorPayment.objects.filter(
                normalized_title=norm_title, doi=norm_doi, amount_paid=amount
            ).exists():
                continue

            PriorPayment.objects.create(
                faculty_name=_s(_cell(row, "Faculty Name"), 255),
                employee_id=_s(_cell(row, "Faculty ID"), 64),
                paper_title=title,
                normalized_title=norm_title,
                doi=norm_doi,
                issn=_s(_cell(row, "ISSN"), 32),
                journal_title=_s(_cell(row, "Source Title"), 512),
                amount_paid=amount,
                paid_at=timezone.now() if payout else None,
                claim_ref=_s(_cell(row, "Overall S No"), 255),
                raw_json=json.dumps(
                    {k: str(v)[:200] if v is not None else None for k, v in row.items()}
                )[:50000],
                import_batch=batch,
            )
            if payout and amount is not None:
                PaidLedger.objects.create(
                    claim=None,
                    payout_month=payout,
                    department=_s(_cell(row, "Department"), 128),
                    faculty_name=_s(_cell(row, "Faculty Name"), 255),
                    staff_id=_s(_cell(row, "Faculty ID"), 64),
                    biometric_id=_id(_cell(row, "Biometric ID"), 64),
                    paper_title=title,
                    journal_title=_s(_cell(row, "Source Title"), 512),
                    amount=amount,
                    voucher_number=None,
                    raw_json=json.dumps(
                        {k: str(v)[:100] if v is not None else None for k, v in row.items()}
                    )[:20000],
                )
            n += 1
            if limit and n >= limit:
                break
        batch.row_count = n
        batch.save()
        return n

    def _import_sjr(self, ws, year: int, limit: int) -> int:
        """Replace the year's Scimago rows in batches.

        This was an update_or_create per row: a round trip each way, which
        measured 38 rows a minute against Neon -- fourteen hours for one sheet.
        The dump is a full snapshot for a year, so the year is cleared and
        rewritten in batches instead.
        """
        ScimagoJournal.objects.filter(year=year).delete()
        batch: list[ScimagoJournal] = []
        seen: set[str] = set()
        n = 0
        for row in _rows(ws):
            title = _s(_cell(row, "Title"), 512)
            if not title:
                continue
            issn_raw = _s(_cell(row, "Issn", "ISSN"))
            issn = issn_raw.split(",")[0].strip()[:32] if issn_raw else None
            key = (issn or f"TITLE:{title[:24]}")[:32]
            if key in seen:
                continue
            seen.add(key)
            batch.append(
                ScimagoJournal(
                    issn=key,
                    year=year,
                    source_id=_s(_cell(row, "Sourceid", "Source ID"), 64),
                    title=title,
                    eissn=None,
                    sjr=_f(_cell(row, "SJR")),
                    categories_json=json.dumps(
                        parse_categories_field(_s(_cell(row, "Categories")) or "")
                    ),
                )
            )
            n += 1
            if len(batch) >= 1000:
                ScimagoJournal.objects.bulk_create(batch, ignore_conflicts=True)
                batch = []
                self.stdout.write(f"  SJR … {n}")
            if limit and n >= limit:
                break
        if batch:
            ScimagoJournal.objects.bulk_create(batch, ignore_conflicts=True)
        return n

    def _import_snip(self, ws, year: int, limit: int) -> int:
        """Replace the year's SNIP rows in batches, for the same reason."""
        SnipSource.objects.filter(year=year).delete()
        batch: list[SnipSource] = []
        seen: set[str] = set()
        n = 0
        for row in _rows(ws):
            title = _s(_cell(row, "Title"), 512)
            if not title:
                continue
            print_issn = _s(_cell(row, "Print ISSN", "Print_ISSN"), 32)
            e_issn = _s(_cell(row, "E-ISSN", "E_ISSN"), 32)
            source_id = _id(_cell(row, "Scopus Source ID", "Source ID"), 64)
            # print_issn is CharField(32). Truncating to 64 let the 46-character
            # "TITLE:" fallback through, which SQLite accepts and Postgres kills
            # the whole import over — so the bug only ever fired in production.
            key = (print_issn or e_issn or source_id or f"TITLE:{title[:24]}")[:32]
            # The sheet carries one row per subject area, so a source repeats.
            if key in seen:
                continue
            seen.add(key)
            batch.append(
                SnipSource(
                    print_issn=key,
                    year=year,
                    title=title,
                    e_issn=e_issn,
                    snip=_f(_cell(row, "SNIP")),
                    sjr=_f(_cell(row, "SJR")),
                    source_id=source_id,
                )
            )
            n += 1
            if len(batch) >= 1000:
                SnipSource.objects.bulk_create(batch, ignore_conflicts=True)
                batch = []
                self.stdout.write(f"  SNIP … {n}")
            if limit and n >= limit:
                break
        if batch:
            SnipSource.objects.bulk_create(batch, ignore_conflicts=True)
        return n

    # ── historical claims ────────────────────────────────────────────────

    def _resolve_owner(
        self,
        *,
        staff_id: str | None,
        name: str | None,
        department: str | None,
        biometric_id: str | None,
        scopus_author_id: str | None,
        email: str | None = None,
        designation: str | None = None,
    ) -> User:
        """Resolve or create a FACULTY User for the claim owner."""
        user = None
        if email:
            user = User.objects.filter(email__iexact=email.strip()).first()
        if not user and staff_id:
            user = User.objects.filter(staff_id=staff_id).first()
        if not user and staff_id:
            fm = FacultyMaster.objects.filter(staff_id=staff_id).first()
            if fm and fm.email:
                user = User.objects.filter(email__iexact=fm.email.strip()).first()
                if not user:
                    user = User.objects.create_user(
                        email=fm.email.strip().lower(),
                        password=secrets.token_urlsafe(12),
                        name=fm.name or name or staff_id,
                        role=Role.FACULTY,
                        department=fm.department or department,
                        staff_id=fm.staff_id,
                        biometric_id=str(fm.biometric_id) if fm.biometric_id else biometric_id,
                        designation=fm.designation or designation,
                        scopus_author_id=fm.scopus_author_id or scopus_author_id,
                        must_change_password=True,
                    )
                    return user
        if user:
            changed = False
            if staff_id and not user.staff_id:
                user.staff_id = staff_id
                changed = True
            if department and not user.department:
                user.department = department
                changed = True
            if biometric_id and not user.biometric_id:
                user.biometric_id = biometric_id
                changed = True
            if scopus_author_id and not user.scopus_author_id:
                user.scopus_author_id = scopus_author_id
                changed = True
            if changed:
                user.save()
            return user

        # Placeholder faculty (no email in masters)
        sid = staff_id or "UNKNOWN"
        synth = f"erp.{sid.lower().replace(' ', '')}@imported.local"
        existing = User.objects.filter(email=synth).first()
        if existing:
            return existing
        return User.objects.create_user(
            email=synth,
            password=secrets.token_urlsafe(12),
            name=name or sid,
            role=Role.FACULTY,
            department=department,
            staff_id=staff_id,
            biometric_id=biometric_id,
            designation=designation,
            scopus_author_id=scopus_author_id,
            must_change_password=True,
            active=False,
        )

    def _upsert_claim(self, *, sheet: str, sno, defaults: dict, doi: str | None, staff_id: str | None, title: str | None) -> Claim:
        existing = find_existing_claim(doi=doi, staff_id=staff_id, title=title)
        ticket = stable_ticket(sheet, sno)
        if existing:
            for k, v in defaults.items():
                if v is None or v == "":
                    continue
                # Don't clobber richer status with weaker unless PAID forced
                if k == "status" and existing.status == ClaimStatus.PAID and v != ClaimStatus.PAID:
                    continue
                setattr(existing, k, v)
            if not existing.ticket_number:
                # avoid unique collision
                if not Claim.objects.filter(ticket_number=ticket).exclude(pk=existing.pk).exists():
                    existing.ticket_number = ticket
            existing.normalized_title = normalize_title(existing.paper_title)[:512]
            existing.save()
            return existing

        if Claim.objects.filter(ticket_number=ticket).exists():
            ticket = f"{ticket}-{secrets.token_hex(2)}"[:32]
        defaults = {**defaults, "ticket_number": ticket}
        defaults.setdefault("normalized_title", normalize_title(defaults.get("paper_title"))[:512])
        return Claim.objects.create(**defaults)

    def _import_processed(self, ws, limit: int) -> int:
        n = 0
        for row in _rows(ws):
            title = _s(
                _cell(
                    row,
                    "Article Title (Scopus)",
                    "Article Title",
                    "Title of the Paper",
                )
            )
            # Prefer Scopus title column when two "Article Title" exist — openpyxl
            # dict may overwrite; also try Source-adjacent keys already listed.
            if not title:
                # fallback: any non-empty Article Title-like value
                for k, v in row.items():
                    if k and "article title" in k.lower() and _s(v):
                        title = _s(v)
                        break
            if not title:
                continue

            staff_id = _s(_cell(row, "Faculty ID", "Staff ID", "Staff-ID"), 64)
            name = _s(_cell(row, "Faculty Name"), 255)
            department = _s(_cell(row, "Department"), 255)
            biometric_id = _id(_cell(row, "Biometric ID"), 64)
            scopus_author_id = _s(_cell(row, "Scopus ID"), 64)
            doi = normalize_doi(_s(_cell(row, "DOI"), 255))
            raw_status = _s(_cell(row, "Status", "Remarks"))
            status = map_excel_status(raw_status)

            owner = self._resolve_owner(
                staff_id=staff_id,
                name=name,
                department=department,
                biometric_id=biometric_id,
                scopus_author_id=scopus_author_id,
            )

            pub_date = _cell(row, "Publication Date")
            pub_date_s = None
            if isinstance(pub_date, datetime):
                pub_date_s = pub_date.date().isoformat()
            elif pub_date is not None:
                pub_date_s = _s(pub_date, 64)

            rem = _f(_cell(row, "Remuneration", "Amount"))
            snip = _f(_cell(row, "SNIP Value", "SNIP"))
            author_point = _f(_cell(row, "Author Position Points"))
            total_authors = _i(_cell(row, "Total No of Authors")) or 1
            author_position = _i(_cell(row, "Author Position")) or 1

            defaults = {
                "owner": owner,
                "status": status,
                "status_note": raw_status[:255] if raw_status else None,
                "paper_title": title,
                "journal_title": _s(_cell(row, "Source Title", "Journal Name"), 512),
                "issn": _s(_cell(row, "ISSN", "ISSN No"), 32),
                "publication_date": pub_date_s,
                "publication_year": _pub_year(pub_date),
                "publication_type": _s(_cell(row, "Document Type", "Publication Type"), 128),
                "staff_id": staff_id,
                "biometric_id": biometric_id,
                "scopus_author_id": scopus_author_id,
                "doi": doi,
                "eid": _s(_cell(row, "Scopus EID"), 64),
                "scopus_url": _s(_cell(row, "Scopus Article Link")),
                "indexing_status": _s(_cell(row, "Indexing Status"), 64),
                "linkage_status": _s(_cell(row, "Scopus ID Link Status"), 64),
                "quartile": _s(_cell(row, "SJR Quartile", "Journal quartile"), 16),
                "subject_category": _s(_cell(row, "Subject Area"), 255),
                "snip": snip,
                "engineering_class": _s(_cell(row, "Engineering Classification"), 64),
                "total_authors": total_authors,
                "author_position": author_position,
                "author_point": author_point,
                "remuneration": rem,
                "sec_refs": _s(_cell(row, "SEC Reference")),
                "is_student_publication": _truthy(_cell(row, "Student Publication")),
                "affiliation_ok": not (
                    _s(_cell(row, "Affiliation"))
                    and str(_cell(row, "Affiliation")).strip().lower() in ("no", "false", "0")
                ),
                "duplicate_warning": _truthy(_cell(row, "Duplicate Submission")),
                "verification_ok": True,
            }
            if status == ClaimStatus.PAID:
                defaults["paid_at"] = timezone.now()
            if status != ClaimStatus.DRAFT:
                defaults["submitted_at"] = timezone.now()

            sno = _cell(row, "S.No", "S No", "Overall S No")
            self._upsert_claim(
                sheet="PROCESSED",
                sno=sno if sno is not None else n + 1,
                defaults=defaults,
                doi=doi,
                staff_id=staff_id,
                title=title,
            )
            n += 1
            if limit and n >= limit:
                break
            if n % 200 == 0:
                self.stdout.write(f"  Processed claims … {n}")
        return n

    def _import_raw_data(self, ws, limit: int) -> int:
        """Create claims for Raw_Data rows not already imported; fill faculty fields on matches."""
        n = 0
        for row in _rows(ws):
            title = _s(_cell(row, "Title of the Paper", "Article Title", "paper_title"))
            if not title:
                continue
            staff_id = _s(_cell(row, "Staff ID", "Staff-ID", "Faculty ID"), 64)
            email = _s(_cell(row, "Email Address", "Email ID", "Email"), 254)
            name = _s(_cell(row, "Faculty Name", "Name of the Staff"), 255)
            department = _s(_cell(row, "Department"), 255)
            biometric_id = _id(_cell(row, "Biometric ID", "Bio-ID"), 64)
            doi = normalize_doi(_s(_cell(row, "DOI"), 255))

            existing = find_existing_claim(doi=doi, staff_id=staff_id, title=title)
            faculty_fields = {
                "paper_title": title,
                "journal_title": _s(_cell(row, "Journal Name", "Source Title"), 512),
                "issn": _s(_cell(row, "ISSN No", "ISSN"), 32),
                "self_reported_quartile": _s(_cell(row, "Journal quartile", "SJR Quartile"), 32),
                "yukthi_id": _s(_cell(row, "Yukthi ID"), 64),
                "publication_type": _s(_cell(row, "Publication Type"), 128),
                "indexing_level": _s(_cell(row, "Journal Indexing Level"), 128),
                # The column literally instructs "else mark NA", so a stored "NA"
                # is a filled answer, not a blank one.
                "indexing_ref": _s(
                    _cell(
                        row,
                        "If Journal Indexing Level is AU Annexure / UGC care , Provide Ref No , else mark NA",
                        "indexing_ref",
                    ),
                    255,
                    drop_na=False,
                ),
                "proof_url": _s(
                    _cell(row, "Upload Proof (Full length Published Paper as pdf )", "Upload Proof")
                ),
                "sec_proof_url": _s(_cell(row, "Upload reference papers with SEC affiliation?")),
                "sec_refs": _s(
                    _cell(
                        row,
                        "List the reference articles cited with saveetha affiliation?",
                        "Reference Number of Articles cited with saveetha affiliation",
                    )
                ),
                "scopus_author_url": _s(_cell(row, "Author SCOPUS Link")),
                "impact_factor": _s(_cell(row, "Impact Factor"), 64),
                "designation": _s(_cell(row, "Designation"), 128),
                "total_authors": _i(_cell(row, "Total No of Authors")) or 1,
                "author_position": _i(_cell(row, "Author Position")) or 1,
                "snip": _f(_cell(row, "SNIP", "SNIP Value")),
                "staff_id": staff_id,
                "biometric_id": biometric_id,
            }
            pub_date = _cell(row, "Date of Publication", "Publication Date")
            if isinstance(pub_date, datetime):
                faculty_fields["publication_date"] = pub_date.date().isoformat()
                faculty_fields["publication_year"] = pub_date.year
            elif pub_date is not None:
                faculty_fields["publication_date"] = _s(pub_date, 64)
                faculty_fields["publication_year"] = _pub_year(pub_date)

            if existing:
                for k, v in faculty_fields.items():
                    if v is None or v == "":
                        continue
                    cur = getattr(existing, k, None)
                    if cur is None or cur == "" or cur == 0:
                        setattr(existing, k, v)
                existing.save()
                n += 1
            else:
                owner = self._resolve_owner(
                    staff_id=staff_id,
                    name=name,
                    department=department,
                    biometric_id=biometric_id,
                    scopus_author_id=None,
                    email=email,
                    designation=faculty_fields.get("designation"),
                )
                ts = _cell(row, "Timestamp")
                defaults = {
                    **faculty_fields,
                    "owner": owner,
                    "status": ClaimStatus.SUBMITTED,
                    "status_note": "Imported from Raw_Data",
                    "doi": doi,
                    "verification_ok": True,
                    "submitted_at": ts if isinstance(ts, datetime) else timezone.now(),
                }
                sno = n + 1
                self._upsert_claim(
                    sheet="RAW",
                    sno=sno,
                    defaults=defaults,
                    doi=doi,
                    staff_id=staff_id,
                    title=title,
                )
                n += 1

            if limit and n >= limit:
                break
        return n

    def _import_accounts_claims(self, ws, actor, limit: int) -> int:
        n = 0
        payout_fallback = timezone.now().date().replace(day=1)
        for row in _rows(ws):
            title = _s(
                _cell(row, "Article Title (Scopus)", "Article Title", "Scopus Article Title", "Title of the Paper")
            )
            if not title:
                continue
            staff_id = _s(_cell(row, "Faculty ID", "Staff ID"), 64)
            name = _s(_cell(row, "Faculty Name"), 255)
            department = _s(_cell(row, "Department"), 255)
            biometric_id = _id(_cell(row, "Biometric ID"), 64)
            doi = normalize_doi(_s(_cell(row, "DOI"), 255))
            amount = _f(_cell(row, "Amount", "(SNIP * 55000)+QF"))
            qf = _f(_cell(row, "QF Amount"))
            rem = amount if amount is not None else _f(_cell(row, "Remuneration"))
            raw_status = _s(_cell(row, "Status"))
            status = map_excel_status(raw_status, force_paid=True)

            owner = self._resolve_owner(
                staff_id=staff_id,
                name=name,
                department=department,
                biometric_id=biometric_id,
                scopus_author_id=_s(_cell(row, "Scopus ID"), 64),
            )

            defaults = {
                "owner": owner,
                "status": status,
                "status_note": (raw_status or "Accounts")[:255],
                "paper_title": title,
                "journal_title": _s(_cell(row, "Source Title"), 512),
                "issn": _s(_cell(row, "ISSN"), 32),
                "staff_id": staff_id,
                "biometric_id": biometric_id,
                "doi": doi,
                "eid": _s(_cell(row, "Scopus EID"), 64),
                "scopus_url": _s(_cell(row, "Scopus Article Link")),
                "snip": _f(_cell(row, "SNIP Value", "SNIP")),
                "quartile": _s(_cell(row, "SJR Quartile"), 16),
                "engineering_class": _s(_cell(row, "Engineering Classification"), 64),
                "subject_category": _s(_cell(row, "Subject Area"), 255),
                "indexing_status": _s(_cell(row, "Indexing Status"), 64),
                "linkage_status": _s(_cell(row, "Scopus ID Link Status"), 64),
                "total_authors": _i(_cell(row, "Total No of Authors")) or 1,
                "author_position": _i(_cell(row, "Author Position")) or 1,
                "author_point": _f(_cell(row, "Author Position Points")),
                "qf_amount": qf,
                "remuneration": rem,
                "verification_ok": True,
                "paid_at": timezone.now(),
                "submitted_at": timezone.now(),
                "payout_month": payout_fallback,
            }
            sno = _cell(row, "S.No", "S No")
            claim = self._upsert_claim(
                sheet="ACCOUNTS",
                sno=sno if sno is not None else n + 1,
                defaults=defaults,
                doi=doi,
                staff_id=staff_id,
                title=title,
            )

            if rem is not None:
                # One ledger row per claim, full stop. Keying the guard on the
                # import month meant re-running the workbook in a later calendar
                # month wrote a second row and double-counted the payment.
                payout = _month(_cell(row, "Month", "Payout Month", "Paid Month")) or payout_fallback
                exists = PaidLedger.objects.filter(claim=claim).exists()
                if not exists:
                    PaidLedger.objects.create(
                        claim=claim,
                        payout_month=payout,
                        department=department,
                        faculty_name=name,
                        staff_id=staff_id,
                        biometric_id=biometric_id,
                        paper_title=title,
                        journal_title=_s(_cell(row, "Source Title"), 512),
                        amount=rem,
                        voucher_number=None,
                        raw_json=json.dumps(
                            {k: str(v)[:100] if v is not None else None for k, v in row.items()}
                        )[:20000],
                    )
            n += 1
            if limit and n >= limit:
                break
        return n
