"""The tables the college has to file, built from what it already holds.

NAAC and NIRF both want publication data every year, and it is currently
assembled by hand from exports. The columns below are the ones those
frameworks ask for, in their order, so the sheet can be lifted rather than
retyped.

What this deliberately does not do is invent the parts we cannot know:

- NIRF's quality-of-publication metrics rest on citation counts and
  top-percentile journal placement. No citation data exists anywhere in this
  system, so those columns are left out entirely rather than filled with a
  number somebody might submit.
- NAAC 3.4.3 asks whether the journal is on the UGC-CARE list. That is
  answered only once a UGC-CARE list has been loaded; until then the column
  reads "Not checked", which is the truth.

Every sheet is accompanied by a Notes sheet saying where the figures came
from and what was excluded, because a number in a submission needs to be
defensible when somebody asks a year later.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any

from django.db.models import Count, Q, Sum

from core.models import Claim, ClaimStatus, JournalStanding, User
from core.services.scimago import issn_variants

#: NAAC's own column order for metric 3.4.3.
NAAC_343_COLUMNS = [
    "Sl. No.",
    "Title of paper",
    "Name of the author/s",
    "Department of the teacher",
    "Name of journal",
    "Year of publication",
    "ISSN number",
    "Link to the article / paper / abstract of the article",
    "Is it listed in UGC CARE list",
]

NIRF_COLUMNS = [
    "Year of publication",
    "Publications (Scopus)",
    "Publications (Web of Science)",
    "Total publications",
    "Q1",
    "Q2",
    "Q3",
    "Q4",
    "No quartile recorded",
    "Faculty who published",
    "Publications per publishing faculty",
]

DEPARTMENT_COLUMNS = [
    "Department",
    "Publications",
    "Paid claims",
    "Amount paid",
    "Q1",
    "Q2",
    "Q3",
    "Q4",
    "Faculty who published",
]

FACULTY_COLUMNS = [
    "Name",
    "Department",
    "Designation",
    "Staff ID",
    "Publications",
    "Paid claims",
    "Amount paid",
    "Q1",
    "Q2",
    "First-author papers",
]


def _ugc_care_status() -> tuple[bool, dict[str, bool]]:
    """ISSNs on the UGC-CARE list, if one has been loaded."""
    rows = JournalStanding.objects.filter(source=JournalStanding.Source.UGC_CARE)
    if not rows.exists():
        return False, {}
    return True, {r.issn: r.listed for r in rows}


def _quartile_of(claim_quartile: str | None) -> str:
    q = (claim_quartile or "").strip().upper()
    return q if q in {"Q1", "Q2", "Q3", "Q4"} else ""


def build_pack(*, year: int | None, scope) -> dict[str, Any]:
    """Every sheet, as rows, for one publication year (or all of them)."""
    claims = scope.exclude(status=ClaimStatus.DRAFT).select_related("owner")
    if year:
        claims = claims.filter(publication_year=year)
    claims = claims.order_by("owner__department", "owner__name", "-publication_year")

    have_ugc, ugc = _ugc_care_status()

    naac: list[list[Any]] = []
    for i, c in enumerate(claims, start=1):
        listed = "Not checked"
        if have_ugc:
            hit = next((ugc[v] for v in issn_variants(c.issn) if v in ugc), None)
            listed = "Yes" if hit else "No"
        naac.append([
            i,
            c.paper_title or "",
            c.owner.name if c.owner_id else "",
            c.owner.department if c.owner_id else "",
            c.journal_title or "",
            c.publication_year or "",
            c.issn or "",
            c.scopus_url or (f"https://doi.org/{c.doi}" if c.doi else ""),
            listed,
        ])

    # ---- NIRF: counts by year, no citation metrics -----------------------
    nirf: list[list[Any]] = []
    years = sorted(
        {y for y in claims.values_list("publication_year", flat=True) if y}, reverse=True
    )
    for y in years:
        rows = claims.filter(publication_year=y)
        indexed = list(rows.values_list("indexing_level", flat=True))
        scopus = sum(1 for i in indexed if "scopus" in (i or "").lower())
        wos = sum(
            1
            for i in indexed
            if any(t in (i or "").upper() for t in ("SCIE", "ESCI", "SSCI", "WEB OF SCIENCE"))
        )
        quartiles = [_quartile_of(q) for q in rows.values_list("quartile", flat=True)]
        people = len({o for o in rows.values_list("owner_id", flat=True) if o})
        total = rows.count()
        nirf.append([
            y,
            scopus,
            wos,
            total,
            quartiles.count("Q1"),
            quartiles.count("Q2"),
            quartiles.count("Q3"),
            quartiles.count("Q4"),
            sum(1 for q in quartiles if not q),
            people,
            round(total / people, 2) if people else 0,
        ])

    # ---- department and faculty summaries -------------------------------
    dept: list[list[Any]] = []
    for row in (
        claims.values("owner__department")
        .annotate(
            publications=Count("id"),
            paid=Count("id", filter=Q(status=ClaimStatus.PAID)),
            amount=Sum("remuneration", filter=Q(status=ClaimStatus.PAID)),
            q1=Count("id", filter=Q(quartile__iexact="Q1")),
            q2=Count("id", filter=Q(quartile__iexact="Q2")),
            q3=Count("id", filter=Q(quartile__iexact="Q3")),
            q4=Count("id", filter=Q(quartile__iexact="Q4")),
            people=Count("owner_id", distinct=True),
        )
        .order_by("-publications")
    ):
        dept.append([
            row["owner__department"] or "Not recorded",
            row["publications"],
            row["paid"],
            round(row["amount"] or 0, 2),
            row["q1"], row["q2"], row["q3"], row["q4"],
            row["people"],
        ])

    faculty: list[list[Any]] = []
    for row in (
        claims.values(
            "owner_id", "owner__name", "owner__department",
            "owner__designation", "owner__staff_id",
        )
        .annotate(
            publications=Count("id"),
            paid=Count("id", filter=Q(status=ClaimStatus.PAID)),
            amount=Sum("remuneration", filter=Q(status=ClaimStatus.PAID)),
            q1=Count("id", filter=Q(quartile__iexact="Q1")),
            q2=Count("id", filter=Q(quartile__iexact="Q2")),
            first_author=Count("id", filter=Q(author_position=1)),
        )
        .order_by("-publications")
    ):
        faculty.append([
            row["owner__name"] or "",
            row["owner__department"] or "",
            row["owner__designation"] or "",
            row["owner__staff_id"] or "",
            row["publications"],
            row["paid"],
            round(row["amount"] or 0, 2),
            row["q1"], row["q2"],
            row["first_author"],
        ])

    # ---- what a reader needs in order to defend these figures -----------
    total_faculty = User.objects.filter(role="FACULTY", active=True).count()
    notes = [
        ["Generated", date.today().isoformat()],
        ["Publication year", str(year) if year else "All years on record"],
        ["Rows", str(len(naac))],
        ["Source", "Claims filed in the publication-remuneration portal, drafts excluded"],
        ["", ""],
        ["What is counted", "One row per claim. A paper with several SEC authors appears once per author, which is how NAAC 3.4.3 asks for it (per teacher)."],
        ["Quartile", "As verified against the Scimago dump for the year of publication where that year is held, otherwise the nearest year available."],
        ["UGC-CARE column", "Answered only when a UGC-CARE list has been loaded into the system. 'Not checked' means no list is present, not that the journal is absent from it."],
        ["", ""],
        ["Deliberately absent", "Citation counts and top-percentile placement. No citation data is held anywhere in this system, so NIRF's quality-of-publication metrics cannot be produced here and must come from Scopus or Web of Science directly."],
        ["Active faculty on record", str(total_faculty)],
    ]

    return {
        "NAAC 3.4.3": {"columns": NAAC_343_COLUMNS, "rows": naac},
        "NIRF publications": {"columns": NIRF_COLUMNS, "rows": nirf},
        "By department": {"columns": DEPARTMENT_COLUMNS, "rows": dept},
        "By faculty": {"columns": FACULTY_COLUMNS, "rows": faculty},
        "Notes": {"columns": ["Item", "Detail"], "rows": notes},
    }


def pack_workbook(pack: dict[str, Any]) -> bytes:
    """One sheet per table, sized so nobody has to widen a column by hand."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, table in pack.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(table["columns"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in table["rows"]:
            # A leading "=" would be read as a formula by Excel.
            ws.append([
                f"'{v}" if isinstance(v, str) and v.startswith("=") else v for v in row
            ])
        ws.freeze_panes = "A2"
        for column in ws.columns:
            longest = max(
                (len(str(c.value)) for c in column[:200] if c.value is not None),
                default=10,
            )
            ws.column_dimensions[column[0].column_letter].width = min(60, max(12, longest + 2))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
