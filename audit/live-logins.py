"""One workbook: every live account, what it is, and whether we hold its password.

Built from `account-credentials-status.csv`, which is itself read from
production rather than from any local copy. The distinction the sheet is
careful about:

- A password can never be read back. Django stores a PBKDF2 hash, so the only
  passwords in here are ones this system issued and that nobody has changed
  since -- checked against the live `must_change_password` flag, and for the
  handful where the flag alone is ambiguous, checked by actually signing in.
- A blank password is not a missing column. It means the account cannot be
  handed to anybody as it stands and has to be reset first, which is a
  different job from writing a number down.

Tabs are by what the account is for, because "who are the four people who can
release money" is the question this gets opened for.
"""
from __future__ import annotations

import csv
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE = "account-credentials-status.csv"
OUT = "live-logins.xlsx"

#: Plain words for the roles, and the order a reader wants them in.
ROLE_LABELS = [
    ("SUPER_ADMIN", "Super admin — the research cell"),
    ("PRINCIPAL", "Principal"),
    ("FINANCE", "Finance"),
    ("HOD", "Heads of department"),
    ("RESEARCH_CELL", "Research cell"),
    ("FACULTY", "Faculty"),
]

COLUMNS = [
    ("email", "Username (email)", 38),
    ("role_label", "Account type", 30),
    ("name", "Name", 30),
    ("department", "Department", 16),
    ("staff_id", "Staff ID", 12),
    ("active", "Active", 8),
    ("password_if_known", "Password", 20),
    ("password_state", "Is this password current?", 52),
]

HEAD = PatternFill("solid", fgColor="1E293B")
WARN = PatternFill("solid", fgColor="FEF3C7")


def sheet(wb, title, rows):
    ws = wb.create_sheet(title[:31])
    ws.append([label for _, label, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(key, "") for key, _, _ in COLUMNS])
        if not r.get("password_if_known"):
            # The rows somebody will otherwise try to hand out.
            for cell in ws[ws.max_row]:
                cell.fill = WARN
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            if cell.font.name != "Arial":
                cell.font = Font(name="Arial", bold=cell.font.bold, color=cell.font.color)
    ws.freeze_panes = "A2"
    return ws


def main() -> int:
    with open(SOURCE, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))

    labels = dict(ROLE_LABELS)
    for r in rows:
        r["role_label"] = labels.get(r.get("role", ""), r.get("role", ""))

    by_role = defaultdict(list)
    for r in rows:
        by_role[r.get("role", "")].append(r)

    wb = Workbook()
    wb.remove(wb.active)

    # ---- what the reader needs to know before using any of it -----------
    notes = wb.create_sheet("Read this first")
    tally = Counter(r["password_state"] for r in rows)
    lines = [
        ("What this is", ""),
        ("", "Every account on the live system, read from production, with the "
             "password where this system still holds a working one."),
        ("", ""),
        ("Why some passwords are blank", ""),
        ("", "No password can be read back out of the database — they are stored "
             "as PBKDF2 hashes, which is the point of them. A password appears "
             "here only if this system issued it and nobody has changed it since."),
        ("", "A blank one means the account needs a reset before anybody can be "
             "given it. Those rows are shaded."),
        ("", ""),
        ("Accounts", str(len(rows))),
    ]
    for state, count in tally.most_common():
        lines.append((state, str(count)))
    lines += [
        ("", ""),
        ("Handle it like a password list", ""),
        ("", "This file is generated locally and is git-ignored. It should not be "
             "emailed, committed, or put on shared storage."),
    ]
    for a, b in lines:
        notes.append([a, b])
    notes.column_dimensions["A"].width = 52
    notes.column_dimensions["B"].width = 96
    for row in notes.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=(cell.column == 1 and bool(cell.value) and not row[1].value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # ---- staff first, faculty last --------------------------------------
    for role, label in ROLE_LABELS:
        if by_role.get(role):
            sheet(wb, label, sorted(by_role[role], key=lambda r: (r.get("name") or "")))
    leftover = [r for r in rows if r.get("role") not in dict(ROLE_LABELS)]
    if leftover:
        sheet(wb, "Other", leftover)
    sheet(wb, "Everyone", sorted(rows, key=lambda r: (r.get("role") or "", r.get("name") or "")))

    wb.save(OUT)
    print(f"{len(rows)} accounts -> {OUT}")
    for role, label in ROLE_LABELS:
        n = len(by_role.get(role, []))
        if n:
            have = sum(1 for r in by_role[role] if r.get("password_if_known"))
            print(f"  {label:<34} {n:>4} accounts, {have:>4} with a usable password")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
