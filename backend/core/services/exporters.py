"""One table, five ways out.

Every export in this system was a fresh piece of openpyxl code, and every one
of them could produce a workbook and a CSV and nothing else. That is fine for
somebody who is going to do more arithmetic, and wrong for everybody else:
a principal signing a submission wants a page she can read and initial, an
office assembling a NAAC file wants something that pastes into a document with
a covering paragraph, and another system wants JSON.

A "sheet" here is the smallest thing worth exporting: a name, an ordered list
of column headings, and rows of plain values. A "pack" is an ordered mapping
of those. Both the accreditation tables and the ordinary reports reduce to
that shape, so neither has to know anything about file formats.

What each format is for, and what it costs:

- xlsx  every sheet, live column widths, one tab each. The working format.
- csv   one sheet only, because a CSV has no concept of a second one. Asking
        for CSV over a multi-sheet pack takes the first sheet and says so.
- json  the same structure the API returns, for another system to read.
- pdf   laid out for reading and signing, not for re-importing. Wide tables
        are truncated by column, and the page says which columns it dropped
        rather than silently losing them.
- docx  a heading and a real Word table per sheet, so it can be pasted into a
        submission and restyled.
"""
from __future__ import annotations

import csv
import io
import json
from typing import Any, Iterable

#: What a caller may ask for.
FORMATS = ("xlsx", "csv", "json", "pdf", "docx")

CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def render(
    pack: dict[str, dict[str, Any]],
    fmt: str,
    *,
    title: str = "Report",
    subtitle: str = "",
) -> bytes:
    """One pack, in the asked-for shape.

    `pack` maps a sheet name to `{"columns": [...], "rows": [[...], ...]}`.
    """
    if fmt not in FORMATS:
        raise ValueError(f"Unknown format {fmt!r}")
    return {
        "xlsx": _xlsx,
        "csv": _csv,
        "json": _json,
        "pdf": _pdf,
        "docx": _docx,
    }[fmt](pack, title, subtitle)


# ---------------------------------------------------------------- xlsx ----


#: Column headings whose values are rupees. Matched on the heading rather
#: than guessed from the values, because a column of amounts that happens to
#: be all zeroes this month is still a column of amounts.
_MONEY_WORDS = ("amount", "paid", "spent", "value", "remuneration", "total", "\u20b9")

#: ...and headings that are counts. A count formatted as currency is the
#: single most obvious way for an export to look careless.
_COUNT_WORDS = ("count", "papers", "publications", "claims", "people", "number")

_HEADER_FILL = "FF1F2430"
_BAND_FILL = "FFF6F7F9"
_RULE = "FFD9DCE1"


def _column_kind(heading: str, rows: list, index: int) -> str:
    """money | number | text, decided once per column.

    Deciding per *cell* is what produces a column where three figures are
    right-aligned with a rupee sign and the rest are left-aligned text,
    which is the tell that a spreadsheet was generated rather than made.
    """
    lowered = _text(heading).lower()
    if any(w in lowered for w in _MONEY_WORDS) and not any(
        w in lowered for w in _COUNT_WORDS
    ):
        return "money"
    if any(w in lowered for w in _COUNT_WORDS):
        return "number"
    sample = [r[index] for r in rows[:80] if len(r) > index and r[index] not in (None, "")]
    if sample and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in sample):
        return "number"
    return "text"


def _xlsx(pack: dict[str, dict[str, Any]], title: str, subtitle: str) -> bytes:
    """A workbook somebody can open in front of a board without apologising.

    The difference between this and a dump of the same rows is entirely in the
    things a reader never consciously notices: that the numbers line up on
    their last digit, that the header stays put when they scroll, that a
    column of rupees reads as rupees, that the printed version fits the width
    of a page instead of spilling one column onto a second sheet of paper.

    A cover sheet leads, because a spreadsheet emailed on to somebody else
    arrives with no context at all -- and "which year is this?" is the first
    question anybody asks of a table of figures.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin", color=_RULE)
    header_font = Font(bold=True, color="FFFFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    band_fill = PatternFill("solid", fgColor=_BAND_FILL)

    # ---- cover -----------------------------------------------------------
    cover = wb.create_sheet("Cover")
    cover.sheet_view.showGridLines = False
    cover["B2"] = title
    cover["B2"].font = Font(bold=True, size=20)
    cover["B3"] = subtitle
    cover["B3"].font = Font(size=11, color="FF5A6472")
    cover["B5"] = "Contents"
    cover["B5"].font = Font(bold=True, size=12)
    line = 6
    for name, sheet in pack.items():
        cover.cell(row=line, column=2, value=name)
        count = len(sheet["rows"])
        cover.cell(row=line, column=3, value=count).number_format = "#,##0"
        cover.cell(row=line, column=4, value="row" if count == 1 else "rows").font = Font(
            color="FF5A6472"
        )
        line += 1
    cover.column_dimensions["A"].width = 3
    cover.column_dimensions["B"].width = 46
    cover.column_dimensions["C"].width = 12
    cover.column_dimensions["D"].width = 10

    # ---- one tab per sheet ----------------------------------------------
    for name, sheet in pack.items():
        # Excel refuses a tab name over 31 characters or containing []:*?/\.
        safe = name[:31]
        for bad in "[]:*?/\\":
            safe = safe.replace(bad, "-")
        ws = wb.create_sheet(safe)
        ws.sheet_view.showGridLines = False

        columns = list(sheet["columns"])
        rows = list(sheet["rows"])
        kinds = [_column_kind(c, rows, i) for i, c in enumerate(columns)]

        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
            cell.border = Border(bottom=thin)
        ws.row_dimensions[1].height = 26

        for r_index, row in enumerate(rows, start=2):
            ws.append(list(row))
            for c_index, kind in enumerate(kinds, start=1):
                cell = ws.cell(row=r_index, column=c_index)
                # Banding rather than a rule between every row: the eye needs
                # help tracking across a wide table, and forty horizontal
                # lines is not help.
                if r_index % 2 == 0:
                    cell.fill = band_fill
                if kind == "money":
                    cell.number_format = '\u20b9#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif kind == "number":
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(vertical="top", wrap_text=False)

        for i, column in enumerate(columns, start=1):
            widest = max(
                [len(_text(column))]
                + [len(_text(r[i - 1])) for r in rows[:200] if len(r) >= i]
            )
            # Money needs room for the symbol, the separators and the paise
            # that a raw digit count does not know about yet.
            floor = 14 if kinds[i - 1] == "money" else 10
            ws.column_dimensions[get_column_letter(i)].width = min(
                max(widest + 3, floor), 60
            )

        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = (
                f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
            )

        # Printing is not an afterthought here: a principal signs a printed
        # page, and a table that needs two sheets of paper per row to be read
        # is a table nobody signs.
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
        ws.oddHeader.left.text = title
        ws.oddFooter.right.text = "Page &P of &N"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------- csv ----


def _csv(pack: dict[str, dict[str, Any]], title: str, subtitle: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    names = list(pack)
    first = pack[names[0]]
    if len(names) > 1:
        # Said out loud, in the file. A CSV that silently contained one of
        # five tables is how the wrong table gets submitted.
        writer.writerow([f"{names[0]} — CSV holds one table; this pack has {len(names)}."])
        writer.writerow([f"Other tables in this pack: {', '.join(names[1:])}. Use Excel for all of them."])
        writer.writerow([])
    writer.writerow(list(first["columns"]))
    for row in first["rows"]:
        writer.writerow([_text(v) for v in row])
    # Excel on Windows needs the BOM to read UTF-8 in a .csv.
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------- json ----


def _json(pack: dict[str, dict[str, Any]], title: str, subtitle: str) -> bytes:
    payload = {
        "title": title,
        "subtitle": subtitle,
        "tables": [
            {
                "name": name,
                "columns": list(sheet["columns"]),
                # Objects rather than positional arrays: a consumer reading
                # row[7] has to be told what column seven is, and will be
                # wrong the first time a column is inserted.
                "rows": [
                    {str(c): v for c, v in zip(sheet["columns"], row)}
                    for row in sheet["rows"]
                ],
                "row_count": len(sheet["rows"]),
            }
            for name, sheet in pack.items()
        ],
    }
    return json.dumps(payload, indent=2, default=str).encode("utf-8")


# ----------------------------------------------------------------- pdf ----

#: Past this many columns a landscape A4 page stops being readable.
_PDF_MAX_COLUMNS = 8
#: Rows past this are summarised rather than printed — a 3,000-row PDF is
#: not a document anybody reads, and it is a file nobody can email.
_PDF_MAX_ROWS = 400


def _pdf(pack: dict[str, dict[str, Any]], title: str, subtitle: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "cell", parent=styles["BodyText"], fontSize=7, leading=8.5, alignment=TA_LEFT
    )
    head_style = ParagraphStyle(
        "head", parent=cell_style, fontName="Helvetica-Bold", textColor=colors.white
    )
    note_style = ParagraphStyle(
        "note", parent=styles["BodyText"], fontSize=8, textColor=colors.HexColor("#666666")
    )

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
    )
    story: list[Any] = [Paragraph(title, styles["Title"])]
    if subtitle:
        story.append(Paragraph(subtitle, note_style))
    story.append(Spacer(1, 6 * mm))

    for i, (name, sheet) in enumerate(pack.items()):
        if i:
            story.append(PageBreak())
        story.append(Paragraph(name, styles["Heading2"]))

        columns = list(sheet["columns"])
        dropped = columns[_PDF_MAX_COLUMNS:]
        columns = columns[:_PDF_MAX_COLUMNS]
        rows = sheet["rows"]
        shown = rows[:_PDF_MAX_ROWS]

        caveats = []
        if dropped:
            caveats.append(
                f"{len(dropped)} column{'s' if len(dropped) > 1 else ''} not shown here "
                f"({', '.join(_text(c) for c in dropped)}) — the Excel copy has them."
            )
        if len(rows) > len(shown):
            caveats.append(
                f"Showing the first {len(shown):,} of {len(rows):,} rows. "
                "The Excel and CSV copies have every one."
            )
        if caveats:
            story.append(Paragraph(" ".join(caveats), note_style))
        story.append(Spacer(1, 3 * mm))

        if not shown:
            story.append(Paragraph("No rows.", note_style))
            continue

        data = [[Paragraph(_text(c), head_style) for c in columns]]
        for row in shown:
            data.append(
                [Paragraph(_text(v)[:300], cell_style) for v in list(row)[: len(columns)]]
            )
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(table)

    doc.build(story)
    return out.getvalue()


# ---------------------------------------------------------------- docx ----

#: Word tables past this get slow to open and impossible to page through.
_DOCX_MAX_ROWS = 1000


def _docx(pack: dict[str, dict[str, Any]], title: str, subtitle: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9)

    for name, sheet in pack.items():
        doc.add_heading(name, level=1)
        columns = list(sheet["columns"])
        rows = sheet["rows"][:_DOCX_MAX_ROWS]
        if len(sheet["rows"]) > len(rows):
            note = doc.add_paragraph(
                f"Showing the first {len(rows):,} of {len(sheet['rows']):,} rows. "
                "The Excel copy has every one."
            )
            for run in note.runs:
                run.font.size = Pt(8)

        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        for cell, column in zip(table.rows[0].cells, columns):
            cell.text = _text(column)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(8)
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, list(row)[: len(columns)]):
                cell.text = _text(value)[:500]
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def filename(stem: str, fmt: str) -> str:
    return f"{stem}.{fmt}"


def as_pack(name: str, columns: Iterable[Any], rows: Iterable[Iterable[Any]]) -> dict:
    """A one-table pack, for the many callers that only have one."""
    return {name: {"columns": list(columns), "rows": [list(r) for r in rows]}}
