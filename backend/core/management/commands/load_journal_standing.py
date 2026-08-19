"""Load a journal list — Scopus discontinued titles, or UGC-CARE.

Neither list ships with the system, because neither is ours to redistribute:
Elsevier publishes the discontinued-titles workbook on its Scopus support
pages, and the UGC-CARE list is published by the UGC. Point this at whichever
file the college holds and it will read the columns it recognises.

Until a list is loaded, the standing check reports "not checked" rather than
"fine", so nothing here silently vouches for a journal it knows nothing about.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import JournalStanding
from core.services.normalize import normalize_issn

#: Column headings these files use, lowercased. Each list is tried in order.
COLUMNS = {
    "issn": ["issn", "print issn", "p-issn", "issn (print)", "journal issn"],
    "eissn": ["eissn", "e-issn", "online issn", "issn (online)"],
    "title": ["title", "source title", "journal title", "name of the journal"],
    "date": [
        "date of discontinuation",
        "discontinued date",
        "removed on",
        "date of removal",
        "coverage discontinued",
    ],
    "reason": ["reason", "reason for discontinuation", "remarks", "notes"],
}


def _pick(row: dict, names: list[str]) -> str | None:
    for n in names:
        for key, value in row.items():
            if (key or "").strip().lower() == n and str(value or "").strip():
                return str(value).strip()
    return None


def _as_date(text: str | None) -> date | None:
    if not text:
        return None
    text = text.strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%b-%y", "%B %Y", "%Y"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.date() if fmt != "%Y" else date(parsed.year, 12, 31)
        except ValueError:
            continue
    return None


class Command(BaseCommand):
    help = "Load a Scopus discontinued or UGC-CARE journal list from CSV or XLSX."

    def add_arguments(self, parser):
        parser.add_argument("path", help="CSV or XLSX file")
        parser.add_argument(
            "--source",
            required=True,
            choices=[c.value for c in JournalStanding.Source],
        )
        parser.add_argument(
            "--listed",
            action="store_true",
            help=(
                "The file lists journals that ARE recognised (a UGC-CARE list). "
                "Without it, every row is treated as a removal, which is what a "
                "discontinued-titles file contains."
            ),
        )
        parser.add_argument("--sheet", default=None, help="XLSX sheet name")

    def _rows(self, path: str, sheet: str | None):
        if path.lower().endswith((".xlsx", ".xlsm")):
            try:
                from openpyxl import load_workbook
            except ImportError as e:  # pragma: no cover - openpyxl ships with the app
                raise CommandError("openpyxl is needed to read a workbook") from e
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[sheet] if sheet else wb.active
            rows = ws.iter_rows(values_only=True)
            headers = [str(h or "").strip() for h in next(rows)]
            for values in rows:
                yield dict(zip(headers, values))
        else:
            with io.open(path, encoding="utf-8-sig", newline="") as fh:
                yield from csv.DictReader(fh)

    def handle(self, *args, **opts):
        source = opts["source"]
        listed = bool(opts["listed"])
        seen = 0
        written = 0
        no_issn = 0
        batch: list[JournalStanding] = []

        for row in self._rows(opts["path"], opts["sheet"]):
            seen += 1
            raw_issn = _pick(row, COLUMNS["issn"]) or _pick(row, COLUMNS["eissn"])
            issn = normalize_issn(raw_issn) if raw_issn else None
            if not issn:
                no_issn += 1
                continue
            batch.append(
                JournalStanding(
                    source=source,
                    issn=issn,
                    title=(_pick(row, COLUMNS["title"]) or "")[:512] or None,
                    listed=listed,
                    changed_on=None if listed else _as_date(_pick(row, COLUMNS["date"])),
                    reason=(_pick(row, COLUMNS["reason"]) or "")[:255] or None,
                )
            )

        with transaction.atomic():
            # A list is a snapshot: rows that have gone from it are no longer
            # part of it, so the source is replaced rather than merged.
            removed = JournalStanding.objects.filter(source=source).delete()[0]
            for i in range(0, len(batch), 1000):
                JournalStanding.objects.bulk_create(batch[i : i + 1000])
                written += len(batch[i : i + 1000])

        self.stdout.write(
            self.style.SUCCESS(
                f"{written:,} journals loaded for {source} "
                f"(replaced {removed:,}); {seen:,} rows read, {no_issn:,} had no usable ISSN"
            )
        )
        if not listed:
            dated = sum(1 for b in batch if b.changed_on)
            self.stdout.write(
                f"{dated:,} carry a removal date. Rows without one are treated as "
                "removed at an unknown time, which flags every paper in them."
            )
